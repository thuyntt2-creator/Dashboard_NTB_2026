import openpyxl
import gspread
from google.oauth2.service_account import Credentials
import os
import sys
from dotenv import load_dotenv

# Configure output encoding for Vietnamese characters
sys.stdout.reconfigure(encoding='utf-8')

load_dotenv()

JSON_FILE = os.environ.get('GOOGLE_APPLICATION_CREDENTIALS') or r'C:\Users\lap4all\Desktop\Backlog_Automation\credentials.json'
if not os.path.exists(JSON_FILE):
    JSON_FILE = 'credentials.json'

SHEET_ID = os.environ.get('SHEET_ID') or '1j6Xm7JRemUGRSfbL-wc8DMwt7qfR7j79w9q79_snVnU'

# Path to the corrected Excel workbook containing calculations and formulas
paths_to_check = [
    r"c:\Users\lap4all\Desktop\NTB_Bao_Cao_Van_Hanh_Co_Bieu_Do.xlsx",
    r"c:\Users\lap4all\Desktop\NTB_Bao_Cao_Van_Hanh_Corrected.xlsx",
    r"c:\Users\lap4all\Desktop\New folder\NTB_Bao_Cao_Van_Hanh_Co_Bieu_Do.xlsx",
    r"c:\Users\lap4all\Desktop\New folder\NTB_Bao_Cao_Van_Hanh_Corrected.xlsx",
]

EXCEL_PATH = None
latest_time = 0
for path in paths_to_check:
    if os.path.exists(path):
        mtime = os.path.getmtime(path)
        if mtime > latest_time:
            latest_time = mtime
            EXCEL_PATH = path

if EXCEL_PATH is None:
    print("❌ ERROR: Could not find any NTB_Bao_Cao_Van_Hanh excel file!")
    sys.exit(1)

import time
print(f"Opening Excel workbook: {EXCEL_PATH} (last modified: {time.ctime(latest_time)})")
wb_excel = openpyxl.load_workbook(EXCEL_PATH, data_only=False)

print("Connecting to Google Sheets...")
creds = Credentials.from_service_account_file(
    JSON_FILE, 
    scopes=['https://www.googleapis.com/auth/spreadsheets', 'https://www.googleapis.com/auth/drive']
)
gc = gspread.authorize(creds)
sh = gc.open_by_key(SHEET_ID)
print(f"Successfully connected to spreadsheet: '{sh.title}'")

def sync_range(sheet_name, start_row, start_col, end_row, end_col):
    print(f"Syncing range {openpyxl.utils.get_column_letter(start_col)}{start_row}:{openpyxl.utils.get_column_letter(end_col)}{end_row} for sheet '{sheet_name}'...")
    ws_excel = wb_excel[sheet_name]
    ws_gsheet = sh.worksheet(sheet_name)
    
    import datetime
    values = []
    for r in range(start_row, end_row + 1):
        row_vals = []
        for c in range(start_col, end_col + 1):
            val = ws_excel.cell(row=r, column=c).value
            if val is None:
                val = ""
            elif isinstance(val, (datetime.datetime, datetime.date)):
                val = val.strftime('%Y-%m-%d')
            row_vals.append(val)
        values.append(row_vals)
        
    start_cell = f"{openpyxl.utils.get_column_letter(start_col)}{start_row}"
    ws_gsheet.update(range_name=start_cell, values=values, value_input_option='USER_ENTERED')
    print("Range synced successfully!")

# Define the exact ranges to update for each sheet
sync_jobs = [
    # 1. sản lượng
    ('sản lượng', 3, 1, 23, 7),     # A3:G23
    ('sản lượng', 3, 9, 23, 14),    # I3:N23
    ('sản lượng', 27, 1, 27, 1),    # A27
    
    # 2. gtcnew
    ('gtcnew', 4, 1, 24, 10),       # A4:J24
    ('gtcnew', 4, 13, 24, 22),      # M4:V24
    ('gtcnew', 27, 1, 27, 1),       # A27
    
    # 3. gtc + tồn
    ('gtc + tồn', 4, 1, 24, 10),    # A4:J24
    ('gtc + tồn', 4, 11, 24, 19),   # K4:S24
    
    # 4. ca2
    ('ca2', 4, 1, 24, 10),          # A4:J24
    ('ca2', 4, 11, 24, 20),         # K4:T24
    
    # 5. LTC
    ('LTC', 4, 1, 24, 10),          # A4:J24
    ('LTC', 4, 12, 24, 21),         # L4:U24
    
    # 6. ODR
    ('ODR', 3, 1, 8, 6),            # A3:F8
    ('ODR', 3, 8, 8, 13),           # H3:M8
    ('ODR', 17, 8, 22, 13),         # H17:M22
    
    # 7. Chart Grid Fixes & Text Clears
    ('gtcnew', 35, 1, 35, 3),       # Clear gtcnew A35:C35
    ('gtc + tồn', 31, 1, 36, 3),    # Sync gtc + tồn A31:C36
    ('ca2', 31, 1, 36, 3),          # Sync ca2 A31:C36
    ('LTC', 30, 1, 35, 3),          # Sync LTC A30:C35
]

for sheet_name, start_row, start_col, end_row, end_col in sync_jobs:
    try:
        sync_range(sheet_name, start_row, start_col, end_row, end_col)
    except Exception as e:
        print(f"❌ Error syncing {sheet_name}: {e}")

print("="*60)
print("🎉 GOOGLE SHEET SYNC COMPLETED SUCCESSFULLY!")
print("="*60)
