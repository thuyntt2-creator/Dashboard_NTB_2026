import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference

file_path = r"c:\Users\lap4all\Desktop\New folder\NTB_Bao_Cao_Van_Hanh_Corrected.xlsx"
wb = openpyxl.load_workbook(file_path)

# 1. Chart for sản lượng (Volume)
print("Adding native Volume chart...")
ws_sl = wb['sản lượng']
# Write support table at A29
ws_sl['A29'] = "Bảng dữ liệu biểu đồ"
ws_sl['A30'] = "Tuần"
ws_sl['B30'] = "Tổng vùng NTB"
ws_sl['C30'] = "Tuyến TTS"

weeks_cols = [("W21", "B", "J"), ("W22", "C", "K"), ("W23", "D", "L"), ("W24", "E", "M")]
for idx, (wk, col_tot, col_tts) in enumerate(weeks_cols, start=31):
    ws_sl[f'A{idx}'] = wk
    ws_sl[f'B{idx}'] = f"={col_tot}23"
    ws_sl[f'C{idx}'] = f"={col_tts}23"

chart_sl = BarChart()
chart_sl.type = "col"
chart_sl.style = 10
chart_sl.title = "SẢN LƯỢNG GIAO TỔNG & TTS VÙNG NTB (W21 - W24)"
chart_sl.y_axis.title = "Sản lượng (đơn)"
chart_sl.x_axis.title = "Tuần"

data_sl = Reference(ws_sl, min_col=2, min_row=30, max_col=3, max_row=34)
cats_sl = Reference(ws_sl, min_col=1, min_row=31, max_row=34)
chart_sl.add_data(data_sl, titles_from_data=True)
chart_sl.set_categories(cats_sl)
chart_sl.width = 16
chart_sl.height = 10
ws_sl.add_chart(chart_sl, "A36")


# 2. Chart for GTC (gtcnew)
print("Adding native GTC chart...")
ws_gtc = wb['gtcnew']
ws_gtc['A29'] = "Bảng dữ liệu biểu đồ"
ws_gtc['A30'] = "Tuần"
ws_gtc['B30'] = "GTC Tổng"
ws_gtc['C30'] = "GTC TTS"

gtc_cols = [("W21", "C", "O"), ("W22", "E", "Q"), ("W23", "G", "S"), ("W24", "I", "U")]
for idx, (wk, col_tot, col_tts) in enumerate(gtc_cols, start=31):
    ws_gtc[f'A{idx}'] = wk
    ws_gtc[f'B{idx}'] = f"={col_tot}24"
    ws_gtc[f'C{idx}'] = f"={col_tts}24"

chart_gtc = LineChart()
chart_gtc.title = "XU HƯỚNG TỈ LỆ GIAO THÀNH CÔNG MỚI (W21 - W24)"
chart_gtc.y_axis.title = "Tỉ lệ (%)"
chart_gtc.x_axis.title = "Tuần"

data_gtc = Reference(ws_gtc, min_col=2, min_row=30, max_col=3, max_row=34)
cats_gtc = Reference(ws_gtc, min_col=1, min_row=31, max_row=34)
chart_gtc.add_data(data_gtc, titles_from_data=True)
chart_gtc.set_categories(cats_gtc)
chart_gtc.width = 16
chart_gtc.height = 10
ws_gtc.add_chart(chart_gtc, "A36")


# 3. Chart for LTC (LTC)
print("Adding native LTC chart...")
ws_ltc = wb['LTC']
ws_ltc['A29'] = "Bảng dữ liệu biểu đồ"
ws_ltc['A30'] = "Tuần"
ws_ltc['B30'] = "LTC Tổng"
ws_ltc['C30'] = "LTC TTS"

ltc_cols = [("W21", "C", "N"), ("W22", "E", "P"), ("W23", "G", "R"), ("W24", "I", "T")]
for idx, (wk, col_tot, col_tts) in enumerate(ltc_cols, start=31):
    ws_ltc[f'A{idx}'] = wk
    ws_ltc[f'B{idx}'] = f"={col_tot}24"
    ws_ltc[f'C{idx}'] = f"={col_tts}24"

chart_ltc = LineChart()
chart_ltc.title = "XU HƯỚNG TỈ LỆ LẤY THÀNH CÔNG (W21 - W24)"
chart_ltc.y_axis.title = "Tỉ lệ (%)"
chart_ltc.x_axis.title = "Tuần"

data_ltc = Reference(ws_ltc, min_col=2, min_row=30, max_col=3, max_row=34)
cats_ltc = Reference(ws_ltc, min_col=1, min_row=31, max_row=34)
chart_ltc.add_data(data_ltc, titles_from_data=True)
chart_ltc.set_categories(cats_ltc)
chart_ltc.width = 16
chart_ltc.height = 10
ws_ltc.add_chart(chart_ltc, "A36")


# 4. Chart for ODR (ODR)
print("Adding native ODR chart...")
ws_odr = wb['ODR']
ws_odr['A25'] = "Bảng dữ liệu biểu đồ"
ws_odr['A26'] = "Tuần"
ws_odr['B26'] = "Bình Thuận"
ws_odr['C26'] = "Khánh Hòa"
ws_odr['D26'] = "Lâm Đồng"
ws_odr['E26'] = "Ninh Thuận"
ws_odr['F26'] = "Đắk Nông"

odr_cols = ["B", "C", "D", "E"]
for idx, wk in enumerate(["W21", "W22", "W23", "W24"]):
    row_idx = 27 + idx
    col_letter = odr_cols[idx]
    ws_odr[f'A{row_idx}'] = wk
    ws_odr[f'B{row_idx}'] = f"={col_letter}3"
    ws_odr[f'C{row_idx}'] = f"={col_letter}4"
    ws_odr[f'D{row_idx}'] = f"={col_letter}5"
    ws_odr[f'E{row_idx}'] = f"={col_letter}6"
    ws_odr[f'F{row_idx}'] = f"={col_letter}7"

chart_odr = LineChart()
chart_odr.title = "XU HƯỚNG ODR THEO TỈNH VÙNG NTB (W21 - W24)"
chart_odr.y_axis.title = "Tỉ lệ (%)"
chart_odr.x_axis.title = "Tuần"

data_odr = Reference(ws_odr, min_col=2, min_row=26, max_col=6, max_row=30)
cats_odr = Reference(ws_odr, min_col=1, min_row=27, max_row=30)
chart_odr.add_data(data_odr, titles_from_data=True)
chart_odr.set_categories(cats_odr)
chart_odr.width = 16
chart_odr.height = 10
ws_odr.add_chart(chart_odr, "A32")

# Save workbook
wb.save(file_path)
print("Saved corrected workbook with native Excel charts!")
