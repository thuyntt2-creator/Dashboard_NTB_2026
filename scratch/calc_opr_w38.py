import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W38_2026.xlsx', data_only=True)
ws = wb['08_OPR TTS']
ams = []
for r in range(6, 24):
    name = ws.cell(r, 1).value
    if name:
        v_day = ws.cell(r, 2).value or 0
        w37_d = ws.cell(r, 3).value or 0
        w38_d = ws.cell(r, 4).value or 0
        v_night = ws.cell(r, 6).value or 0
        w37_n = ws.cell(r, 7).value or 0
        w38_n = ws.cell(r, 8).value or 0
        v_tot = ws.cell(r, 10).value or 0
        w37_tot = ws.cell(r, 11).value or 0
        w38_tot = ws.cell(r, 12).value or 0
        diff_tot = ws.cell(r, 13).value or 0
        ams.append({
            'am': name, 'v_day': v_day, 'w37_d': w37_d, 'w38_d': w38_d,
            'v_night': v_night, 'w37_n': w37_n, 'w38_n': w38_n,
            'v_tot': v_tot, 'w37_tot': w37_tot, 'w38_tot': w38_tot, 'diff_tot': diff_tot
        })

tot_day = sum(a['v_day'] for a in ams)
tot_night = sum(a['v_night'] for a in ams)
tot_all = sum(a['v_tot'] for a in ams)
opr_w38_day = sum(a['v_day'] * a['w38_d'] for a in ams) / tot_day
opr_w37_day = sum(a['v_day'] * a['w37_d'] for a in ams) / tot_day
opr_w38_night = sum(a['v_night'] * a['w38_n'] for a in ams) / tot_night
opr_w37_night = sum(a['v_night'] * a['w37_n'] for a in ams) / tot_night
opr_w38_tot = sum(a['v_tot'] * a['w38_tot'] for a in ams) / tot_all
opr_w37_tot = sum(a['v_tot'] * a['w37_tot'] for a in ams) / tot_all

print(f"Total đơn TTS: {tot_all} (Day: {tot_day}, Night: {tot_night})")
print(f"OPR W37: Tot={opr_w37_tot:.2%}, Day={opr_w37_day:.2%}, Night={opr_w37_night:.2%}")
print(f"OPR W38: Tot={opr_w38_tot:.2%}, Day={opr_w38_day:.2%}, Night={opr_w38_night:.2%}, Diff={(opr_w38_tot-opr_w37_tot):+.2%}")
failed = [a for a in ams if a['w38_tot'] < 0.80]
print(f"Failed AMs ({len(failed)}/{len(ams)}):")
for f in sorted(failed, key=lambda x: x['w38_tot']):
    tre = round(f['v_tot'] * (1 - f['w38_tot']))
    print(f"  {f['am']}: {f['w38_tot']:.1%} (Trễ: {tre} đơn)")
