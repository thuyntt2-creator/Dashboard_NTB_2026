import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document("KICH_BAN_THUYET_TRINH_W38_INSIGHT_CHUYEN_SAU.docx")
print("Total paragraphs in doc:", len(doc.paragraphs))
print("Total tables in doc:", len(doc.tables))

for i, p in enumerate(doc.paragraphs[:20]):
    print(f"P{i}: {p.text[:60]}")

print("\n--- TABLES INSPECT ---")
for i, t in enumerate(doc.tables):
    cell = t.rows[0].cells[0]
    first_p = cell.paragraphs[0].text if cell.paragraphs else ""
    second_p = cell.paragraphs[1].text if len(cell.paragraphs) > 1 else ""
    print(f"Table {i+1}: {first_p[:40]} | {second_p[:40]}")
