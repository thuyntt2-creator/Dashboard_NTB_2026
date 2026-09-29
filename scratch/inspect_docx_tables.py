import docx
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
doc_path = os.path.join(BASE_DIR, 'KICH_BAN_THUYET_TRINH_MOI_NHAT.docx')
if os.path.exists(doc_path):
    doc = docx.Document(doc_path)
    print(f"Loaded {doc_path} with {len(doc.tables)} tables.")
    for idx, table in enumerate(doc.tables):
        first_text = table.rows[0].cells[0].text[:80] if len(table.rows) > 0 else ""
        print(f"Table {idx}: {first_text}")
        if 'hoàn' in table.rows[0].cells[0].text.lower() or 'fd' in table.rows[0].cells[0].text.lower() or 'quảng tín' in table.rows[0].cells[0].text.lower():
            print(f"Found %FD in Table {idx}!")
