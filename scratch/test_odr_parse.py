import openpyxl
import json
import sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W39_2026.xlsx', data_only=True)
ws = wb['05_ODR']
weeks = ['W36', 'W37', 'W38', 'W39']

def extract_rows(r_start, r_end):
    items = []
    for r in range(r_start, r_end + 1):
        name = ws.cell(r, 1).value
        if name and str(name).strip() not in ['None', '', 'AM', 'Tỉnh', 'THEO AM - TTS', 'THEO TỈNH - Full hàng', 'THEO TỈNH - TTS']:
            w_vals = [ws.cell(r, c).value or 0 for c in range(3, 7)]
            diff = ws.cell(r, 7).value or 0
            vol = ws.cell(r, 2).value or w_vals[-1]
            item = {
                'am': str(name).strip(),
                'tinh': str(name).strip(),
                'vol': vol,
                'diff': diff
            }
            for idx, w in enumerate(weeks):
                item[w.lower()] = w_vals[idx] if idx < len(w_vals) else 0
            items.append(item)
    return items

am_tts = extract_rows(34, 51)
print(f"=== THEO AM - TTS (Found {len(am_tts)} AMs) ===")
for row in am_tts:
    w36_pct = row['w36'] * 100
    w37_pct = row['w37'] * 100
    w38_pct = row['w38'] * 100
    w39_pct = row['w39'] * 100
    diff_pct = row['diff'] * 100
    print(f"{row['am']:<22} | SL={row['vol']:>5} | W36={w36_pct:>5.1f}% | W37={w37_pct:>5.1f}% | W38={w38_pct:>5.1f}% | W39={w39_pct:>5.1f}% | Diff={diff_pct:>+5.1f}%")

print("\n=== THEO TINH - Full hang ===")
tinh_full = extract_rows(56, 60)
for row in tinh_full:
    w36_pct = row['w36'] * 100
    w37_pct = row['w37'] * 100
    w38_pct = row['w38'] * 100
    w39_pct = row['w39'] * 100
    diff_pct = row['diff'] * 100
    print(f"{row['tinh']:<15} | SL={row['vol']:>8.0f} | W36={w36_pct:>5.1f}% | W37={w37_pct:>5.1f}% | W38={w38_pct:>5.1f}% | W39={w39_pct:>5.1f}% | Diff={diff_pct:>+5.1f}%")

print("\n=== THEO TINH - TTS ===")
tinh_tts = extract_rows(65, 69)
for row in tinh_tts:
    w36_pct = row['w36'] * 100
    w37_pct = row['w37'] * 100
    w38_pct = row['w38'] * 100
    w39_pct = row['w39'] * 100
    diff_pct = row['diff'] * 100
    print(f"{row['tinh']:<15} | SL={row['vol']:>8.0f} | W36={w36_pct:>5.1f}% | W37={w37_pct:>5.1f}% | W38={w38_pct:>5.1f}% | W39={w39_pct:>5.1f}% | Diff={diff_pct:>+5.1f}%")
