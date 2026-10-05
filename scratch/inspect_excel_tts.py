import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W40_2026.xlsx', data_only=True)
print("Sheet names:", wb.sheetnames)

for name in wb.sheetnames:
    if 'gtc' in name.lower() or 'tts' in name.lower():
        ws = wb[name]
        print(f"\n--- SHEET: {name} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
        for r in range(1, 15):
            row_vals = [str(ws.cell(r, c).value or '')[:20] for c in range(1, 12)]
            if any(row_vals):
                print(f"Row {r:2d}: {row_vals}")
