import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W40_2026.xlsx', data_only=True)

def dump_sheet(sheet_name, max_r=30, max_c=10):
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name} not found")
        return
    ws = wb[sheet_name]
    print(f"\n==================== SHEET: {sheet_name} ====================")
    for r in range(1, max_r + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, max_c + 1)]
        if any(row_vals):
            vals_str = [f"{v}" if v is not None else "" for v in row_vals]
            print(f"R{r:02d}: " + " | ".join(vals_str[:7]))

# 1. Sheet 07_Gan
dump_sheet('07_Gan', 40, 8)

# 2. Sheet 05_ODR
dump_sheet('05_ODR', 35, 8)

# 3. Sheet 08_OPR TTS
dump_sheet('08_OPR TTS', 30, 8)

# 4. Sheet 09_Rot LC
dump_sheet('09_Rot LC', 25, 8)

# 5. Sheet 12_FD
dump_sheet('12_FD', 30, 8)
