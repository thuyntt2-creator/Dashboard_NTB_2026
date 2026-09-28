import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('KICH_BAN_CANH_BAO_BUU_CUC_W38_VS_W39.docx')
print("=== PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P{i}: {p.text.strip()}")

print("\n=== TABLES ===")
for t_idx, t in enumerate(doc.tables):
    print(f"\n--- TABLE {t_idx} ({len(t.rows)} rows, {len(t.columns)} cols) ---")
    for r_idx, r in enumerate(t.rows):
        row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
        print(f"R{r_idx}: {' | '.join(row_txt)}")
