import docx, sys, shutil
sys.stdout.reconfigure(encoding='utf-8')

src = 'KICH_BAN_THUYET_TRINH_MOI_NHAT.docx'
doc = docx.Document(src)

# TABLE 2: MỤC III - TỶ LỆ GIAO THÀNH CÔNG (%GTC TỔNG) W39 CHUẨN XÁC 100%
cell2 = doc.tables[2].rows[0].cells[0]
for p in cell2.paragraphs[1:]:
    p._p.getparent().remove(p._p)

cell2.paragraphs[0].text = "🗣️ ĐÁNH GIÁ TỶ LỆ GIAO THÀNH CÔNG (%GTC) VÀ ĐỘ LỆCH THEO TỈNH & AM (W39 vs W38):"

sec3 = [
    '📍 1. BẢNG CHỈ SỐ GTC 5 TỈNH THÀNH (W39 vs W38):',
    '• Top 1 - Ninh Thuận: Full W39 đạt 67,96% (+2,28%p) | TTS W39 đạt 68,09% (+5,00%p) ➔ Quán quân toàn vùng cả Full và TikTok Shop!',
    '• Top 2 - Bình Thuận: Full W39 đạt 67,49% (-1,71%p) | TTS W39 đạt 67,94% (-0,15%p) ➔ Vững vàng trên 67%, giữ chuẩn SLA xanh.',
    '• Top 3 - Khánh Hòa: Full W39 đạt 58,56% (+3,57%p) | TTS W39 đạt 59,60% (+5,78%p) ➔ Bước nhảy vọt ấn tượng nhất vùng, TTS tăng gần 6%p.',
    '• Top 4 - Đắk Nông: Full W39 đạt 48,81% (+1,97%p) | TTS W39 đạt 49,56% (+5,07%p) ➔ Cải thiện tốt nhưng vẫn chưa vượt ngưỡng 50%.',
    '• Top 5 - Lâm Đồng: Full W39 đạt 47,38% (-0,60%p) | TTS W39 đạt 48,73% (+1,79%p) ➔ Báo động đỏ: Thấp nhất toàn vùng, cả Full và TTS đều dưới 50%.',
    '➔ TOÀN VÙNG W39: %GTC Full hàng đạt 56,60% (+0,92%p so với W38: 55,68%) | TikTok Shop đạt 57,45% (+3,55%p so với W38: 53,89% ➔ Lần đầu tiên vượt Full hàng).',
    '',
    '🎙️ LỜI THOẠI THUYẾT TRÌNH GTC TỔNG (KHI BẬT TAB 3: GTC TỔNG):',
    '\"Dạ tiếp theo em xin phép bấm chuyển qua Tab 3 là phần Tỷ lệ Giao thành công (%GTC Tổng):',
    '',
    'Nhìn vào màn hình Tab GTC tuần này, chỉ số toàn vùng có khởi sắc nhẹ: GTC Full đạt 56.6%, tăng gần 1% so với tuần trước. Điểm sáng nhất là GTC TikTok Shop tăng vọt lên 57.5%, tăng tới 3.55% và lần đầu tiên vượt qua luôn cả tỷ lệ giao của hàng Full.',
    '',
    'Đầu tiên, nhìn vào bảng xếp hạng 5 tỉnh:',
    '• Tuyên dương Ninh Thuận tuần này quá xuất sắc: Anh Nghĩa với anh Long kéo Ninh Thuận lên quán quân toàn mạng cả hàng Full (68.0%) lẫn TikTok Shop (68.1%).',
    '• Bình Thuận bám sát nút ở vị trí thứ hai với 67.5% Full và 67.9% TTS, tiếp tục giữ vững chuẩn SLA xanh.',
    '• Khánh Hòa là tỉnh có bước nhảy vọt mạnh nhất tuần này: GTC Full tăng 3.6% và TTS tăng gần 6%, kéo tỷ lệ giao của tỉnh lên sát mốc 60%.',
    '• Tuy nhiên, báo động đỏ nằm ở Lâm Đồng và Đắk Nông: Cả 2 tỉnh này tỷ lệ giao đều lẹt đẹt dưới 50% (Đắk Nông 48.8%, Lâm Đồng 47.4%), trực tiếp kéo tụt chỉ số trung bình của cả vùng mình xuống.',
    '',
    'Nhìn xuống biểu đồ phân hóa chi tiết của 18 AM bên dưới:',
    '',
    'Nhóm dẫn đầu giữ vững phong độ trên 69%:',
    '• Đứng đầu toàn mạng vẫn là anh Nguyễn Ngọc Khánh (75.0% Full, TTS lên tới 76.7%) — anh Khánh tiếp tục là quán quân chất lượng của vùng.',
    '• Tiếp theo là anh Nguyễn Đỗ Minh Nghĩa đạt gần 71% và anh Nguyễn Duy Long đạt hơn 69%, cả 2 anh đều có mức tăng trưởng GTC rất tốt.',
    '',
    'Nhưng mà, nhìn xuống đáy bảng xếp hạng, em xin phép cảnh báo khẩn cấp 3 AM đang rơi vào vùng nguy hiểm:',
    '• Báo động nặng nhất tuần này là anh Lê Văn Trường ở Lâm Đồng: Tỷ lệ GTC Full của anh Trường rơi tự do từ 41.2% tuần trước xuống chỉ còn 35.5% tuần này, tức là giảm sâu hơn 5.6%. Đây là mức rơi mạnh nhất toàn mạng, chủ yếu do bưu cục Lâm Viên 2 bị tê liệt ca chiều, kéo tụt cả cụm Đà Lạt.',
    '• Người thứ hai là anh Lê Minh Lợi: GTC Full chỉ đạt 30.6%, đặc biệt GTC TikTok Shop rớt thảm hại xuống 21.3% — thấp nhất toàn mạng lưới.',
    '• Người thứ ba là anh Trương Quang Linh: Dù tuần này có nỗ lực tăng được 9% nhưng GTC Full vẫn đang đội sổ toàn vùng ở mức 25.8%.',
    '',
    '🎯 Về nhiệm vụ tác chiến cho phần GTC tuần W40:',
    '• Em đề nghị anh Trường, anh Lợi và anh Linh bắt buộc phải kiểm soát lại tỷ lệ xuất tuyến ca chiều và giải quyết dứt điểm các bưu cục đang bị nghẽn đơn tồn, quyết tâm tuần tới phải kéo GTC vượt qua mốc 45%.',
    '• Mục tiêu toàn vùng tuần W40: Đưa GTC Full vượt mốc 58% và GTC TikTok Shop chạm ngưỡng 60% Sếp giao.',
    '',
    'Dạ, bây giờ em xin chuyển qua Tab tiếp theo là phần GTC Ca 1 sáng TikTok Shop ạ.\"'
]

for line in sec3:
    p = cell2.add_paragraph()
    p.text = line

doc.save(src)
for c in ['KICH_BAN_THUYET_TRINH_W39_INSIGHT_CHUYEN_SAU.docx', 'KICH_BAN_THUYET_TRINH_W39_NAM_TRUNG_BO.docx']:
    shutil.copyfile(src, c)

print('Updated Table 2 (GTC TONG) with exact W39 data successfully!')
