import docx, sys, shutil
sys.stdout.reconfigure(encoding='utf-8')

src = 'KICH_BAN_THUYET_TRINH_MOI_NHAT.docx'
doc = docx.Document(src)

# TABLE 0: TỔNG QUAN GÃY GỌN
cell0 = doc.tables[0].rows[0].cells[0]
for p in cell0.paragraphs[1:]:
    p._p.getparent().remove(p._p)

sec1 = [
    '📍 1. BẢNG CHỈ SỐ NHANH 5 TỈNH (W39 vs W38):',
    '• Khánh Hòa: 92.7k đơn (-1,0%) | GTC 58,6% (+3,6%p) | ODR 91,8% | TTS: 18.0k (+449 đơn) ➔ Top 1 toàn vùng.',
    '• Lâm Đồng: 88.3k đơn (-7,8%) | GTC 47,4% | ODR 84,7% | TTS: 19.3k (-304 đơn, top 1 TTS vùng).',
    '• Bình Thuận: 80.8k đơn (-5,5%) | GTC 67,5% | ODR 96,3% | TTS: 18.1k (+2.273 đơn, +14,3%!).',
    '• Đắk Nông: 33.9k đơn (-2,8%) | GTC 48,8% | ODR 88,5% | TTS: 9.0k (+702 đơn, +8,4%).',
    '• Ninh Thuận: 33.2k đơn (-2,0%) | GTC 68,0% (Quán quân) | ODR 96,4% (Quán quân) | TTS: 8.3k (+935 đơn, +12,7%).',
    '➔ TOÀN VÙNG: 328.9k đơn Full (-4,3%) | TTS lập đỉnh 72.8k đơn (+5,9%, chiếm 22,1%) | GTC Full 56,6% | GTC TTS 57,5% (+3,6%p) | Rớt LC giảm sâu còn 1,52% (-1,8%p) | Truy thu: 306,9 Tr ₫ / 4.013 đơn.',
    '',
    '🎙️ LỜI THOẠI THUYẾT TRÌNH TỔNG QUAN (1 PHÚT GÃY GỌN):',
    '"Kính thưa Ban Giám Đốc và các anh chị AM, tổng quan tuần W39 tóm gọn trong 3 ý chính:',
    '1. Sản lượng: Đạt 328.9k đơn (-4,3%), giảm ở phân khúc B2C nhưng TikTok Shop bùng nổ lập đỉnh 72.8k đơn (+5,9%), chiếm 22,1% cơ cấu toàn vùng.',
    '2. Điểm sáng chất lượng: GTC TikTok Shop tăng vọt lên 57,5% (+3,6%p, vượt GTC Full); Tỷ lệ rớt luân chuyển KTC giảm hơn một nửa, chỉ còn 1,52%.',
    '3. Hai điểm nóng cần xử lý: Bưu cục cảnh báo tăng lên 15 điểm (kẹt tại Lâm Đồng, Đắk Nông); và phát sinh 306,9 triệu truy thu do đơn tồn đọng giao hàng và luân chuyển trả."',
    '',
    '🎯 TRỌNG TÂM W40: Giải tỏa 15 bưu cục nghẽn, siết tồn đọng để chặn truy thu, giữ nhịp rớt KTC dưới 2%.'
]
for line in sec1:
    p = cell0.add_paragraph()
    p.text = line

# TABLE 1: SẢN LƯỢNG GÃY GỌN
cell1 = doc.tables[1].rows[0].cells[0]
for p in cell1.paragraphs[1:]:
    p._p.getparent().remove(p._p)

sec2 = [
    '📍 1. SỐ LIỆU SẢN LƯỢNG 5 TỈNH:',
    '• Khánh Hòa (92.7k, -1,0%) | Lâm Đồng (88.3k, -7,8%) | Bình Thuận (80.8k, -5,5%) | Đắk Nông (33.9k, -2,8%) | Ninh Thuận (33.2k, -2,0%).',
    '• Toàn vùng: Full 328.925 đơn (-4,3%) | TikTok Shop 72.781 đơn (+5,9%, tỷ trọng 22,1%).',
    '',
    '🎙️ LỜI THOẠI THUYẾT TRÌNH SẢN LƯỢNG (TAB II — 3 CHARTS):',
    '"Kính thưa Ban Giám Đốc, nhìn vào 3 biểu đồ sản lượng, xin báo cáo 4 điểm trọng tâm:',
    '',
    '1. Dòng hàng dịch chuyển sang TikTok Shop:',
    'Đơn truyền thống giảm nhưng TTS tăng trưởng ở 4/5 tỉnh (+4.055 đơn, +5,9%). Bình Thuận (+14%), Ninh Thuận (+13%) và Đắk Nông (+8%) là 3 động lực kéo tải.',
    '',
    '2. Hai AM tăng trưởng & gánh tải tốt nhất:',
    '• AM Thái Thị Thanh Thư (Khánh Hòa): Quán quân tăng trưởng (+2.227 đơn, đạt 34.8k đơn) — Là AM DUY NHẤT trong 18 AM tăng trưởng dương Full hàng, đưa Thư lên Top 2 sản lượng vùng.',
    '• AM Nguyễn Duy Long (Bình Thuận): Gánh tải số 1 (42.5k đơn Full, chiếm 13% vùng) và là AM đầu tiên vượt mốc 10k đơn TTS/tuần (10.6k đơn, +15%).',
    '',
    '3. Hai AM sụt giảm sâu cần lưu tâm:',
    '• AM Nguyễn Thanh Long (Cam Ranh): Sụt giảm TTS nặng nhất mạng (-1.186 đơn TTS, -40,5%, rơi từ 2.9k còn 1.7k đơn). Đây là điểm bất thường lớn khi toàn vùng TTS đang tăng.',
    '• AM Lê Văn Trường (Lâm Đồng): Sụt Full sâu nhất vùng (-2.132 đơn, còn 26.8k đơn) — Đi liền với điểm nghẽn bưu cục Lâm Viên 2 (GTC 19,5%) và 50,6 Tr truy thu từ 722 đơn tồn đọng.',
    '• Lưu ý: AM Trương Quang Linh (34,8%) và Huỳnh Thúc Duân (29,0%) có tỷ trọng đơn sàn rất cao tại vùng núi, cần xuất tuyến sớm ca sáng để bảo vệ GTC.',
    '',
    '4. Nhiệm vụ tác chiến W40:',
    '• AM Long Cam Ranh: Rà soát đầu mối gửi hàng, kéo sản lượng TTS về lại mốc 2.500 đơn/tuần.',
    '• AM Trường: Thông luồng 722 đơn tồn Đà Lạt, khôi phục sản lượng Full về 28.000 đơn.',
    '• 18 AM: Tiếp tục ưu tiên tiến độ lấy và xử lý luồng đơn TikTok Shop trong ngày."',
    '',
    '🎯 GIAO VIỆC: Hoàn thành rà soát nguyên nhân Cam Ranh & giải tỏa tồn Lâm Viên trong 48h tới.'
]
for line in sec2:
    p = cell1.add_paragraph()
    p.text = line

doc.save(src)
for c in ['KICH_BAN_THUYET_TRINH_W39_INSIGHT_CHUYEN_SAU.docx', 'KICH_BAN_THUYET_TRINH_W39_NAM_TRUNG_BO.docx']:
    shutil.copyfile(src, c)

print('Updated docx successfully!')
