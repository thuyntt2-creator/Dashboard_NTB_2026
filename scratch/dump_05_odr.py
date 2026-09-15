# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'c:\Users\lap4all\Desktop\New folder\BaoCao_Tuan_NTB_W37_2026.xlsx', data_only=True)
ws = wb['05_ODR']
print(f"Max rows: {ws.max_row}, Max cols: {ws.max_column}")
for r in range(1, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if any(vals):
        # format percentages
        formatted = []
        for v in vals:
            if isinstance(v, float):
                formatted.append(f"{v:.4f}")
            else:
                formatted.append(str(v) if v is not None else '')
        print(f"Row {r:2d}: " + " | ".join(formatted[:8]))
