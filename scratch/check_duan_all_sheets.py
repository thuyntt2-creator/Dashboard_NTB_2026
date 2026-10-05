import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W40_2026.xlsx', data_only=True)

for s in ['02_GTC tong', '03_GTC Ca1+Ton', '03b_GTC Ca1 thuan', '04_GTC Ca2', '05_ODR', '06_LTC', '07_Gan', '09_Rot LC', '12_FD']:
    if s in wb.sheetnames:
        ws = wb[s]
        # find where THEO AM - TTS starts and ends
        print(f"\n--- {s} ---")
        for r in range(10, 56):
            c1 = str(ws.cell(r, 1).value or '').strip()
            if 'duân' in c1.lower() or 'long' in c1.lower():
                print(f"  Row {r:2d}: {c1}")
            if 'theo' in c1.lower():
                print(f"  Header Row {r:2d}: {c1}")
