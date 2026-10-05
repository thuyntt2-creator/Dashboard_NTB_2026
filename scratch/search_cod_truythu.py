import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W40_2026.xlsx', data_only=True)
print("=== ALL SHEETS ===")
for name in wb.sheetnames:
    print("-", name)

def search_text_in_sheets(keywords):
    print("\n=== SEARCHING KEYWORDS:", keywords)
    for name in wb.sheetnames:
        ws = wb[name]
        found = []
        for r in range(1, min(ws.max_row+1, 50)):
            for c in range(1, min(ws.max_column+1, 20)):
                v = str(ws.cell(r, c).value or "")
                if any(k.lower() in v.lower() for k in keywords):
                    found.append(f"R{r}C{c}: {v[:40]}")
        if found:
            print(f"\nSheet [{name}]:")
            for item in found[:10]:
                print("  ", item)

search_text_in_sheets(["cod", "truy thu", "nợ", "công nợ", "tiền"])
