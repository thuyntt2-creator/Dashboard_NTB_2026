import docx, sys, shutil
sys.stdout.reconfigure(encoding='utf-8')

src = 'KICH_BAN_THUYET_TRINH_MOI_NHAT.docx'
doc = docx.Document(src)

# TABLE 3: MỤC IV - GTC CA 1 SÁNG TIKTOK SHOP W39 CHUẨN XÁC 100%
cell3 = doc.tables[3].rows[0].cells[0]
for p in cell3.paragraphs[1:]:
    p._p.getparent().remove(p._p)

cell3.paragraphs[0].text = "🗣️ MỔ XẺ HIỆU SUẤT GIAO CA 1 SÁNG TIKTOK SHOP (TARGET SLA ≥ 76.0%) W39 vs W38:"

sec4 = [
    '📍 1. BẢNG CHỈ SỐ GTC CA 1 SÁNG TIKTOK SHOP 5 TỈNH (W39 vs W38):',
    '• Top 1 - Ninh Thuận: Ca 1 TTS đạt 84,75% (W38: 81,71%, +3,04%p) ➔ Quán quân toàn vùng, vượt xa chuẩn SLA ≥76%!',
    '• Top 2 - Bình Thuận: Ca 1 TTS đạt 84,64% (W38: 85,54%, -0,90%p) ➔ Duy trì phong độ xuất sắc trên 84%, vượt xa chuẩn.',
    '• Top 3 - Khánh Hòa: Ca 1 TTS đạt 75,53% (W38: 67,58%, +7,95%p) ➔ Bước nhảy vọt mạnh nhất toàn mạng (+8,0%p WoW), áp sát mốc 76%.',
    '• Top 4 - Đắk Nông: Ca 1 TTS đạt 67,81% (W38: 65,11%, +2,71%p) ➔ Chưa đạt target 76%, còn cách đích hơn 8%p.',
    '• Top 5 - Lâm Đồng: Ca 1 TTS đạt 66,85% (W38: 63,96%, +2,89%p) ➔ Báo động đỏ: Thấp nhất vùng, cách chuẩn SLA gần 10%p.',
    '➔ TOÀN VÙNG W39: Ca 1 thuần sáng đạt 75,10% (+3,84%p WoW so với W38: 71,26%, tiệm cận sát nút Target SLA 76.0%) ➔ Sang Ca 2 rơi tự do về 46,66% (chênh lệch gãy cánh gần 28,5%p) ➔ Kéo tỷ lệ cả ngày TTS xuống 57,45%.',
    '',
    '🎙️ LỜI THOẠI THUYẾT TRÌNH CA 1 SÁNG TIKTOK SHOP (TAB 4):',
    '\"Dạ tiếp theo em xin phép chuyển qua Tab 4 là phần Hiệu suất Giao Ca 1 sáng TikTok Shop (mục tiêu SLA của công ty là phải đạt từ 76% trở lên):',
    '',
    'Nhìn vào bức tranh toàn vùng tuần W39, Ca 1 thuần sáng của TikTok Shop đã có bước tiến bộ rất rõ rệt: Kéo từ 71.3% tuần trước lên 75.1% tuần này, tăng gần 4% và đang áp sát sát nút mốc chuẩn 76%. Tuy nhiên, điểm nhức nhối nhất của vùng mình vẫn là bài toán gãy cánh ca chiều: Sáng giao được hơn 75% nhưng chiều (Ca 2) lại rớt tự do về 46.7%, tức là chênh lệch tới gần 28.5%! Chính ca chiều đã kéo cả ngày của TikTok Shop xuống còn 57.5%.',
    '',
    'Đầu tiên, nhìn vào bảng xếp hạng 5 tỉnh:',
    '• Tuyên dương Ninh Thuận (84.8%) và Bình Thuận (84.6%): Cả 2 tỉnh duyên hải này đều vượt xa chuẩn SLA 76%, giữ tỷ lệ giao ca sáng trên 84%.',
    '• Khánh Hòa tuần này có bước nhảy vọt ngoạn mục nhất: Ca 1 TTS tăng vọt gần 8%p (từ 67.6% lên 75.5%), chỉ còn cách vạch đích chuẩn đúng 0.5%p.',
    '• Ngược lại, Đắk Nông (67.8%) và Lâm Đồng (66.9%) vẫn chưa qua được mốc 70%, đang cách chuẩn SLA của công ty từ 8 đến gần 10%p.',
    '',
    'Bóc tách chi tiết theo 18 anh chị AM:',
    '',
    'Top AM giao Ca 1 TikTok Shop xuất sắc nhất (vượt xa chuẩn 76%):',
    '• Quán quân tuần này là anh Nguyễn Ngọc Khánh ở Khánh Hòa đạt tới 89.1%!',
    '• Tiếp theo là anh Nguyễn Đỗ Minh Nghĩa đạt 86.9% (+3.9%p), chị Cao Thị Thanh Thủy đạt 86.0%, anh Nguyễn Duy Long đạt 85.5% (trên khối lượng lớn gần 7.800 đơn) và chị Nguyễn Thị Tuyết Thơ đạt 81.6%. Đây là nhóm 5 AM giữ vững kỷ luật giao ca sáng rất đáng biểu dương.',
    '',
    'Nhưng mà, nhìn xuống đáy bảng xếp hạng, có 3 AM đang kéo lùi chỉ số Ca 1 của cả vùng:',
    '• Chỗ anh Lê Văn Trường ở Lâm Đồng: Ca 1 TTS chỉ đạt 51.2% (giảm -1.7%p WoW trên khối lượng lớn hơn 4.000 đơn). Ca sáng mà chỉ giao được một nửa thì chắc chắn ca chiều sẽ bị vỡ trận.',
    '• Hai AM ở Đắk Nông là anh Trương Quang Linh (39.2%) và anh Lê Minh Lợi (38.7%): Cả 2 anh đều chưa chạm nổi mốc 40% ở ca sáng.',
    '',
    '🎯 Về nhiệm vụ tác chiến cho Ca 1 TTS tuần W40:',
    '• Em đề nghị anh Trường, anh Linh và anh Lợi bắt buộc bưu tá phải xuất tuyến lượt 1 trước 08h30 sáng, ưu tiên giao 100% đơn TikTok Shop trước 11h00 trưa để kéo Ca 1 tuần tới vượt qua mốc 60%.',
    '• Toàn vùng quyết tâm tuần W40 đưa Ca 1 thuần sáng vượt mốc 76.0% SLA cam kết.',
    '',
    'Dạ, bây giờ em xin chuyển qua Tab tiếp theo là phần Kiểm soát Tỷ lệ gán đơn giao ạ.\"'
]

for line in sec4:
    p = cell3.add_paragraph()
    p.text = line

doc.save(src)
for c in ['KICH_BAN_THUYET_TRINH_W39_INSIGHT_CHUYEN_SAU.docx', 'KICH_BAN_THUYET_TRINH_W39_NAM_TRUNG_BO.docx']:
    shutil.copyfile(src, c)

print('Updated Table 3 (GTC CA 1 TTS) with exact W39 data successfully!')
