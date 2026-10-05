import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

file_path = r"C:\Users\lap4all\Downloads\KICH_BAN_THUYET_TRINH_W39_NAM_TRUNG_BO (16).docx"
doc = docx.Document(file_path)

def print_table_full(idx):
    if idx >= len(doc.tables):
        return
    tbl = doc.tables[idx]
    print(f"\n==================== TABLE {idx+1} FULL TEXT ====================")
    for r in tbl.rows:
        for c in r.cells:
            print(c.text.strip())

print_table_full(4) # Table 5 (Gán)
print_table_full(5) # Table 6 (ODR)
print_table_full(8) # Table 9 (FD)
print_table_full(10) # Table 11 (COD)
print_table_full(11) # Table 12 (Truy thu)
