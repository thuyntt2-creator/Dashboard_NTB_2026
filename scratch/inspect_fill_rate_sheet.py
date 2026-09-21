import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook("BaoCao_Tuan_NTB_W38_2026.xlsx", data_only=True)
sheet_name = "14_NTB_Fill_Rate_Report"
if sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"Sheet {sheet_name}: {ws.max_row} rows, {ws.max_column} cols")
    for r in range(1, min(35, ws.max_row + 1)):
        row_vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
        if any(v is not None for v in row_vals):
            print(f"Row {r:2d}: {row_vals}")
else:
    print(f"Sheet {sheet_name} not found in workbook. Available sheets:", wb.sheetnames)
