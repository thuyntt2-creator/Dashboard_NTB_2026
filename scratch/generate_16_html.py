import docx
import re
import os
import sys

def docx_to_html(docx_path, html_path):
    doc = docx.Document(docx_path)
    
    html = """<!DOCTYPE html>
<html lang="vi">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Kịch Bản Thuyết Trình W38 — 16 Phần Toàn Diện | GHN Nam Trung Bộ</title>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap" rel="stylesheet">
    <style>
        :root {
            --primary: #0F4C81;
            --primary-dark: #0A3258;
            --accent: #EA580C;
            --accent-light: #FFF7ED;
            --accent-border: #F97316;
            --text-main: #1F2937;
            --text-muted: #6B7280;
            --bg-page: #F8FAFC;
            --card-bg: #FFFFFF;
        }
        * { box-sizing: border-box; margin: 0; padding: 0; }
        body {
            font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
            background: var(--bg-page);
            color: var(--text-main);
            line-height: 1.6;
            padding: 24px 16px;
        }
        .container {
            max-width: 960px;
            margin: 0 auto;
        }
        .header-card {
            background: white;
            border-radius: 16px;
            padding: 32px 24px;
            text-align: center;
            box-shadow: 0 4px 20px -2px rgba(0,0,0,0.06);
            margin-bottom: 24px;
            border-top: 5px solid var(--primary);
        }
        .badge-region {
            display: inline-block;
            background: #E0F2FE;
            color: var(--primary);
            font-weight: 700;
            font-size: 12px;
            padding: 4px 14px;
            border-radius: 9999px;
            text-transform: uppercase;
            letter-spacing: 0.5px;
            margin-bottom: 8px;
        }
        h1 {
            color: var(--primary);
            font-size: 26px;
            font-weight: 800;
            margin-bottom: 6px;
        }
        .date-range {
            color: var(--text-muted);
            font-size: 14px;
            font-style: italic;
            margin-bottom: 12px;
        }
        .subtitle {
            color: var(--accent);
            font-size: 14px;
            font-weight: 600;
            max-width: 720px;
            margin: 0 auto 20px auto;
        }
        .action-bar {
            display: flex;
            gap: 12px;
            justify-content: center;
            flex-wrap: wrap;
        }
        .btn-download {
            background: linear-gradient(135deg, var(--primary) 0%, #1D4ED8 100%);
            color: white;
            padding: 10px 20px;
            border-radius: 10px;
            text-decoration: none;
            font-weight: 600;
            font-size: 14px;
            display: inline-flex;
            align-items: center;
            gap: 8px;
            box-shadow: 0 4px 12px rgba(15, 76, 129, 0.25);
            transition: all 0.2s;
        }
        .btn-download:hover {
            transform: translateY(-1px);
            box-shadow: 0 6px 16px rgba(15, 76, 129, 0.35);
        }
        .btn-back {
            background: white;
            color: var(--text-main);
            border: 1px solid #D1D5DB;
            padding: 10px 18px;
            border-radius: 10px;
            text-decoration: none;
            font-weight: 600;
            font-size: 14px;
            display: inline-flex;
            align-items: center;
            gap: 6px;
        }
        .btn-back:hover { background: #F3F4F6; }
        
        .toc-card {
            background: white;
            border-radius: 12px;
            padding: 16px 20px;
            margin-bottom: 24px;
            box-shadow: 0 2px 8px rgba(0,0,0,0.04);
        }
        .toc-title {
            font-weight: 700;
            font-size: 14px;
            color: var(--primary);
            margin-bottom: 12px;
            display: flex;
            align-items: center;
            gap: 6px;
        }
        .toc-grid {
            display: grid;
            grid-template-columns: repeat(auto-fill, minmax(280px, 1fr));
            gap: 8px;
        }
        .toc-link {
            font-size: 12.5px;
            color: #374151;
            text-decoration: none;
            padding: 6px 10px;
            border-radius: 6px;
            background: #F9FAFB;
            display: block;
            border-left: 3px solid transparent;
            transition: all 0.15s;
        }
        .toc-link:hover {
            background: #EFF6FF;
            color: var(--primary);
            border-left-color: var(--primary);
        }
        
        .section-card {
            background: white;
            border-radius: 14px;
            padding: 24px 20px;
            margin-bottom: 20px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.04);
            scroll-margin-top: 20px;
        }
        .section-header {
            color: var(--primary);
            font-size: 17px;
            font-weight: 800;
            margin-bottom: 14px;
            padding-bottom: 8px;
            border-bottom: 1.5px solid #F3F4F6;
            display: flex;
            align-items: center;
            gap: 8px;
        }
        .callout-box {
            background: var(--accent-light);
            border-left: 4px solid var(--accent-border);
            border-radius: 0 10px 10px 0;
            padding: 18px 20px;
        }
        .speech-title {
            color: var(--accent);
            font-weight: 700;
            font-size: 15px;
            margin-bottom: 12px;
            display: flex;
            align-items: center;
            gap: 6px;
        }
        .speech-text {
            font-size: 14.5px;
            color: #1F2937;
            margin-bottom: 12px;
            text-align: justify;
        }
        .speech-text:last-of-type { margin-bottom: 14px; }
        
        .insight-tag {
            font-weight: 700;
            font-size: 13.5px;
            color: var(--primary);
            margin-top: 14px;
            margin-bottom: 6px;
            display: flex;
            align-items: center;
            gap: 6px;
        }
        .warning-tag {
            font-weight: 700;
            font-size: 13.5px;
            color: #DC2626;
            margin-top: 14px;
            margin-bottom: 6px;
            display: flex;
            align-items: center;
            gap: 6px;
        }
        .action-tag {
            font-weight: 700;
            font-size: 13.5px;
            color: #15803D;
            margin-top: 14px;
            margin-bottom: 6px;
            display: flex;
            align-items: center;
            gap: 6px;
        }
        .bullet-item {
            font-size: 13.5px;
            margin-left: 18px;
            margin-bottom: 4px;
            color: #374151;
        }
        .bullet-item.warning { color: #991B1B; }
        .bullet-item.action { color: #14532D; }
        
        footer {
            text-align: center;
            padding: 24px;
            color: var(--text-muted);
            font-size: 13px;
        }
    </style>
</head>
<body>
    <div class="container">
        <div class="header-card">
            <div class="badge-region">GHN Express — Vùng Nam Trung Bộ</div>
            <h1>BÁO CÁO VẬN HÀNH & KINH DOANH TUẦN W38</h1>
            <div class="date-range">(Chu kỳ dữ liệu: 14/09/2026 – 20/09/2026)</div>
            <div class="subtitle">Kịch bản thuyết trình toàn diện 16 phần chuẩn hóa số liệu thực tế, mổ xẻ Insight bản chất vận hành & Nghịch lý điều hành</div>
            <div class="action-bar">
                <a href="/download-doc" class="btn-download">
                    📥 Tải File Word (.docx) 16 Phần Chuẩn
                </a>
                <a href="/hop" class="btn-back">
                    📊 Quay lại Dashboard Điều Hành
                </a>
            </div>
        </div>

        <div class="toc-card">
            <div class="toc-title">📑 Mục Lục 16 Chuyên Đề Báo Cáo Tuần W38:</div>
            <div class="toc-grid">
"""

    toc_items = []
    sections_html = []
    
    for i, t in enumerate(doc.tables):
        sec_idx = 4 + i * 2
        sec_title = doc.paragraphs[sec_idx].text if sec_idx < len(doc.paragraphs) else f"Phần {i+1}"
        cell = t.rows[0].cells[0]
        paragraphs = [p.text for p in cell.paragraphs if p.text.strip()]
        
        speech_title = paragraphs[0] if paragraphs else f"Phần {i+1}"
        
        # Clean title for TOC
        clean_toc = re.sub(r'^[^\w\s]+', '', sec_title).strip()
        clean_toc = clean_toc.strip('[]')
        toc_items.append(f'<a href="#sec-{i+1}" class="toc-link"><strong>{i+1}.</strong> {clean_toc}</a>')
        
        # Parse body
        body_paras = []
        insights = []
        warnings = []
        actions = []
        mode = 'body'
        
        for p in paragraphs[1:]:
            if "INSIGHT BẢN CHẤT" in p:
                mode = 'insight'
                continue
            elif "CẢNH BÁO ĐỎ" in p:
                mode = 'warning'
                continue
            elif "QUYẾT SÁCH HÀNH ĐỘNG" in p or "MỆNH LỆNH TÁC CHIẾN" in p:
                mode = 'action'
                continue
                
            if mode == 'body':
                body_paras.append(p)
            elif mode == 'insight':
                insights.append(p.lstrip('•- '))
            elif mode == 'warning':
                warnings.append(p.lstrip('•- '))
            elif mode == 'action':
                actions.append(p.lstrip('•- '))

        sec_card = f"""
        <div class="section-card" id="sec-{i+1}">
            <div class="section-header">{sec_title}</div>
            <div class="callout-box">
                <div class="speech-title">{speech_title}</div>
        """
        for bp in body_paras:
            sec_card += f'<p class="speech-text">{bp}</p>\n'
            
        if insights:
            sec_card += '<div class="insight-tag">🔍 INSIGHT BẢN CHẤT & GỐC RỄ NGUYÊN NHÂN:</div>\n'
            for ins in insights:
                sec_card += f'<div class="bullet-item">• {ins}</div>\n'
                
        if warnings:
            sec_card += '<div class="warning-tag">⚠️ CẢNH BÁO ĐỎ & NGUY CƠ TIỀM ẨN:</div>\n'
            for w in warnings:
                sec_card += f'<div class="bullet-item warning">• {w}</div>\n'
                
        if actions:
            sec_card += '<div class="action-tag">🎯 QUYẾT SÁCH HÀNH ĐỘNG & MỆNH LỆNH TÁC CHIẾN:</div>\n'
            for a in actions:
                sec_card += f'<div class="bullet-item action">• {a}</div>\n'
                
        sec_card += """
            </div>
        </div>
        """
        sections_html.append(sec_card)

    html += "\n".join(toc_items)
    html += """
            </div>
        </div>
    """
    html += "\n".join(sections_html)
    html += """
        <footer>
            GHN Express Vùng Nam Trung Bộ — Hệ thống Báo cáo & Điều hành Tự động W38 © 2026
        </footer>
    </div>
</body>
</html>
"""

    with open(html_path, 'w', encoding='utf-8') as f:
        f.write(html)
    print(f"Generated HTML successfully: {html_path} ({os.path.getsize(html_path)} bytes)")

if __name__ == "__main__":
    docx_to_html("KICH_BAN_THUYET_TRINH_W38_INSIGHT_CHUYEN_SAU.docx", "KICH_BAN_THUYET_TRINH_W38_NAM_TRUNG_BO.html")
