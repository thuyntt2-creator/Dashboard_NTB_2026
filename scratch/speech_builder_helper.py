import os, sys, docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.path.insert(0, os.path.join(os.getcwd(), 'scratch'))
from doc_builder_helpers import format_run, set_cell_borders_and_shading

def add_speech_section(doc, sec_title, speech_heading, paragraphs_text, insights=None, warnings=None, actions=None):
    # Section Heading
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(14)
    p_sec.paragraph_format.space_after = Pt(6)
    p_sec.paragraph_format.keep_with_next = True
    r_sec = p_sec.add_run(sec_title)
    format_run(r_sec, font_size_pt=13, bold=True, color_rgb=(0x0F, 0x4C, 0x81))
    
    # 1x1 Callout Table
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_borders_and_shading(cell, border_color="F97316", fill_color="FFF7ED")
    
    # Speech Title
    p_first = cell.paragraphs[0]
    p_first.paragraph_format.space_before = Pt(2)
    p_first.paragraph_format.space_after = Pt(4)
    p_first.paragraph_format.line_spacing = 1.15
    r_title = p_first.add_run(speech_heading)
    format_run(r_title, font_size_pt=11.5, bold=True, color_rgb=(0xEA, 0x58, 0x0C))
    
    # Body Paragraphs (Spoken text with natural flow)
    for p_text in paragraphs_text:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(p_text)
        format_run(r, font_size_pt=10.5, color_rgb=(0x1F, 0x29, 0x37))
        
    # Insights
    if insights:
        p_in_h = cell.add_paragraph()
        p_in_h.paragraph_format.space_before = Pt(6)
        p_in_h.paragraph_format.space_after = Pt(2)
        r_ih = p_in_h.add_run("🔍 INSIGHT BẢN CHẤT & GỐC RỄ NGUYÊN NHÂN:")
        format_run(r_ih, font_size_pt=10.5, bold=True, color_rgb=(0x0F, 0x4C, 0x81))
        for ins in insights:
            p_in = cell.add_paragraph()
            p_in.paragraph_format.space_before = Pt(1)
            p_in.paragraph_format.space_after = Pt(2)
            p_in.paragraph_format.left_indent = Inches(0.15)
            r_in = p_in.add_run(f"• {ins}")
            format_run(r_in, font_size_pt=10.0, color_rgb=(0x1F, 0x29, 0x37))
            
    # Warnings
    if warnings:
        p_w_h = cell.add_paragraph()
        p_w_h.paragraph_format.space_before = Pt(6)
        p_w_h.paragraph_format.space_after = Pt(2)
        r_wh = p_w_h.add_run("⚠️ CẢNH BÁO ĐỎ & NGUY CƠ TIỀM ẨN:")
        format_run(r_wh, font_size_pt=10.5, bold=True, color_rgb=(0xDC, 0x26, 0x26))
        for w in warnings:
            p_w = cell.add_paragraph()
            p_w.paragraph_format.space_before = Pt(1)
            p_w.paragraph_format.space_after = Pt(2)
            p_w.paragraph_format.left_indent = Inches(0.15)
            r_w = p_w.add_run(f"• {w}")
            format_run(r_w, font_size_pt=10.0, color_rgb=(0x99, 0x1B, 0x1B))
            
    # Actions
    if actions:
        p_a_h = cell.add_paragraph()
        p_a_h.paragraph_format.space_before = Pt(6)
        p_a_h.paragraph_format.space_after = Pt(2)
        r_ah = p_a_h.add_run("🎯 QUYẾT SÁCH HÀNH ĐỘNG & MỆNH LỆNH TÁC CHIẾN TUẦN W39:")
        format_run(r_ah, font_size_pt=10.5, bold=True, color_rgb=(0x04, 0x78, 0x57))
        for a in actions:
            p_a = cell.add_paragraph()
            p_a.paragraph_format.space_before = Pt(1)
            p_a.paragraph_format.space_after = Pt(2)
            p_a.paragraph_format.left_indent = Inches(0.15)
            r_a = p_a.add_run(f"• {a}")
            format_run(r_a, font_size_pt=10.0, color_rgb=(0x06, 0x5F, 0x46))

print("Speech section builder ready.")
