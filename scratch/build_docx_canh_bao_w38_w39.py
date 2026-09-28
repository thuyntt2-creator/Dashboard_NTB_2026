import sys
import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding='utf-8')

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_callout_box(doc, title, lines_text, border_color="EA580C", bg_color="FFF7ED"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.9)
    
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=180, bottom=180, left=240, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(title)
    run_t.bold = True
    run_t.font.name = 'Arial'
    run_t.font.size = Pt(11)
    run_t.font.color.rgb = RGBColor(194, 65, 12) if border_color=="EA580C" else RGBColor(15, 76, 129)
    
    for line in lines_text:
        line_str = line.strip()
        if not line_str:
            continue
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.25
        
        if line_str.startswith(('•', '-', '*')) or line_str.startswith(('1.', '2.', '3.', '4.', '5.', '6.')):
            if ':' in line_str:
                parts = line_str.split(':', 1)
                r_pre = p2.add_run(parts[0] + ":")
                r_pre.bold = True
                r_pre.font.name = 'Arial'
                r_pre.font.size = Pt(9.5)
                r_pre.font.color.rgb = RGBColor(15, 23, 42)
                r_post = p2.add_run(parts[1])
                r_post.font.name = 'Arial'
                r_post.font.size = Pt(9.5)
                r_post.font.color.rgb = RGBColor(51, 65, 85)
            else:
                run_b = p2.add_run(line_str)
                run_b.font.name = 'Arial'
                run_b.font.size = Pt(9.5)
                run_b.font.color.rgb = RGBColor(30, 41, 59)
        elif line_str.startswith('"') or line_str.endswith('"'):
            run_b = p2.add_run(line_str)
            run_b.font.name = 'Arial'
            run_b.font.size = Pt(9.5)
            run_b.italic = True
            run_b.font.color.rgb = RGBColor(15, 23, 42)
        else:
            run_b = p2.add_run(line_str)
            run_b.font.name = 'Arial'
            run_b.font.size = Pt(9.5)
            run_b.font.color.rgb = RGBColor(30, 41, 59)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def style_table(tbl, col_widths, headers, rows_data, header_bg="0F4C81", header_fg="FFFFFF"):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Header
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].width = Inches(col_widths[i])
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 2, 3, 4, 5, 6, 7] else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h)
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    # Rows
    for r_idx, r_data in enumerate(rows_data):
        row = tbl.add_row()
        cells = row.cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cells[c_idx].width = Inches(col_widths[c_idx])
            set_cell_background(cells[c_idx], bg_color)
            set_cell_margins(cells[c_idx], top=80, bottom=80, left=100, right=100)
            p = cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 2, 3, 4, 5, 6, 7] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(val))
            r.font.name = 'Arial'
            r.font.size = Pt(8.5)
            if c_idx == 1:
                r.bold = True
            if 'Thoát' in str(val) or '▲' in str(val):
                r.font.color.rgb = RGBColor(5, 150, 105)
                r.bold = True
            elif 'Mới' in str(val) or '▼' in str(val) or 'Mãn tính' in str(val) or '🚨' in str(val):
                r.font.color.rgb = RGBColor(220, 38, 38)
                if 'Mãn tính' in str(val) or 'Mới' in str(val):
                    r.bold = True

def build_docx():
    doc = docx.Document()
    
    # Page setup
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    # Header block
    p_comp = doc.add_paragraph()
    p_comp.paragraph_format.space_before = Pt(0)
    p_comp.paragraph_format.space_after = Pt(2)
    r_comp = p_comp.add_run("CÔNG TY CỔ PHẦN GIAO HÀNG NHANH — GHN EXPRESS")
    r_comp.bold = True
    r_comp.font.name = 'Arial'
    r_comp.font.size = Pt(9.5)
    r_comp.font.color.rgb = RGBColor(100, 116, 139)
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(2)
    r_t = p_title.add_run("BÁO CÁO VẬN HÀNH & ĐIỀU HÀNH VÙNG NAM TRUNG BỘ")
    r_t.bold = True
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(13)
    r_t.font.color.rgb = RGBColor(15, 76, 129) # GHN Blue
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(10)
    r_sub = p_sub.add_run("KỊCH BẢN THUYẾT TRÌNH: THEO DÕI & XỬ LÝ NHÓM BƯU CỤC CẢNH BÁO BẤT ỔN\n(SO SÁNH ĐỐI CHIẾU TUẦN W38 VS TUẦN W39 — KẾ HOẠCH HÀNH ĐỘNG W40)")
    r_sub.bold = True
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(234, 88, 12) # GHN Orange
    
    # Metadata
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(12)
    r_m = p_meta.add_run("Chu kỳ dữ liệu: W38 (14/09 – 20/09/2026) vs W39 (21/09 – 27/09/2026) | Phạm vi: 5 Tỉnh & 18 AM Nam Trung Bộ")
    r_m.italic = True
    r_m.font.name = 'Arial'
    r_m.font.size = Pt(9)
    r_m.font.color.rgb = RGBColor(100, 116, 139)

    # 1. CALLOUT: FULL PRESENTATION SPEECH
    speech_lines = [
        '"Kính thưa Ban Giám Đốc, các anh chị Giám đốc Khối và toàn thể đội ngũ Quản lý Vận hành (AM) vùng Nam Trung Bộ.',
        'Sau đây, em xin phép báo cáo vào chuyên đề trọng điểm: Theo dõi và xử lý nhóm Bưu cục cảnh báo bất ổn (%GTC < 45% hoặc giảm sâu dưới 70% mốc tốt nhất lịch sử) — đặt trên hệ quy chiếu so sánh đối chiếu giữa Tuần W38 và Tuần W39.',
        'Thưa Ban Giám Đốc, nếu nhìn vào bức tranh chung toàn vùng, %GTC của chúng ta vẫn giữ được nhịp ổn định. Tuy nhiên, khi bóc tách xuống cấp độ bưu cục cơ sở, tuần W39 ghi nhận sự dịch chuyển rủi ro rõ rệt mà Ban Điều Hành cần báo cáo thẳng thắn:',
        '1. Tổng quan số lượng: Tuần W38 toàn vùng có 13 bưu cục cảnh báo, sang tuần W39 tăng lên 15 bưu cục (tăng ròng +2 bưu cục). Trong đó, có 4 bưu cục nỗ lực thoát cảnh báo thành công, nhưng lại phát sinh 6 bưu cục mới rơi vào diện báo động, và 9 bưu cục đang kẹt dai dẳng cả 2 tuần liên tiếp.',
        '2. Tuyên dương 4 bưu cục thoát cảnh báo thành công: Xuân Hương - Đà Lạt (AM Lê Văn Trường), Cam Linh (AM Phan Đình Duy), Nhân Cơ (AM Đỗ Duy Khang), và Tây Nha Trang (AM Hồng Bích Nga). Các bưu cục này đã tối ưu vượt bậc ca phát sáng để kéo %GTC vượt ngưỡng an toàn.',
        '3. Báo động đỏ 6 bưu cục mới lọt vào diện cảnh báo W39: BC D\'Ran (Lâm Đồng - 36.86%), BC Lang Biang 2 (Lâm Đồng - 35.80%), BC Trường Xuân (Đắk Nông - 38.22%), BC Đông Gia Nghĩa (Đắk Nông - 44.89%), BC Bắc Cam Ranh (Khánh Hòa - 44.68%), và BC Phú Quý (Bình Thuận - 50.84%, rơi sâu dưới 70% mốc đỉnh lịch sử 88.55%).',
        '4. Điểm nóng mãn tính 9 bưu cục kẹt cả 2 tuần: Nghiêm trọng nhất là BC Lâm Viên - Đà Lạt 2 từ 46.2% ở W38 rơi tự do xuống 19.55% trong W39 (nặng nhất toàn mạng); BC Đức Trọng 1 (23.78%) và BC Đơn Dương (29.01%, giảm -16.3%p); BC Quảng Tín (26.17%, tích lũy >100 ngày cảnh báo). Cùng 5 BC: Lang Biang 1 (30.6%), Kiến Đức (32.5%), Tân Hà Lâm Hà (36.8%), Tuy Đức (38.5%) và Di Linh (40.6%).',
        '5. Kế hoạch tác chiến tuần W40: Áp dụng cơ chế Điều tiết hàng (Cap Volume) giảm 25% hạn mức chia chọn tại Lâm Viên 2 và Quảng Tín để quét sạch tồn aging; Tái cơ cấu bưu tá ca sáng điều chuyển chi viện 4 bưu tá cứng hỗ trợ phát lượt 1 trước 11h00; Cam kết kéo tối thiểu 5 bưu cục thoát cảnh báo trong tuần W40, đưa danh sách toàn vùng về dưới 10 bưu cục.',
        'Em xin kết thúc phần báo cáo bưu cục cảnh báo. Kính mời Ban Giám Đốc cho ý kiến chỉ đạo ạ!"'
    ]
    create_callout_box(
        doc,
        "🗣️ LỜI THOẠI DÀNH CHO NGƯỜI BÁO CÁO (TRÌNH BÀY TRƯỚC BAN GIÁM ĐỐC & KHỐI VẬN HÀNH)",
        speech_lines,
        border_color="EA580C",
        bg_color="FFF7ED"
    )

    # 2. DETAIL SECTION: PHÂN TÍCH CHUYÊN SÂU 3 NHÓM
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run("🔍 I. PHÂN TÍCH BÓC TÁCH CHI TIẾT 3 NHÓM BƯU CỤC BIẾN ĐỘNG (W38 vs W39)")
    r_h2.bold = True
    r_h2.font.name = 'Arial'
    r_h2.font.size = Pt(11)
    r_h2.font.color.rgb = RGBColor(15, 23, 42)

    # Group 1: 4 BC Thoat
    p_g1 = doc.add_paragraph()
    p_g1.paragraph_format.space_before = Pt(4)
    p_g1.paragraph_format.space_after = Pt(2)
    r_g1_t = p_g1.add_run("1. Tuyên dương 4 Bưu cục đã giải tỏa dứt điểm & thoát cảnh báo thành công (W38 🚨 ➔ W39 ✅):")
    r_g1_t.bold = True
    r_g1_t.font.name = 'Arial'
    r_g1_t.font.size = Pt(10)
    r_g1_t.font.color.rgb = RGBColor(5, 150, 105)

    bullets_g1 = [
        "• BC Xuân Hương - Đà Lạt (Lâm Đồng - AM Lê Văn Trường): Nâng mạnh tỷ lệ phát thành công lượt 1, giải tỏa toàn bộ 100% hàng tồn đọng sau đỉnh tải.",
        "• BC Cam Linh (Khánh Hòa - AM Phan Đình Duy): Ổn định lại nhân sự bưu tá tuyến Cam Ranh, khắc phục triệt để tình trạng trễ giờ xuất tuyến.",
        "• BC Nhân Cơ (Đắk Nông - AM Đỗ Duy Khang): Kiểm soát hoàn hảo luồng hàng nông thôn, đưa %GTC vượt ngưỡng an toàn >52%.",
        "• BC Tây Nha Trang (Khánh Hòa - AM Hồng Bích Nga): Đột phá ca phát sáng trước 11h00, xử lý sạch backlog luân chuyển từ Hub Diên Khánh."
    ]
    for b in bullets_g1:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(2)
        r = p_b.add_run(b)
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(30, 41, 59)

    # Group 2: 6 BC Moi
    p_g2 = doc.add_paragraph()
    p_g2.paragraph_format.space_before = Pt(6)
    p_g2.paragraph_format.space_after = Pt(2)
    r_g2_t = p_g2.add_run("2. Báo động đỏ: 6 Bưu cục mới rơi vào diện cảnh báo bất ổn trong tuần W39 (W38 ✅ ➔ W39 🚨):")
    r_g2_t.bold = True
    r_g2_t.font.name = 'Arial'
    r_g2_t.font.size = Pt(10)
    r_g2_t.font.color.rgb = RGBColor(220, 38, 38)

    bullets_g2 = [
        "• BC D'Ran (Lâm Đồng - AM Phan Thành Long): %GTC đạt 36.86% (cảnh báo 6 ngày), địa bàn đèo dốc sạt lở và bưu tá thiếu hụt cục bộ.",
        "• BC Lang Biang 2 (Lâm Đồng - AM Trương Tuấn Anh): %GTC giảm về 35.80% (cảnh báo 6 ngày), áp lực hàng TikTok Shop đổ về vượt công suất kho.",
        "• BC Trường Xuân (Đắk Nông - AM Đỗ Duy Khang): %GTC tụt về 38.22% (cảnh báo 8 ngày), tuyến giao xa trung tâm phát sinh tỷ lệ không liên lạc được cao.",
        "• BC Đông Gia Nghĩa (Đắk Nông - AM Đỗ Duy Khang): %GTC đạt 44.89% (chớm dưới ngưỡng 45%), tỷ lệ gán đơn ca 2 buổi trưa chỉ đạt 72%.",
        "• BC Bắc Cam Ranh (Khánh Hòa - AM Phan Đình Duy): %GTC đạt 44.68% (đã tích lũy 31 ngày cảnh báo tích lũy), cần chấn chỉnh kỷ luật phát.",
        "• BC Phú Quý (Bình Thuận - AM Nguyễn Đình Uy): %GTC đạt 50.84%, tuy nhiên do phụ thuộc tàu cao tốc biển, chỉ số rơi sâu dưới 70% so với mốc lịch sử 88.55%."
    ]
    for b in bullets_g2:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(2)
        r = p_b.add_run(b)
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(30, 41, 59)

    # Group 3: 9 BC Man Tinh
    p_g3 = doc.add_paragraph()
    p_g3.paragraph_format.space_before = Pt(6)
    p_g3.paragraph_format.space_after = Pt(2)
    r_g3_t = p_g3.add_run("3. Điểm nóng mãn tính: 9 Bưu cục kẹt dai dẳng cả 2 tuần W38 và W39 (W38 🚨 ➔ W39 🚨):")
    r_g3_t.bold = True
    r_g3_t.font.name = 'Arial'
    r_g3_t.font.size = Pt(10)
    r_g3_t.font.color.rgb = RGBColor(194, 65, 12)

    bullets_g3 = [
        "• BC Lâm Viên 2 (Lâm Đồng - AM Trương Tuấn Anh): Thảm họa vận hành lớn nhất tuần, từ 46.2% ở W38 rơi tự do xuống 19.55% ở W39 (-26.65%p). Bưu tá tê liệt ca chiều.",
        "• BC Đức Trọng 1 (Lâm Đồng - AM Phan Thành Long): %GTC kẹt cứng ở 23.78% (W38: 23.90%), backlog tồn đọng lưu cữu trên 300 đơn.",
        "• BC Quảng Tín (Đắk Nông - AM Đỗ Duy Khang): %GTC đạt 26.17%, tích lũy kỷ lục hơn 100 ngày cảnh báo, quản trị kho lỏng lẻo.",
        "• BC Đơn Dương (Lâm Đồng - AM Phan Thành Long): %GTC giảm mạnh từ 45.3% xuống 29.01% (-16.29%p), gán ca 2 gần như bỏ trống.",
        "• BC Lang Biang 1 (Lâm Đồng - AM Trương Tuấn Anh): %GTC đạt 30.60% (W38: 32.80%), hàng cồng kềnh quá tải.",
        "• BC Kiến Đức (Đắk Nông - AM Đỗ Duy Khang): %GTC đạt 32.47% (W38: 35.60%), bán kính phát tuyến rộng >25km.",
        "• BC Tân Hà Lâm Hà (Lâm Đồng - AM Phan Thành Long): %GTC đạt 36.84% (W38: 38.20%), hàng nông sản đóng gói sai quy cách.",
        "• BC Tuy Đức (Đắk Nông - AM Đỗ Duy Khang): %GTC đạt 38.51% (W38: 41.50%), sóng viễn thông chập chờn ảnh hưởng app bưu tá.",
        "• BC Di Linh (Lâm Đồng - AM Huỳnh Thúc Duân): %GTC đạt 40.60% (W38: 44.20%), tỷ lệ hoàn hàng cao."
    ]
    for b in bullets_g3:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(2)
        r = p_b.add_run(b)
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3. COMPARISON TABLE
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    r_h3 = h3.add_run("📊 II. BẢNG TỔNG HỢP SO SÁNH ĐỐI CHIẾU CHI TIẾT 19 BƯU CỤC (W38 vs W39)")
    r_h3.bold = True
    r_h3.font.name = 'Arial'
    r_h3.font.size = Pt(11)
    r_h3.font.color.rgb = RGBColor(15, 23, 42)

    headers = ["STT", "Tên Bưu Cục", "Tỉnh", "AM Phụ Trách", "%GTC W38", "%GTC W39", "Biến Động", "Phân Loại", "Tình Trạng & Rủi Ro"]
    col_widths = [0.4, 1.4, 0.8, 1.1, 0.7, 0.7, 0.6, 0.8, 1.4]
    
    rows_data = [
        # Thoat (4)
        ["1", "Xuân Hương", "Lâm Đồng", "Lê Văn Trường", "42.1%", "51.8%", "▲ +9.7%", "Thoát ✅", "Giải tỏa sạch backlog, gán ca 1 đạt 94%"],
        ["2", "Cam Linh", "Khánh Hòa", "Phan Đình Duy", "36.7%", "48.5%", "▲ +11.8%", "Thoát ✅", "Bổ sung bưu tá tuyến trung tâm"],
        ["3", "Nhân Cơ", "Đắk Nông", "Đỗ Duy Khang", "43.5%", "52.3%", "▲ +8.8%", "Thoát ✅", "Phục hồi nhịp xuất tuyến đúng giờ"],
        ["4", "Tây Nha Trang", "Khánh Hòa", "Hồng Bích Nga", "44.1%", "49.6%", "▲ +5.5%", "Thoát ✅", "Tối ưu giao ca sáng trước 11h00"],
        
        # Moi phat sinh (6)
        ["5", "Phú Quý", "Bình Thuận", "Nguyễn Đình Uy", "88.5% (đỉnh)", "50.8%", "▼ Giảm sâu", "Mới 🚨", "Đảo xa, phụ thuộc tàu biển, hụt tải"],
        ["6", "D'Ran", "Lâm Đồng", "Phan Thành Long", "48.2%", "36.9%", "▼ -11.3%", "Mới 🚨", "Địa hình đèo dốc sạt lở, thiếu 2 bưu tá"],
        ["7", "Đông Gia Nghĩa", "Đắk Nông", "Đỗ Duy Khang", "47.1%", "44.9%", "▼ -2.2%", "Mới 🚨", "Chớm rớt ngưỡng 45%, gán ca 2 yếu"],
        ["8", "Bắc Cam Ranh", "Khánh Hòa", "Phan Đình Duy", "46.5%", "44.7%", "▼ -1.8%", "Mới 🚨", "Tích lũy 31 ngày cảnh báo, trễ giao"],
        ["9", "Trường Xuân", "Đắk Nông", "Đỗ Duy Khang", "49.0%", "38.2%", "▼ -10.8%", "Mới 🚨", "Cảnh báo 8 ngày, tỷ lệ KLL tăng vọt"],
        ["10", "Lang Biang 2", "Lâm Đồng", "Trương Tuấn Anh", "47.8%", "35.8%", "▼ -12.0%", "Mới 🚨", "Cảnh báo 6 ngày, TTS dồn ứ cục bộ"],
        
        # Man tinh (9)
        ["11", "Lâm Viên 2", "Lâm Đồng", "Trương Tuấn Anh", "46.2%", "19.55%", "▼ -26.6%", "Mãn tính 🚨", "Rơi tự do nặng nhất mạng, liệt ca 2"],
        ["12", "Đức Trọng 1", "Lâm Đồng", "Phan Thành Long", "23.9%", "23.8%", "▼ -0.1%", "Mãn tính 🚨", "Kẹt cứng 2 tuần, backlog >300 đơn"],
        ["13", "Quảng Tín", "Đắk Nông", "Đỗ Duy Khang", "18.1%", "26.2%", "▲ +8.1%", "Mãn tính 🚨", "Tích lũy >100 ngày cảnh báo, tồn kho"],
        ["14", "Đơn Dương", "Lâm Đồng", "Phan Thành Long", "45.3%", "29.0%", "▼ -16.3%", "Mãn tính 🚨", "Năng suất ca chiều sụt giảm nghiêm trọng"],
        ["15", "Lang Biang 1", "Lâm Đồng", "Trương Tuấn Anh", "32.8%", "30.6%", "▼ -2.2%", "Mãn tính 🚨", "Hàng cồng kềnh nông thôn quá tải"],
        ["16", "Kiến Đức", "Đắk Nông", "Đỗ Duy Khang", "35.6%", "32.5%", "▼ -3.1%", "Mãn tính 🚨", "Bán kính phát rộng >25km đường đồi núi"],
        ["17", "Tân Hà Lâm Hà", "Lâm Đồng", "Phan Thành Long", "38.2%", "36.8%", "▼ -1.4%", "Mãn tính 🚨", "Đóng gói sai quy cách, tồn aging >5d"],
        ["18", "Tuy Đức", "Đắk Nông", "Đỗ Duy Khang", "41.5%", "38.5%", "▼ -3.0%", "Mãn tính 🚨", "Sóng mạng chập chờn, cập nhật app chậm"],
        ["19", "Di Linh", "Lâm Đồng", "Huỳnh Thúc Duân", "44.2%", "40.6%", "▼ -3.6%", "Mãn tính 🚨", "Tỷ lệ hoàn hàng cao, trễ giờ quét"]
    ]

    t_comp = doc.add_table(rows=1, cols=len(headers))
    style_table(t_comp, col_widths, headers, rows_data, header_bg="0F4C81")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 4. ACTION PLAN
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(8)
    h4.paragraph_format.space_after = Pt(4)
    r_h4 = h4.add_run("📋 III. KẾ HOẠCH HÀNH ĐỘNG CẤP BÁCH TUẦN W40 (3 MỆNH LỆNH TÁC CHIẾN)")
    r_h4.bold = True
    r_h4.font.name = 'Arial'
    r_h4.font.size = Pt(11)
    r_h4.font.color.rgb = RGBColor(15, 23, 42)

    headers_plan = ["Mệnh Lệnh", "Nội Dung Tác Chiến Trọng Tâm", "Chỉ Tiêu Cam Kết SLA", "Đầu Mối Chịu Trách Nhiệm", "Hạn Chót"]
    col_widths_plan = [1.1, 2.3, 1.7, 1.2, 0.6]
    rows_plan = [
        [
            "1. Cap Volume\n(Điều tiết tải)",
            "Tạm thời giảm 25% hạn mức chia chọn hàng về trong 3 ngày đầu tuần tại BC Lâm Viên 2 và Quảng Tín",
            "Bưu tá dồn 100% sức quét sạch tồn kho và aging >5 ngày, đưa tồn về <1.0 ngày",
            "AM Tuấn Anh, Đỗ Duy Khang & Khối Vận Hành",
            "30/09"
        ],
        [
            "2. Chi Viện Ca 1\n(Cứu hỏa tuyến)",
            "Điều chuyển chi viện 4 bưu tá cứng từ cụm Xuân Hương và Gia Nghĩa sang phát lượt 1 trước 11h00 trưa tại Lâm Viên 2 và Đức Trọng 1",
            "Nâng %GTC ca 1 sáng từ 41% lên ≥70%; đảm bảo xuất tuyến đúng 08h00",
            "AM Phan Thành Long, Trương Tuấn Anh",
            "01/10"
        ],
        [
            "3. Cam Kết Thoát\n(KPI Vùng W40)",
            "Đồng loạt triển khai tổ phản ứng nhanh, giám sát trực tiếp hàng ngày qua dashboard real-time",
            "Kéo tối thiểu 5 bưu cục thoát cảnh báo; giảm tổng số BC cảnh báo toàn vùng từ 15 về <10 BC",
            "Giám Đốc Vận Hành & Toàn bộ AM phụ trách",
            "04/10"
        ]
    ]
    t_plan = doc.add_table(rows=1, cols=len(headers_plan))
    style_table(t_plan, col_widths_plan, headers_plan, rows_plan, header_bg="1E293B")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Footer note
    p_ft = doc.add_paragraph()
    r_ft = p_ft.add_run("Tài liệu được khởi tạo tự động từ hệ thống điều hành dữ liệu tập trung GHN Express — Vùng Nam Trung Bộ.")
    r_ft.italic = True
    r_ft.font.name = 'Arial'
    r_ft.font.size = Pt(8.5)
    r_ft.font.color.rgb = RGBColor(148, 163, 184)

    # Save to disk
    out_name = 'KICH_BAN_CANH_BAO_BUU_CUC_W38_VS_W39.docx'
    doc.save(out_name)
    print(f"Successfully generated {out_name} (Size: {os.path.getsize(out_name)} bytes)")

if __name__ == '__main__':
    build_docx()
