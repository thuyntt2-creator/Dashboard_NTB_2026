import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.path.insert(0, os.path.join(os.getcwd(), 'scratch'))
from doc_builder_helpers import format_run, set_cell_borders_and_shading

def add_redesigned_callout(doc, section_title, speech_title, province_block, am_analysis_block, insights=None, warnings=None, actions=None):
    # Section title
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(14)
    p_sec.paragraph_format.space_after = Pt(6)
    p_sec.paragraph_format.keep_with_next = True
    r_sec = p_sec.add_run(section_title)
    format_run(r_sec, font_size_pt=13, bold=True, color_rgb=(0x0F, 0x4C, 0x81)) # Navy
    
    # 1x1 table
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_borders_and_shading(cell, border_color="F97316", fill_color="FFF7ED")
    
    # First paragraph: Speech title
    p_first = cell.paragraphs[0]
    p_first.paragraph_format.space_before = Pt(2)
    p_first.paragraph_format.space_after = Pt(4)
    p_first.paragraph_format.line_spacing = 1.15
    r_title = p_first.add_run(speech_title)
    format_run(r_title, font_size_pt=11.5, bold=True, color_rgb=(0xEA, 0x58, 0x0C)) # Orange
    
    # --- KHỐI 1: CHỈ SỐ 5 TỈNH THÀNH (ĐẶT LÊN ĐẦU TIÊN) ---
    if province_block:
        p_prov_header = cell.add_paragraph()
        p_prov_header.paragraph_format.space_before = Pt(4)
        p_prov_header.paragraph_format.space_after = Pt(2)
        p_prov_header.paragraph_format.line_spacing = 1.15
        r_ph = p_prov_header.add_run("📍 1. BẢNG CHỈ SỐ TOÀN VÙNG & 5 TỈNH THÀNH (W38 vs W37):")
        format_run(r_ph, font_size_pt=10.5, bold=True, color_rgb=(0x0F, 0x4C, 0x81))
        
        for pline in province_block:
            p_p = cell.add_paragraph()
            p_p.paragraph_format.space_before = Pt(1)
            p_p.paragraph_format.space_after = Pt(2)
            p_p.paragraph_format.line_spacing = 1.15
            p_p.paragraph_format.left_indent = Inches(0.15)
            r_p = p_p.add_run(pline)
            format_run(r_p, font_size_pt=10.5, color_rgb=(0x1F, 0x29, 0x37))

    # --- KHỐI 2: PHÂN TÍCH CHI TIẾT THEO 18 AM ---
    if am_analysis_block:
        p_am_header = cell.add_paragraph()
        p_am_header.paragraph_format.space_before = Pt(6)
        p_am_header.paragraph_format.space_after = Pt(2)
        p_am_header.paragraph_format.line_spacing = 1.15
        r_ah = p_am_header.add_run("👤 2. BÓC TÁCH CHI TIẾT HIỆU SUẤT THEO QUẢN LÝ VẬN HÀNH (AM):")
        format_run(r_ah, font_size_pt=10.5, bold=True, color_rgb=(0x0F, 0x4C, 0x81))
        
        for aline in am_analysis_block:
            p_a = cell.add_paragraph()
            p_a.paragraph_format.space_before = Pt(2)
            p_a.paragraph_format.space_after = Pt(3)
            p_a.paragraph_format.line_spacing = 1.15
            r_a = p_a.add_run(aline)
            format_run(r_a, font_size_pt=10.5, color_rgb=(0x1F, 0x29, 0x37))
            
    # Insights
    if insights:
        p_in = cell.add_paragraph()
        p_in.paragraph_format.space_before = Pt(6)
        p_in.paragraph_format.space_after = Pt(2)
        p_in.paragraph_format.line_spacing = 1.15
        r_in_tag = p_in.add_run("🔍 INSIGHT & ĐIỂM NGHẼN THỰC TẾ TẠI CÁC AM:\n")
        format_run(r_in_tag, font_size_pt=10.5, bold=True, color_rgb=(0x0F, 0x4C, 0x81))
        for ins in insights:
            p_b = cell.add_paragraph()
            p_b.paragraph_format.space_before = Pt(1)
            p_b.paragraph_format.space_after = Pt(2)
            p_b.paragraph_format.line_spacing = 1.15
            p_b.paragraph_format.left_indent = Inches(0.2)
            r_b = p_b.add_run(f"• {ins}")
            format_run(r_b, font_size_pt=10, color_rgb=(0x1F, 0x29, 0x37))

    # Warnings
    if warnings:
        p_w = cell.add_paragraph()
        p_w.paragraph_format.space_before = Pt(6)
        p_w.paragraph_format.space_after = Pt(2)
        p_w.paragraph_format.line_spacing = 1.15
        r_w_tag = p_w.add_run("⚠️ CẢNH BÁO ĐỎ TẬP TRUNG THEO AM:\n")
        format_run(r_w_tag, font_size_pt=10.5, bold=True, color_rgb=(0xDC, 0x26, 0x26))
        for w in warnings:
            p_wb = cell.add_paragraph()
            p_wb.paragraph_format.space_before = Pt(1)
            p_wb.paragraph_format.space_after = Pt(2)
            p_wb.paragraph_format.line_spacing = 1.15
            p_wb.paragraph_format.left_indent = Inches(0.2)
            r_wb = p_wb.add_run(f"• {w}")
            format_run(r_wb, font_size_pt=10, color_rgb=(0x99, 0x1B, 0x1B))

    # Actions
    if actions:
        p_act = cell.add_paragraph()
        p_act.paragraph_format.space_before = Pt(6)
        p_act.paragraph_format.space_after = Pt(2)
        p_act.paragraph_format.line_spacing = 1.15
        r_act_tag = p_act.add_run("🎯 GIAO VIỆC & NHIỆM VỤ TÁC CHIẾN ĐÍCH DANH CHO TỪNG AM:\n")
        format_run(r_act_tag, font_size_pt=10.5, bold=True, color_rgb=(0x15, 0x80, 0x3D))
        for a in actions:
            p_ab = cell.add_paragraph()
            p_ab.paragraph_format.space_before = Pt(1)
            p_ab.paragraph_format.space_after = Pt(2)
            p_ab.paragraph_format.line_spacing = 1.15
            p_ab.paragraph_format.left_indent = Inches(0.2)
            r_ab = p_ab.add_run(f"• {a}")
            format_run(r_ab, font_size_pt=10, color_rgb=(0x14, 0x53, 0x2D))

    # Spacer after table
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(4)

print("Redesigned callout helper ready.")
