# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'c:\Users\lap4all\Desktop\New folder\Bao_Cao_Hop_Van_Hanh_13h_NTB.xlsx', data_only=True)
ws = wb['ODR Tuần 37']
print("ODR Tuần 37 max rows:", ws.max_row, "cols:", ws.max_column)
for r in range(1, min(ws.max_row + 1, 40)):
    vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 15))]
    if any(vals):
        print(f"Row {r:2d}: {vals}")
