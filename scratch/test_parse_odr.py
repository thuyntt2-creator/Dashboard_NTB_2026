import openpyxl, json, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W38_2026.xlsx', data_only=True)
ws = wb['05_ODR']

# 1. Overview
overview = []
r_full = [ws.cell(6, c).value for c in range(1, 8)]
r_tts = [ws.cell(7, c).value for c in range(1, 8)]
overview.append({
    'label': 'Full hàng',
    'w35': r_full[1],
    'w36': r_full[2],
    'w37': r_full[3],
    'w38': r_full[4],
    'diff': r_full[5]
})
overview.append({
    'label': 'TTS',
    'w35': r_tts[1],
    'w36': r_tts[2],
    'w37': r_tts[3],
    'w38': r_tts[4],
    'diff': r_tts[5]
})
print("Parsed Overview:")
print(json.dumps(overview, indent=2, ensure_ascii=False))

# Let's check AM province mapping from cocau or existing am_full in data.json
with open('data.json', encoding='utf-8') as f:
    d = json.load(f)

am_prov_map = {}
for item in d.get('odr', {}).get('am_full', []) + d.get('san_luong', {}).get('am_full', []):
    if item.get('am') and item.get('tinh'):
        am_prov_map[item['am']] = item['tinh']

# 2. AM Full
am_full = []
for r in range(12, 30):
    am = ws.cell(r, 1).value
    if not am: continue
    sl = ws.cell(r, 2).value or 0
    w35 = ws.cell(r, 3).value
    w36 = ws.cell(r, 4).value
    w37 = ws.cell(r, 5).value
    w38 = ws.cell(r, 6).value
    diff = ws.cell(r, 7).value
    tinh = am_prov_map.get(am, '')
    am_full.append({
        'am': am,
        'tinh': tinh,
        'vol': sl,
        'w35': w35,
        'w36': w36,
        'w37': w37,
        'w38': w38,
        'diff': diff
    })

print(f"\nParsed {len(am_full)} AM Full rows. Sample:")
print(json.dumps(am_full[:2], indent=2, ensure_ascii=False))

# 3. AM TTS
am_tts = []
for r in range(34, 52):
    am = ws.cell(r, 1).value
    if not am: continue
    sl = ws.cell(r, 2).value or 0
    w35 = ws.cell(r, 3).value
    w36 = ws.cell(r, 4).value
    w37 = ws.cell(r, 5).value
    w38 = ws.cell(r, 6).value
    diff = ws.cell(r, 7).value
    tinh = am_prov_map.get(am, '')
    am_tts.append({
        'am': am,
        'tinh': tinh,
        'vol': sl,
        'w35': w35,
        'w36': w36,
        'w37': w37,
        'w38': w38,
        'diff': diff
    })
print(f"\nParsed {len(am_tts)} AM TTS rows. Sample:")
print(json.dumps(am_tts[:2], indent=2, ensure_ascii=False))

# 4. Tinh Full
tinh_full = []
for r in range(56, 61):
    tinh = ws.cell(r, 1).value
    if not tinh: continue
    sl = ws.cell(r, 2).value or 0
    w35 = ws.cell(r, 3).value
    w36 = ws.cell(r, 4).value
    w37 = ws.cell(r, 5).value
    w38 = ws.cell(r, 6).value
    diff = ws.cell(r, 7).value
    tinh_full.append({
        'am': tinh,
        'tinh': tinh,
        'vol': sl,
        'w35': w35,
        'w36': w36,
        'w37': w37,
        'w38': w38,
        'diff': diff
    })
print(f"\nParsed {len(tinh_full)} Tinh Full rows:")
print(json.dumps(tinh_full, indent=2, ensure_ascii=False))

# 5. Tinh TTS
tinh_tts = []
for r in range(65, 70):
    tinh = ws.cell(r, 1).value
    if not tinh: continue
    sl = ws.cell(r, 2).value or 0
    w35 = ws.cell(r, 3).value
    w36 = ws.cell(r, 4).value
    w37 = ws.cell(r, 5).value
    w38 = ws.cell(r, 6).value
    diff = ws.cell(r, 7).value
    tinh_tts.append({
        'am': tinh,
        'tinh': tinh,
        'vol': sl,
        'w35': w35,
        'w36': w36,
        'w37': w37,
        'w38': w38,
        'diff': diff
    })
print(f"\nParsed {len(tinh_tts)} Tinh TTS rows:")
print(json.dumps(tinh_tts, indent=2, ensure_ascii=False))
