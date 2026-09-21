import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document("KICH_BAN_THUYET_TRINH_W38_INSIGHT_CHUYEN_SAU.docx")

print(f"Total tables: {len(doc.tables)}")

for i, t in enumerate(doc.tables):
    cell = t.rows[0].cells[0]
    full_text = "\n".join([p.text for p in cell.paragraphs])
    has_am = "AM " in full_text or "am " in full_text.lower()
    # Find AM names mentioned
    ams_found = []
    for am in ["Khánh", "Thủy", "Nghĩa", "Thư", "Phi", "Nhựt", "Long", "Lợi", "Linh", "Trường", "Tiến", "Nga", "Chi", "Duân", "Nhung", "Thơ", "Vũ", "Duy"]:
        if am in full_text:
            ams_found.append(am)
    print(f"Section {i+1}: Has AM: {has_am} | AMs: {', '.join(ams_found[:6])} | Length: {len(full_text)} chars")
