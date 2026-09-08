import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W36_2026.xlsx', data_only=True)
ws = wb['08_OPR TTS']

ams = []
for r in range(6, 23):
    am = ws.cell(r, 1).value
    vol_day = ws.cell(r, 2).value or 0
    opr_day = ws.cell(r, 4).value or 0
    if opr_day > 1: opr_day = 0 # Fix Lê Minh Lợi 2 -> 0
    
    vol_night = ws.cell(r, 6).value or 0
    opr_night = ws.cell(r, 8).value or 0
    
    vol_total = ws.cell(r, 10).value or (vol_day + vol_night)
    opr_total = ws.cell(r, 12).value or 0
    
    # Tính số đơn trễ/rớt OPR
    fail_day = round(vol_day * (1 - opr_day))
    fail_night = round(vol_night * (1 - opr_night))
    fail_total = round(vol_total * (1 - opr_total))
    
    ams.append({
        'am': str(am).strip(),
        'vol_total': vol_total,
        'opr_total': opr_total,
        'fail_total': fail_total,
        'vol_day': vol_day,
        'opr_day': opr_day,
        'fail_day': fail_day,
        'vol_night': vol_night,
        'opr_night': opr_night,
        'fail_night': fail_night
    })

total_fail_region = sum(x['fail_total'] for x in ams)
total_vol_region = sum(x['vol_total'] for x in ams)
total_fail_day = sum(x['fail_day'] for x in ams)
total_fail_night = sum(x['fail_night'] for x in ams)

print(f"Toàn Vùng OPR TTS W36:")
print(f"Tổng đơn: {total_vol_region}, Tổng đơn rớt OPR: {total_fail_region} (Tỷ lệ rớt OPR: {total_fail_region/total_vol_region*100:.1f}%)")
print(f"Rớt OPR Ca Ngày (9h-19h): {total_fail_day} đơn")
print(f"Rớt OPR Ca Đêm (19h-9h): {total_fail_night} đơn\n")

print(f"{'AM':22} | {'Tổng đơn':8} | {'%OPR W36':9} | {'Đơn Rớt OPR':11} | {'% Tỷ Trọng Rớt':14}")
print("-" * 75)
for x in sorted(ams, key=lambda a: a['fail_total'], reverse=True):
    rate_fail = (x['fail_total'] / total_fail_region * 100) if total_fail_region > 0 else 0
    print(f"{x['am']:22} | {x['vol_total']:8} | {x['opr_total']*100:8.1f}% | {x['fail_total']:11} | {rate_fail:13.1f}%")
