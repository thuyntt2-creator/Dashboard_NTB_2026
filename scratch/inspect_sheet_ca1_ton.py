import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W40_2026.xlsx', data_only=True)

def inspect_sheet(name):
    if name not in wb.sheetnames:
        print(f"Sheet {name} not found")
        return
    ws = wb[name]
    print(f"\n================ SHEET: {name} ================")
    for r in range(1, 40):
        row_vals = [ws.cell(r, c).value for c in range(1, 15)]
        if any(row_vals):
            compact = [f"C{c}:{v}" for c, v in enumerate(row_vals, 1) if v is not None]
            print(f"R{r:02d}: " + " | ".join(compact[:7]))

inspect_sheet('03_GTC Ca1+Ton')
inspect_sheet('03b_GTC Ca1 thuan')
