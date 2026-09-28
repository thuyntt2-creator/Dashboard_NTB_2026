import docx, sys, shutil
sys.stdout.reconfigure(encoding='utf-8')

src = 'KICH_BAN_THUYET_TRINH_MOI_NHAT.docx'
doc = docx.Document(src)

# TABLE 7: MỤC VIII - HIỆU SUẤT %OPR TIKTOK SHOP W39
cell7 = doc.tables[7].rows[0].cells[0]
for p in cell7.paragraphs[1:]:
    p._p.getparent().remove(p._p)

cell7.paragraphs[0].text = "🗣️ HIỆU SUẤT %OPR TIKTOK SHOP CA NGÀY VÀ CA ĐÊM (TARGET SLA ≥ 80.0%) W39 vs W38:"

sec8 = [
    '📍 1. BẢNG CHỈ SỐ %OPR TIKTOK SHOP 5 TỈNH THÀNH (W39 vs W38):',
    '• Top 1 - Ninh Thuận: OPR Tổng đạt 91,13% (+0,78%p) | Ca ngày: 96,49% | Ca đêm: 77,46% ➔ Quán quân toàn vùng, vượt xa target ≥80%!',
    '• Top 2 - Khánh Hòa: OPR Tổng đạt 83,22% (-2,09%p) | Ca ngày: 92,95% | Ca đêm: 74,01% ➔ Vững vàng đạt chuẩn xanh.',
    '• Top 3 - Lâm Đồng: OPR Tổng đạt 75,39% (+5,41%p) | Ca ngày: 87,74% | Ca đêm: 47,01% ➔ Có cải thiện nhưng chưa đạt chuẩn 80%.',
    '• Top 4 - Bình Thuận: OPR Tổng đạt 74,36% (-13,69%p) | Ca ngày: 88,70% | Ca đêm: 52,37% (-30,3%p) ➔ Báo động sụt giảm ca đêm.',
    '• Top 5 - Đắk Nông: OPR Tổng đạt 50,00% | Ca ngày: 88,50% | Ca đêm: 0,57% ➔ Báo động đỏ: Ca đêm gần như tê liệt hoàn toàn!',
    '➔ TOÀN VÙNG W39: OPR Tổng đạt 80,60% (vừa chạm chuẩn SLA ≥80.0%) | Ca ngày giữ vững phong độ rất cao 91,2% ➔ Ca đêm rớt sâu về 58,4% (chênh lệch tới 32,8%p).',
    '',
    '🎙️ LỜI THOẠI THUYẾT TRÌNH %OPR TIKTOK SHOP (KHI BẬT TAB 8 TRÊN DASHBOARD):',
    '\"Dạ tiếp theo em xin phép chuyển qua Tab 8 là chỉ số %OPR TikTok Shop (tỷ lệ xử lý đơn hàng sàn đúng quy trình, target SLA yêu cầu phải đạt từ 80% trở lên):',
    '',
    'Nhìn vào tỷ lệ toàn vùng tuần W39, OPR Tổng của mình đạt 80.6%, tức là vừa đủ chạm chuẩn xanh 80% của công ty. Tuy nhiên, khi bóc tách giữa Ca ngày và Ca đêm thì bức tranh lại bị phân hóa cực kỳ gay gắt:',
    '',
    '• Ở Ca ngày: Toàn vùng mình làm rất chuẩn chỉ, tỷ lệ OPR đạt tới hơn 91.2%, tỉnh nào cũng đạt từ 88% đến 96%.',
    '• Nhưng tử huyệt lại nằm ở Ca đêm: OPR ca đêm toàn vùng bị rớt tự do về 58.4%, tức là chênh lệch tới gần 33%p so với ca ngày! Đơn sàn dồn về đêm xử lý rất chậm, dẫn đến nguy cơ chậm luân chuyển sang sáng hôm sau.',
    '',
    'Đánh giá theo 5 tỉnh thành:',
    '• Tuyên dương Ninh Thuận tiếp tục là ngọn cờ đầu: OPR Tổng đạt 91.1% (Ca ngày 96.5%, Ca đêm vẫn giữ được 77.5%), vượt xa chuẩn SLA 80%.',
    '• Khánh Hòa giữ vững phong độ Top 2 với 83.2% OPR Tổng.',
    '• Ngược lại, Bình Thuận tuần này bị tụt sâu mất gần 14%p (chỉ còn 74.4%) do ca đêm bị gãy cánh rớt về 52.4%.',
    '• Báo động đỏ nặng nhất là Đắk Nông: OPR Tổng chỉ đạt đúng 50.0%. Đáng nói là Ca ngày Đắk Nông làm được 88.5%, nhưng Ca đêm gần như buông xuôi, rớt thảm hại về 0.57% — tức là ca đêm gần như không xử lý được đơn nào!',
    '',
    'Bóc tách chi tiết theo 18 AM bên dưới:',
    '',
    'Nhóm AM giữ OPR xuất sắc trên 90%:',
    '• Anh Nguyễn Duy Long (OPR Tổng 90.7%, Ca ngày đạt 96.1%).',
    '• Chị Thái Thị Thanh Thư (OPR Tổng 90.2%, Ca ngày đạt 94.5%).',
    '• Và anh Nguyễn Lê Nguyên Vũ ở Ninh Thuận (OPR Tổng đạt 92.3%).',
    '',
    'Các điểm nóng cần chấn chỉnh khẩn cấp về OPR Ca đêm:',
    '• Chỗ chị Trần Thị Nhung: OPR Tổng tuần này chỉ đạt 25.3% do ca đêm bị rớt xuống mức 0.74%!',
    '• Chỗ anh Lê Văn Trường ở Lâm Đồng: OPR Tổng chỉ đạt 53.5% (ca đêm chỉ đạt 24.0%).',
    '• Chỗ anh Nguyễn Hoàng Phi (59.3%) và anh Lê Thanh Nhựt (67.6%): Ca đêm của 2 anh cũng chỉ đạt từ 34% đến 36%.',
    '',
    '🎯 Về nhiệm vụ tác chiến cho chỉ số OPR tuần W40:',
    '• Em đề nghị các AM có OPR ca đêm dưới 50% (đặc biệt là chị Nhung, anh Trường, anh Phi, anh Nhựt và các bưu cục Đắk Nông): Bắt buộc phải bố trí lại nhân sự trực ca đêm tại bưu cục, đẩy nhanh tốc độ quét nhập và xử lý đơn sàn trước 22h00 đêm, không để dồn ứ sang ca sáng hôm sau.',
    '• Quyết tâm tuần W40 đưa OPR Ca đêm toàn vùng vượt qua mốc 70% và OPR Tổng vượt mốc 85%.\"'
]

for line in sec8:
    p = cell7.add_paragraph()
    p.text = line

doc.save(src)
for c in ['KICH_BAN_THUYET_TRINH_W39_INSIGHT_CHUYEN_SAU.docx', 'KICH_BAN_THUYET_TRINH_W39_NAM_TRUNG_BO.docx']:
    shutil.copyfile(src, c)

print('Updated Table 7 (OPR TTS) with exact W39 data successfully!')
