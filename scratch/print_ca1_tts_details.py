import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W40_2026.xlsx', data_only=True)
ws = wb['03b_GTC Ca1 thuan']

print("=== GTC CA 1 THUẦN — TTS (THEO TỈNH) ===")
for r in range(67, 72):
    tinh = ws.cell(r, 1).value
    vol = float(ws.cell(r, 2).value or 0)
    w37 = float(ws.cell(r, 3).value or 0)
    w38 = float(ws.cell(r, 4).value or 0)
    w39 = float(ws.cell(r, 5).value or 0)
    w40 = float(ws.cell(r, 6).value or 0)
    diff = float(ws.cell(r, 7).value or 0)
    print(f"{tinh:<15} | W39: {w39:.2%} -> W40: {w40:.2%} | Δ: {diff*100:+.2f}%p | Vol: {int(vol):,}")

print("\n=== GTC CA 1 THUẦN — TTS (THEO AM) ===")
am_ca1_tts = []
for r in range(35, 54):
    am = ws.cell(r, 1).value
    vol = float(ws.cell(r, 2).value or 0)
    w37 = float(ws.cell(r, 3).value or 0)
    w38 = float(ws.cell(r, 4).value or 0)
    w39 = float(ws.cell(r, 5).value or 0)
    w40 = float(ws.cell(r, 6).value or 0)
    diff = float(ws.cell(r, 7).value or 0)
    if am:
        am_ca1_tts.append({
            'am': am,
            'vol': int(vol),
            'w39': w39,
            'w40': w40,
            'diff': diff
        })

print("\n--- XẾP THEO BIẾN ĐỘNG TĂNG TRƯỞNG Δ GTC CA 1 THUẦN TTS (WoW) ---")
for i, r in enumerate(sorted(am_ca1_tts, key=lambda x: x['diff'], reverse=True)):
    print(f"#{i+1:2d}: {r['am']:<25} | W39: {r['w39']:.2%} -> W40: {r['w40']:.2%} | Δ: {r['diff']*100:+.2f}%p | Vol TTS: {r['vol']:,}")

print("\n--- XẾP THEO %GTC CA 1 THUẦN TTS TUYỆT ĐỐI W40 ---")
for i, r in enumerate(sorted(am_ca1_tts, key=lambda x: x['w40'], reverse=True)):
    print(f"#{i+1:2d}: {r['am']:<25} | GTC Ca 1 TTS W40: {r['w40']:.2%} (W39: {r['w39']:.2%}) | Vol TTS: {r['vol']:,}")
