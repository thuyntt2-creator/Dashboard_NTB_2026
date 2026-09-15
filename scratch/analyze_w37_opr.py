import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W37_2026.xlsx', data_only=True)
ws = wb['08_OPR TTS']

ams = []
for r in range(6, 23):
    am = ws.cell(r, 1).value
    if not am: continue
    tot_vol = ws.cell(r, 10).value or 0
    w36 = ws.cell(r, 11).value or 0
    w37 = ws.cell(r, 12).value or 0
    fail = round(tot_vol * (1 - w37)) if tot_vol else 0
    ams.append({'am': am, 'vol': tot_vol, 'w36': w36, 'w37': w37, 'fail': fail})

print("--- TOP AM GÂY RỚT ĐƠN W37 ---")
ams_fail_sorted = sorted(ams, key=lambda x: x['fail'], reverse=True)
for a in ams_fail_sorted:
    pct_fail = (a['fail'] / 1921.0) * 100 # 8771 - 6850 = 1921 total fails
    if a['fail'] > 0:
        print(f"{a['am']}: W37={a['w37']:.1%}, fail={a['fail']}/{a['vol']} ({pct_fail:.1f}% tổng rớt vùng)")

print("\n--- RANK AM THEO %OPR W37 ---")
ams_w37_sorted = sorted(ams, key=lambda x: x['w37'], reverse=True)
for a in ams_w37_sorted:
    print(f"{a['am']}: W37={a['w37']:.1%}, W36={a['w36']:.1%}, vol={a['vol']}")
