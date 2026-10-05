import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

file_path = r"C:\Users\lap4all\Downloads\kịch bản.docx"
doc = docx.Document(file_path)

def print_tbl(idx):
    if idx < len(doc.tables):
        print(f"\n==================== TABLE {idx+1} FULL ====================")
        for r in doc.tables[idx].rows:
            for c in r.cells:
                print(c.text.strip())

print_tbl(1) # Table 2 (Sản Lượng)
print_tbl(2) # Table 3 (GTC Tổng)
print_tbl(3) # Table 4 (GTC Ca 1 TTS)
print_tbl(4) # Table 5 (Tỷ lệ gán)
