import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W40_2026.xlsx', data_only=True)

def print_sheet_summary(name, max_r=25):
    if name not in wb.sheetnames:
        return
    ws = wb[name]
    print(f"\n================ SHEET: {name} ================")
    for r in range(1, max_r):
        vals = [ws.cell(r, c).value for c in range(1, 10)]
        if any(vals):
            non_empty = [f"C{c}:{str(v)[:20]}" for c, v in enumerate(vals, 1) if v is not None]
            print(f"R{r:02d}: " + " | ".join(non_empty[:6]))

for s in ['05_ODR', '06_LTC', '07_Gan', '08_OPR TTS', '09_Rot LC', '12_FD']:
    print_sheet_summary(s, 20)
