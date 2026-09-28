import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

table12_content = """🗣️ KIỂM SOÁT THU HỘ COD 80,7 TỶ ₫ — TỶ LỆ TIỀN MẶT GIẢM XUỐNG 40,1% (CHUYỂN KHOẢN VIETQR ĐẠT 59,9%) W39 vs W38:

📍 1. BẢNG CHỈ SỐ TOÀN VÙNG (W39 vs W38):
• Tổng COD thu hộ toàn vùng W39: Đạt 80.733,6 Triệu VNĐ (80,7 Tỷ đồng), giảm -4,4% (-3,7 tỷ) so với W38 (84.441,1 Triệu VNĐ).
• Tiền mặt: Đạt 32.404,7 Triệu VNĐ (32,4 Tỷ đồng), chiếm 40,1% tổng dòng tiền, GIẢM MẠNH -3,0%p WoW (W38 là 43,1%). Giá trị tiền mặt lưu thông giảm tới -4.023,1 Triệu VNĐ (-11,0% WoW) ➔ Giảm áp lực rủi ro thâm hụt quỹ rất lớn!
• Chuyển khoản (VietQR / Thanh toán số): Đạt 48.329,0 Triệu VNĐ (48,3 Tỷ đồng), chiếm tới 59,9% (~60% tổng COD toàn vùng), TĂNG TRƯỞNG +3,0%p WoW so với W38 (56,9%).

👤 2. BÓC TÁCH CHI TIẾT THEO QUẢN LÝ VẬN HÀNH (AM) — PHÂN HÓA RÕ RỆT:

🔴 NHÓM TIỀN MẶT CAO BÁO ĐỘNG (≥ 70% TIỀN MẶT — CẦN GIẢM GẤP): Gồm 3 AM:
  1. AM Trương Quang Linh: Tỷ lệ tiền mặt lên tới 80,6% (tăng vọt +6,4%p từ 74,2% ở W38). Đây là AM DUY NHẤT trong nhóm tiền mặt cao bị tăng mạnh ngược xu thế toàn vùng!
  2. AM Lê Thanh Nhựt: Tỷ lệ tiền mặt 77,7% (giảm nhẹ -1,1%p từ 78,9%), vẫn giữ tỷ lệ tiền mặt rất cao tại Bình Thuận.
  3. AM Huỳnh Thúc Duân: Tỷ lệ tiền mặt 71,5% (đã giảm được -5,0%p từ 76,5% ở W38).

🟡 NHÓM CẦN CẢI THIỆN TIẾP (50% - 70% TIỀN MẶT):
  4. AM Huỳnh Thị Kim Chi: 68,6% tiền mặt (giảm -5,3%p từ 73,9%).
  5. AM Trần Thị Nhung: 62,8% tiền mặt (giảm mạnh -7,6%p từ 70,5%).
  6. AM Nguyễn Đỗ Minh Nghĩa: 51,4% tiền mặt (giảm -2,2%p từ 53,6%).
  7. AM Nguyễn Thị Tuyết Thơ: 50,0% tiền mặt (tăng +8,2%p từ 41,8%).

🟢 NHÓM KIỂM SOÁT THANH TOÁN SỐ XUẤT SẮC (TIỀN MẶT DƯỚI 50% — VIETQR ÁP ĐẢO):
  8. AM Lê Văn Trường: Tiền mặt 45,5% (-3,4%p), VietQR đạt 54,5%.
  9. AM Hồng Bích Nga: Tiền mặt 41,2% (-4,9%p), VietQR đạt 58,8%.
  10. AM Phan Đình Duy: Tiền mặt 40,0% (-6,6%p), VietQR đạt 60,0%.
  11. AM Nguyễn Hoàng Phi: Tiền mặt 38,9% (-3,7%p), VietQR đạt 61,1%.
  12. AM Nguyễn Lê Nguyên Vũ: Tiền mặt 38,8% (-7,3%p), VietQR đạt 61,2%.
  13. AM Nguyễn Thanh Long (Cam Ranh): Tiền mặt 35,4% (+3,0%p), VietQR đạt 64,6%.
  14. AM Nguyễn Ngọc Khánh (Nha Trang): Tiền mặt 30,8% (+0,2%p), VietQR đạt 69,2%.
  15. AM Lê Minh Lợi: Tiền mặt giảm sâu từ 37,0% xuống 26,1% (-10,9%p), VietQR đạt 73,9%.
  16. AM Nguyễn Duy Long (Ninh Thuận): Tiền mặt chỉ 23,3% (-1,3%p), VietQR áp đảo tới 76,7%!
  17. AM Cao Thị Thanh Thủy (Bình Thuận): Tiền mặt chỉ 19,7% (+2,5%p), VietQR đạt 80,3%!
  18. AM Thái Thị Thanh Thư (Khánh Hòa): Quán quân toàn mạng GHN về thanh toán số, tiền mặt chỉ còn 4,4% (giảm -1,2%p), Chuyển khoản QR chiếm tới 95,6% trên hàng chục tỷ dòng tiền!

🔍 3. TOP 10 BƯU CỤC CÓ TỶ LỆ THU TIỀN MẶT CAO NHẤT VÙNG:
  1. (KHO) CK Diên Điền (AM Phan Đình Duy): 100,0% tiền mặt (+8,9%p, thu 650,8 Tr tiền mặt, 0đ QR).
  2. (DNO) Quảng Khê (AM Trần Thị Nhung): 99,6% tiền mặt (382,4 Tr).
  3. (DNO) Bắc Gia Nghĩa (AM Huỳnh Thúc Duân): 98,1% tiền mặt (429,4 Tr).
  4. (KHO) Cam Lâm 1 (AM Nguyễn Hoàng Phi): 91,7% tiền mặt (+7,6%p, 558,1 Tr).
  5. (BTH) Hàm Thuận (AM Lê Thanh Nhựt): 88,1% tiền mặt (1.166,7 Tr).
  6. (BTH) Đồng Kho (AM Lê Thanh Nhựt): 88,0% tiền mặt (1.311,2 Tr).
  7. (LDO) Ninh Gia (AM Nguyễn Thị Tuyết Thơ): 85,2% tiền mặt (+12,5%p, 723,2 Tr).
  8. (LDO) Hòa Ninh (AM Hồng Bích Nga): 82,9% tiền mặt (558,5 Tr).
  9. (LDO) Cát Tiên (AM Nguyễn Đỗ Minh Nghĩa): 81,8% tiền mặt (370,3 Tr).
  10. (LDO) Đam Rông 3 (AM Huỳnh Thị Kim Chi): 81,0% tiền mặt (703,1 Tr).

🎯 4. CHỈ ĐẠO HÀNH ĐỘNG & GIAO NHIỆM VỤ ĐÍCH DANH:
• Anh Trương Quang Linh: Yêu cầu giải trình gấp trong 24h vì sao toàn vùng giảm tiền mặt nhưng địa bàn của anh lại tăng vọt +6,4%p (từ 74% lên 81%). Kiểm tra ngay danh sách bưu tá thu tiền mặt tuyến Quảng Tín.
• Anh Lê Thanh Nhựt: Chấn chỉnh bưu cục Hàm Thuận và Đồng Kho — mỗi kho đang ôm trên 1,1 - 1,3 tỷ tiền mặt, tiềm ẩn rủi ro cướp giật, mất mát và chiếm dụng vốn. Bắt buộc trang bị decal VietQR trên toàn bộ phương tiện giao nhận.
• Biểu dương chị Thư (95,6% QR), anh Long Ninh Thuận (76,7% QR), chị Thủy (80,3% QR) và anh Lợi (giảm -10,9%p tiền mặt). Toàn vùng quyết tâm đưa tỷ lệ tiền mặt về dưới 35% trong tháng 10!"""

for fn in ['KICH_BAN_THUYET_TRINH_MOI_NHAT.docx', 'KICH_BAN_THUYET_TRINH_W39_INSIGHT_CHUYEN_SAU.docx']:
    try:
        doc = docx.Document(fn)
        if len(doc.tables) > 12:
            t12 = doc.tables[12]
            t12.rows[0].cells[0].text = table12_content
            doc.save(fn)
            print(f"Updated Table 12 in {fn} successfully.")
        else:
            print(f"File {fn} has less than 13 tables.")
    except Exception as e:
        print(f"Error updating {fn}: {e}")
