import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W39_2026.xlsx', data_only=True)
ws = wb['08_OPR TTS']

ams = []
for r in range(6, 24):
    am = ws.cell(r, 1).value
    if am and str(am).strip() not in ['None', '', 'AM']:
        v_day = ws.cell(r, 2).value or 0
        w38_day = ws.cell(r, 3).value or 0
        w39_day = ws.cell(r, 4).value or 0
        diff_day = ws.cell(r, 5).value or 0
        v_night = ws.cell(r, 6).value or 0
        w38_night = ws.cell(r, 7).value or 0
        w39_night = ws.cell(r, 8).value or 0
        diff_night = ws.cell(r, 9).value or 0
        v_tot = ws.cell(r, 10).value or 0
        w38_tot = ws.cell(r, 11).value or 0
        w39_tot = ws.cell(r, 12).value or 0
        diff_tot = ws.cell(r, 13).value or 0
        fail_tot = round(v_tot * (1 - w39_tot))
        ams.append({
            'am': str(am).strip(),
            'v_day': v_day, 'w38_day': w38_day, 'w39_day': w39_day, 'diff_day': diff_day,
            'v_night': v_night, 'w38_night': w38_night, 'w39_night': w39_night, 'diff_night': diff_night,
            'v_tot': v_tot, 'w38_tot': w38_tot, 'w39_tot': w39_tot, 'diff_tot': diff_tot,
            'fail_tot': fail_tot
        })

print('=== TOP AM ĐẠT KPI OPR W39 (>= 80%) ===')
pass_ams = [a for a in ams if a['w39_tot'] >= 0.8]
pass_ams.sort(key=lambda x: x['w39_tot'], reverse=True)
for a in pass_ams:
    print(f"{a['am']:25} | W39: {a['w39_tot']*100:5.1f}% | W38: {a['w38_tot']*100:5.1f}% | Diff: {a['diff_tot']*100:+5.1f}%p | Ngày: {a['w39_day']*100:5.1f}% | Đêm: {a['w39_night']*100:5.1f}%")

print('\n=== TOP AM CHƯA ĐẠT KPI OPR W39 (< 80%) ===')
fail_ams = [a for a in ams if a['w39_tot'] < 0.8]
fail_ams.sort(key=lambda x: x['fail_tot'], reverse=True)
tot_fail = sum(a['fail_tot'] for a in ams)
for a in fail_ams:
    pct_fail = (a['fail_tot'] / tot_fail * 100) if tot_fail > 0 else 0
    print(f"{a['am']:25} | W39: {a['w39_tot']*100:5.1f}% | Trễ: {a['fail_tot']:4d} đơn ({pct_fail:4.1f}%) | Ngày: {a['w39_day']*100:5.1f}% | Đêm: {a['w39_night']*100:5.1f}%")

# Overall region check
tot_v_day = sum(a['v_day'] for a in ams)
tot_v_night = sum(a['v_night'] for a in ams)
tot_v_all = sum(a['v_tot'] for a in ams)
pass_v_day_w39 = sum(a['v_day'] * a['w39_day'] for a in ams)
pass_v_night_w39 = sum(a['v_night'] * a['w39_night'] for a in ams)
pass_v_all_w39 = sum(a['v_tot'] * a['w39_tot'] for a in ams)

pass_v_day_w38 = sum(a['v_day'] * a['w38_day'] for a in ams)
pass_v_night_w38 = sum(a['v_night'] * a['w38_night'] for a in ams)
pass_v_all_w38 = sum(a['v_tot'] * a['w38_tot'] for a in ams)

print('\n=== TỔNG TOÀN VÙNG (TÍNH TỪ 17 AM CÓ SỐ LIỆU) ===')
print(f"Tổng Đơn: {tot_v_all:,} (Ngày: {tot_v_day:,} | Đêm: {tot_v_night:,})")
print(f"OPR W39: {pass_v_all_w39/tot_v_all*100:.2f}% (Ngày: {pass_v_day_w39/tot_v_day*100:.2f}% | Đêm: {pass_v_night_w39/tot_v_night*100:.2f}%)")
print(f"OPR W38: {pass_v_all_w38/tot_v_all*100:.2f}% (Ngày: {pass_v_day_w38/tot_v_day*100:.2f}% | Đêm: {pass_v_night_w38/tot_v_night*100:.2f}%)")
diff_vung = (pass_v_all_w39/tot_v_all - pass_v_all_w38/tot_v_all) * 100
print(f"Biến động WoW: {diff_vung:+.2f}%p")
