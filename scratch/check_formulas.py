# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'c:\Users\lap4all\Desktop\New folder\BaoCao_Tuan_NTB_W37_2026.xlsx', data_only=False)
ws = wb['05_ODR']
print("Formulas in 05_ODR:")
for r in [5, 6, 7, 11, 12, 13, 33, 34, 55, 56, 64, 65]:
    vals = [ws.cell(r, c).value for c in range(1, 8)]
    print(f"Row {r}: {vals}")

ws_tq = wb['00_Tong quan']
print("\nFormulas in 00_Tong quan:")
for r in [19, 30, 31]:
    vals = [ws_tq.cell(r, c).value for c in range(1, 8)]
    print(f"Row {r}: {vals}")
