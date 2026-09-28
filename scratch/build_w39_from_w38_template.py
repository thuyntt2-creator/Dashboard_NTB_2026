import sys
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import shutil

sys.stdout.reconfigure(encoding='utf-8')

src_template = r"C:\Users\lap4all\Downloads\KICH_BAN_THUYET_TRINH_W38_NAM_TRUNG_BO (3).docx"
if not os.path.exists(src_template):
    print(f"Error: {src_template} not found!")
    sys.exit(1)

doc = docx.Document(src_template)

# 1. Update Title Header paragraphs
for p in doc.paragraphs:
    if "TUẦN W38" in p.text:
        p.text = p.text.replace("TUẦN W38", "TUẦN W39")
    if "14/09/2026 – 20/09/2026" in p.text:
        p.text = p.text.replace("14/09/2026 – 20/09/2026", "21/09/2026 – 27/09/2026")
    if "(W38)" in p.text:
        p.text = p.text.replace("(W38)", "(W39)")

# 2. Update Table 0 (Section I: Tổng Hợp Trọng Tâm Tuần W39)
cell0 = doc.tables[0].rows[0].cells[0]
p0_first = cell0.paragraphs[0]
p0_first.text = "🗣️ LỜI MỞ ĐẦU & TỔNG QUAN ĐIỀU HÀNH VÙNG TUẦN W39:"

sec1_lines = [
    "📍 1. BẢNG CHỈ SỐ TOÀN VÙNG & 5 TỈNH THÀNH (W39 vs W38):",
    "• Khánh Hòa: 92.696 đơn (-1,0% WoW) | %GTC 58,6% (+3,6%p) | %ODR 91,8% | TTS: 18.043 đơn (+449 đơn WoW) ➔ Giữ vững vị trí số 1 sản lượng toàn vùng!",
    "• Lâm Đồng: 88.287 đơn (-7,8% WoW) | Dẫn đầu TTS vùng (19.283 đơn, chiếm 26,5%) | %GTC 47,4% | %ODR 84,7% | Tồn aging cao cần giải tỏa.",
    "• Bình Thuận: 80.817 đơn (-5,5% WoW) | %GTC 67,5% | %ODR 96,3% (Top 2 vùng) | TTS tăng bùng nổ: 18.146 đơn (+2.273 đơn WoW, +14,3%!).",
    "• Đắk Nông: 33.883 đơn (-2,8% WoW) | %GTC 48,8% (+2,0%p) | %ODR 88,5% | TTS: 9.016 đơn (+702 đơn WoW, +8,4%).",
    "• Ninh Thuận: 33.242 đơn (-2,0% WoW) | Quán quân %GTC toàn vùng (68,0%) | Quán quân %ODR toàn vùng (96,4%) | TTS: 8.293 đơn (+935 đơn WoW, +12,7%).",
    "➔ TOÀN VÙNG: 328.925 đơn Full (-4,3%) | TikTok Shop đạt đỉnh 72.781 đơn (+5,9% WoW, chiếm 22,1% tỷ trọng) | %GTC Full 56,60% (+0,92%p) | %GTC TTS 57,45% (+3,55%p) | %ODR 90,79% | Rớt luân chuyển giảm sâu về 1,52% (-1,81%p WoW) | Truy thu phát sinh 306,9 Tr ₫ trên 4.013 đơn.",
    "👤 2. BÓC TÁCH CHI TIẾT HIỆU SUẤT THEO QUẢN LÝ VẬN HÀNH (AM):",
    "Kính chào Ban Giám Đốc và toàn thể 18 anh chị Quản lý Vận hành (AM). Nhìn vào bức tranh tuần W39, sự phân hóa giữa các AM thể hiện rất rõ nét:",
    "• Nhóm AM dẫn đầu sản lượng và gánh tải: AM Nguyễn Duy Long tiếp tục là đầu tàu số 1 với 42.506 đơn Full (12,9% toàn vùng) và 10.584 đơn TTS (vô địch toàn mạng); AM Thái Thị Thanh Thư bứt phá ngoạn mục tăng +2.227 đơn Full (đạt 34.765 đơn, vươn lên Top 2 toàn vùng!).",
    "• Nhóm AM giữ chất lượng giao hàng (%GTC) xuất sắc trên 67%: AM Nguyễn Ngọc Khánh, AM Cao Thị Thanh Thủy, AM Nguyễn Đỗ Minh Nghĩa và AM Nguyễn Duy Long tiếp tục bảo vệ vững chắc chuẩn SLA xanh.",
    "• Điểm sáng vượt trội khâu vận tải: Tỷ lệ rớt luân chuyển toàn vùng giảm ngoạn mục từ 3,32% xuống 1,52% (-1,81%p WoW, giảm hơn 54% lượng đơn rớt) nhờ siết chặt kỷ luật bàn giao ca KTC.",
    "• Hai vấn đề nổi cộm cần can thiệp gấp: Số tiền truy thu phát sinh 306,9 triệu đồng trên 4.013 đơn do buông lỏng cân đo; và danh sách bưu cục cảnh báo tăng lên 15 điểm nóng cần điều tiết tải ngay trong tuần W40.",
    "🔍 INSIGHT & ĐIỂM NGHẼN THỰC TẾ TẠI CÁC AM:\n",
    "• TikTok Shop tăng tốc lên 72.781 đơn (+5,9% WoW, chiếm 22,1% toàn vùng), khẳng định các AM phải ưu tiên tuyệt đối nguồn lực lấy hàng và giao ca sáng cho phân khúc này.",
    "• %GTC TTS tăng vọt +3,55%p lên 57,45% chứng minh giải pháp thúc đẩy giao hàng TMĐT đã đi đúng hướng, nhưng Lâm Đồng và Đắk Nông vẫn dưới 50% do vướng đơn tồn aging.",
    "⚠️ CẢNH BÁO ĐỎ TẬP TRUNG THEO AM:\n",
    "• Khoản tiền truy thu phát sinh 306,9 triệu đồng (tập trung tại Kim Chi 58,3 Tr, Bích Nga 50,7 Tr, Văn Trường 50,6 Tr, Văn Phước 27,1 Tr) đe dọa trực tiếp dòng tiền.",
    "• AM Nguyễn Thanh Long (Cam Ranh) hụt -1.186 đơn TTS (-40% WoW), cần kiểm tra ngay hiện tượng mất shop live sang đối thủ.",
    "🎯 GIAO VIỆC & NHIỆM VỤ TÁC CHIẾN ĐÍCH DANH CHO TỪNG AM:\n",
    "• Toàn thể 18 AM quán triệt ngay 3 trọng tâm W40: Triển khai Cap Volume giải tỏa bưu cục nghẽn, kiểm soát 100% cân đo tại quầy và đẩy mạnh thu tiền COD qua VietQR."
]

# Clear cell0 except first paragraph
for p in cell0.paragraphs[1:]:
    p._p.getparent().remove(p._p)

for line in sec1_lines:
    p = cell0.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.text = line

# 3. Update Table 1 (Section II: Sản Lượng Giao - THE EXACT USER FAVORITE FORMAT!)
cell1 = doc.tables[1].rows[0].cells[0]
p1_first = cell1.paragraphs[0]
p1_first.text = "🗣️ ĐÁNH GIÁ SẢN LƯỢNG 5 TỈNH & XẾP HẠNG TĂNG TRƯỞNG 18 AM:"

sec2_lines = [
    "📍 1. BẢNG CHỈ SỐ TOÀN VÙNG & 5 TỈNH THÀNH (W39 vs W38):",
    "• Top 1 - Khánh Hòa: 92.696 đơn (giảm nhẹ -919 đơn, -1,0% WoW) | TTS: 18.043 đơn (+449 đơn WoW) ➔ Giữ vững vị trí số 1 toàn vùng!",
    "• Top 2 - Lâm Đồng: 88.287 đơn (giảm -7.440 đơn, -7,8% WoW) | TTS: 19.283 đơn (-304 đơn WoW, đầu tàu TTS vùng).",
    "• Top 3 - Bình Thuận: 80.817 đơn (giảm -4.672 đơn, -5,5% WoW) | TTS: 18.146 đơn (+2.273 đơn WoW, +14,3% ➔ Tăng TTS mạnh nhất vùng!).",
    "• Top 4 - Đắk Nông: 33.883 đơn (giảm -965 đơn, -2,8% WoW) | TTS: 9.016 đơn (+702 đơn WoW, +8,4%).",
    "• Top 5 - Ninh Thuận: 33.242 đơn (giảm -676 đơn, -2,0% WoW) | TTS: 8.293 đơn (+935 đơn WoW, +12,7%).",
    "➔ TỔNG TOÀN VÙNG: 328.925 đơn Full hàng (-4,3% WoW) | TikTok Shop: 72.781 đơn (+5,9% WoW, lập đỉnh 22,1% tỷ trọng).",
    "👤 2. BÓC TÁCH CHI TIẾT HIỆU SUẤT THEO QUẢN LÝ VẬN HÀNH (AM):",
    "Kính thưa Ban Giám Đốc, nhìn sâu vào bảng xếp hạng quy mô sản lượng của 18 AM tuần W39:",
    "• Top 5 AM có sản lượng giao lớn nhất vùng:",
    "  1. AM Nguyễn Duy Long: 42.506 đơn (quản lý địa bàn trọng điểm Bình Thuận, -512 đơn WoW | TTS: 10.584 đơn, tăng kỷ lục +1.379 đơn WoW, vô địch toàn mạng!).",
    "  2. AM Thái Thị Thanh Thư: 34.765 đơn (bứt phá ngoạn mục +2.227 đơn WoW, tương ứng +6,8% — AM tăng sản lượng Full tốt nhất vùng, vươn lên Top 2!).",
    "  3. AM Lê Thanh Nhựt: 29.070 đơn (giảm -1.530 đơn WoW | TTS: 7.108 đơn, tăng mạnh +800 đơn WoW).",
    "  4. AM Lê Văn Trường: 26.787 đơn (giảm -2.132 đơn WoW tại cụm Đà Lạt - Đơn Dương | TTS: 5.440 đơn, +131 đơn).",
    "  5. AM Nguyễn Ngọc Khánh: 26.242 đơn (giảm -2.083 đơn WoW tại Nha Trang | TTS: 5.155 đơn, +516 đơn).",
    "• Nhóm AM tăng trưởng dương tuần này:",
    "  - Về Full hàng: Duy nhất AM Thái Thị Thanh Thư (+2.227 đơn); cùng 2 AM giữ nhịp ổn định là AM Trương Quang Linh (+49 đơn, đạt 1.906 đơn) và AM Huỳnh Thúc Duân (+4 đơn, đạt 4.988 đơn).",
    "  - Về TikTok Shop: Bùng nổ 12/18 AM tăng trưởng, dẫn đầu là AM Nguyễn Duy Long (+1.379 đơn), AM Lê Thanh Nhựt (+800 đơn), AM Hồng Bích Nga (+610 đơn, đạt 4.900 đơn), AM Nguyễn Hoàng Phi (+593 đơn, đạt 5.390 đơn), AM Phan Đình Duy (+531 đơn, đạt 5.347 đơn), AM Nguyễn Ngọc Khánh (+516 đơn, đạt 5.155 đơn), AM Cao Thị Thanh Thủy (+513 đơn, đạt 3.592 đơn), AM Thái Thị Thanh Thư (+511 đơn, đạt 5.562 đơn), AM Trần Thị Nhung (+379 đơn, đạt 6.057 đơn).",
    "• Nhóm AM sụt giảm tải lớn cần lưu ý:",
    "  - Về Full hàng: AM Lê Văn Trường (-2.132 đơn), AM Nguyễn Ngọc Khánh (-2.083 đơn), AM Phan Đình Duy (-1.649 đơn), AM Lê Thanh Nhựt (-1.530 đơn).",
    "  - Về TikTok Shop: Báo động đỏ AM Nguyễn Thanh Long tại Cam Ranh rơi tự do -1.186 đơn TTS (-40,5% WoW, từ 2.930 xuống 1.744 đơn) và AM Lê Minh Lợi giảm -637 đơn TTS (từ 844 về 207 đơn).",
    "• Nhóm AM địa bàn nhỏ: AM Lê Minh Lợi (2.539 đơn Full), AM Trương Quang Linh (1.906 đơn Full).",
    "🔍 INSIGHT & ĐIỂM NGHẼN THỰC TẾ TẠI CÁC AM:\n",
    "• AM Thái Thị Thanh Thư lội ngược dòng tăng +2.227 đơn nhờ tóm trọn sức mua tiêu dùng nội thị Nha Trang, trong khi các AM khác hạ nhiệt theo nhịp thị trường.",
    "• Kênh TikTok Shop tăng vọt lên 72.781 đơn (+5,9% WoW) với 12/18 AM tăng trưởng, trở thành phao cứu sinh bù đắp sản lượng cho toàn vùng.",
    "⚠️ CẢNH BÁO ĐỎ TẬP TRUNG THEO AM:\n",
    "• AM Nguyễn Thanh Long (Cam Ranh) mất gần một nửa sản lượng TTS (-1.186 đơn), cần kiểm tra ngay hiện tượng các shop livestream lớn chuyển sang đối thủ.",
    "• 4 AM đầu tàu (Trường, Khánh, Duy, Nhựt) giảm hơn 7.300 đơn Full hàng, ảnh hưởng đến hiệu quả lấp đầy xe tải KTC.",
    "🎯 GIAO VIỆC & NHIỆM VỤ TÁC CHIẾN ĐÍCH DANH CHO TỪNG AM:\n",
    "• Khối KD phối hợp cùng AM Nguyễn Thanh Long làm việc trực tiếp với các shop thời trang Cam Ranh để kéo lại sản lượng livestream.",
    "• Các AM có tỷ trọng TTS trên 24% (Duân, Phi, Long, Nhung, Nhựt) ưu tiên hoàn tất chia chọn trước 08h00 để đảm bảo SLA ca sáng."
]

for p in cell1.paragraphs[1:]:
    p._p.getparent().remove(p._p)

for line in sec2_lines:
    p = cell1.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.text = line

# 4. Update Table 15 (Section XVI: 15 Bưu cục cảnh báo bất ổn W38 vs W39)
cell15 = doc.tables[15].rows[0].cells[0]
p15_first = cell15.paragraphs[0]
p15_first.text = "🗣️ DANH SÁCH 15 BƯU CỤC CẢNH BÁO BẤT ỔN (W38 vs W39) & 3 MỆNH LỆNH TÁC CHIẾN:"

sec16_lines = [
    "📍 1. BIẾN ĐỘNG SỐ LƯỢNG TOÀN VÙNG (TĂNG TỪ 13 LÊN 15 BƯU CỤC):",
    "• Toàn vùng ghi nhận 15 bưu cục rơi vào diện cảnh báo bất ổn (%GTC < 45% hoặc giảm sâu dưới 70% mốc đỉnh lịch sử), tăng ròng +2 bưu cục so với 13 bưu cục ở tuần W38.",
    "• Cơ cấu dịch chuyển: 4 bưu cục đã thoát cảnh báo thành công, 6 bưu cục mới xuất hiện, và 9 bưu cục mãn tính kẹt cả 2 tuần liên tiếp.",
    "👤 2. BÓC TÁCH CHI TIẾT 3 NHÓM BƯU CỤC CẢNH BÁO THEO AM PHỤ TRÁCH:",
    "Kính thưa Ban Giám Đốc, nhìn sâu vào danh sách bưu cục bất ổn tuần W39:",
    "• Nhóm 1: Tuyên dương 4 Bưu cục đã giải tỏa & thoát cảnh báo thành công:",
    "  1. BC Xuân Hương - Đà Lạt (AM Lê Văn Trường): Nâng %GTC lên >50%, giải tỏa 100% backlog.",
    "  2. BC Cam Linh (AM Phan Đình Duy): Phục hồi nhân sự bưu tá và giờ xuất tuyến.",
    "  3. BC Nhân Cơ (AM Đỗ Duy Khang): Kiểm soát tốt luồng hàng nông thôn.",
    "  4. BC Tây Nha Trang (AM Hồng Bích Nga): Tối ưu ca phát sáng trước 11h00.",
    "• Nhóm 2: Báo động đỏ 6 Bưu cục mới rơi vào diện cảnh báo W39:",
    "  1. BC D'Ran (AM Phan Thành Long): %GTC đạt 36,86% (cảnh báo 6 ngày), địa hình đèo dốc sạt lở.",
    "  2. BC Lang Biang 2 (AM Trương Tuấn Anh): %GTC giảm về 35,80% (cảnh báo 6 ngày), TTS dồn ứ cục bộ.",
    "  3. BC Trường Xuân (AM Đỗ Duy Khang): %GTC tụt về 38,22% (cảnh báo 8 ngày), tỷ lệ KLL cao.",
    "  4. BC Đông Gia Nghĩa (AM Đỗ Duy Khang): %GTC đạt 44,89% (chớm dưới ngưỡng 45%), gán ca 2 yếu.",
    "  5. BC Bắc Cam Ranh (AM Phan Đình Duy): %GTC đạt 44,68% (đã tích lũy 31 ngày cảnh báo).",
    "  6. BC Phú Quý (AM Nguyễn Đình Uy): %GTC đạt 50,84%, tuy nhiên rơi sâu dưới 70% so với mốc lịch sử 88,55%.",
    "• Nhóm 3: Điểm nóng mãn tính 9 Bưu cục kẹt dai dẳng cả 2 tuần W38 và W39:",
    "  - Báo động nặng nhất mạng: BC Lâm Viên - Đà Lạt 2 (AM Trương Tuấn Anh) rơi tự do từ 46,2% xuống 19,55% (-26,6%p, bưu tá tê liệt ca chiều).",
    "  - BC Đức Trọng 1 (AM Phan Thành Long): Kẹt ở 23,78% (backlog tồn đọng >300 đơn).",
    "  - BC Quảng Tín (AM Đỗ Duy Khang): Đạt 26,17% (tích lũy kỷ lục >100 ngày cảnh báo).",
    "  - BC Đơn Dương (AM Phan Thành Long): Giảm mạnh về 29,01% (giảm -16,3%p WoW).",
    "  - Cùng 5 BC: Lang Biang 1 (30,6%), Kiến Đức (32,5%), Tân Hà Lâm Hà (36,8%), Tuy Đức (38,5%), Di Linh (40,6%).",
    "🔍 INSIGHT & ĐIỂM NGHẼN THỰC TẾ TẠI CÁC AM:\n",
    "• Tồn đọng và rớt %GTC ở nhóm 9 bưu cục mãn tính chủ yếu do thắt nút cổ chai ở tỷ lệ gán Ca 2 trưa và địa bàn đồi núi xa xôi.",
    "• Việc Lâm Viên 2 tụt xuống 19,55% kéo lùi toàn bộ chỉ số giao hàng của cụm thành phố Đà Lạt.",
    "⚠️ CẢNH BÁO ĐỎ TẬP TRUNG THEO AM:\n",
    "• Nguy cơ vỡ tải dây chuyền tại cụm Đà Lạt - Đơn Dương - Đức Trọng trong các đợt sale đầu tháng 10 nếu không điều tiết tải ngay.",
    "🎯 GIAO VIỆC & NHIỆM VỤ TÁC CHIẾN ĐÍCH DANH CHO TỪNG AM:\n",
    "• 1. Áp dụng cơ chế Cap Volume: Giảm 25% hạn mức chia chọn về Lâm Viên 2 và Quảng Tín trong 3 ngày đầu tuần để bưu tá dọn sạch tồn kho.",
    "• 2. Chi viện bưu tá ca sáng: Điều chuyển 4 bưu tá cứng hỗ trợ phát lượt 1 trước 11h00 tại Lâm Viên 2 và Đức Trọng 1.",
    "• 3. Cam kết W40: Quyết tâm kéo tối thiểu 5 bưu cục thoát cảnh báo, đưa danh sách toàn vùng về dưới 10 bưu cục."
]

for p in cell15.paragraphs[1:]:
    p._p.getparent().remove(p._p)

for line in sec16_lines:
    p = cell15.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.text = line

# Save to destination master
out_master = 'KICH_BAN_THUYET_TRINH_MOI_NHAT.docx'
doc.save(out_master)
print(f"Generated {out_master} successfully! (Size: {os.path.getsize(out_master)} bytes)")

copies = [
    'KICH_BAN_THUYET_TRINH_W39_INSIGHT_CHUYEN_SAU.docx',
    'KICH_BAN_THUYET_TRINH_W39_INSIGHT_CHUYEN_SAU_CHINH_SUA.docx',
    'KICH_BAN_THUYET_TRINH_W39_NAM_TRUNG_BO.docx'
]
for c in copies:
    shutil.copyfile(out_master, c)
    print(f"Copied to {c}")
