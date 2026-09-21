import json
import re
import os
import openpyxl

print("=== Checking data.js ===")
if os.path.exists("data.js"):
    with open("data.js", "r", encoding="utf-8") as f:
        content = f.read()
    # find keys or sections
    slides = re.findall(r'id:\s*["\']([^"\']+)["\'],\s*title:\s*["\']([^"\']+)["\']', content)
    print("Found slides in data.js:", len(slides))
    for i, s in enumerate(slides):
        print(f"  Slide {i+1}: {s[0]} - {s[1]}")

print("\n=== Checking BaoCao_Tuan_NTB_W38_2026.xlsx ===")
if os.path.exists("BaoCao_Tuan_NTB_W38_2026.xlsx"):
    wb = openpyxl.load_workbook("BaoCao_Tuan_NTB_W38_2026.xlsx", read_only=True)
    print("Sheets in W38 excel:", len(wb.sheetnames))
    for i, s in enumerate(wb.sheetnames):
        print(f"  Sheet {i+1}: {s}")

print("\n=== Checking templates/ or html files ===")
for fn in ["index.html", "templates/index.html", "templates/meeting.html"]:
    if os.path.exists(fn):
        with open(fn, "r", encoding="utf-8") as f:
            html = f.read()
        nav_items = re.findall(r'<li[^>]*data-slide=["\']([^"\']+)["\'][^>]*>(.*?)</li>', html)
        if nav_items:
            print(f"{fn} nav items count:", len(nav_items))
            for i, item in enumerate(nav_items):
                print(f"  {i+1}: {item[0]} -> {re.sub('<[^<]+?>', '', item[1]).strip()}")
