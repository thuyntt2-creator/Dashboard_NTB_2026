import docx

doc = docx.Document('KICH_BAN_THUYET_TRINH_W37_INSIGHT_CHUYEN_SAU.docx')
for i, p in enumerate(doc.paragraphs):
    print(f"P{i+1}: style='{p.style.name}' | text='{p.text}'")

print("\n--- TABLE BORDERS / SHADING ---")
for t_idx, tbl in enumerate(doc.tables):
    c = tbl.cell(0, 0)
    tcPr = c._tc.get_or_add_tcPr()
    xml = tcPr.xml
    print(f"Table {t_idx+1} xml snippet: {xml[:300]}")
