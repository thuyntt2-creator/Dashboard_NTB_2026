import openpyxl, sys, json
sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W36_2026.xlsx', data_only=True)
ws = wb['09_Rot LC']

ams_data = []
for r in range(11, 28):
    am = ws.cell(r, 1).value
    vol = ws.cell(r, 2).value or 0
    w35 = ws.cell(r, 3).value or 0
    w36 = ws.cell(r, 4).value or 0
    vol_rot = round(vol * w36)
    ams_data.append({'am': am, 'vol': vol, 'w35': w35, 'w36': w36, 'vol_rot': vol_rot})

tot_rot = sum(x['vol_rot'] for x in ams_data)
tot_vol = sum(x['vol'] for x in ams_data)
print(f'Total Vol cần LC: {tot_vol}, Total Vol rớt: {tot_rot}')

for x in sorted(ams_data, key=lambda a: a['vol_rot'], reverse=True):
    rate_rot = (x['vol_rot'] / tot_rot * 100) if tot_rot > 0 else 0
    print(f"{x['am']:22} | Vol cần: {x['vol']:4} | % Rớt: {x['w36']*100:5.1f}% | Đơn rớt: {x['vol_rot']:3} | Tỷ trọng rớt: {rate_rot:5.1f}%")
