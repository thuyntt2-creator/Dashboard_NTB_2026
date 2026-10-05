import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

file_path = r"C:\Users\lap4all\Downloads\KICH_BAN_THUYET_TRINH_W39_NAM_TRUNG_BO (16).docx"

try:
    doc = docx.Document(file_path)
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    print(f"Total tables: {len(doc.tables)}")
    
    print("\n--- FIRST 40 PARAGRAPHS ---")
    for i, p in enumerate(doc.paragraphs[:40]):
        if p.text.strip():
            print(f"P{i:02d}: {p.text}")
            
    print("\n--- SAMPLE HEADINGS IN W39 ---")
    for i, p in enumerate(doc.paragraphs):
        if p.text.startswith("#") or p.text.startswith("Chuyên đề") or p.text.startswith("PHẦN") or "TAB" in p.text.upper() or p.text.startswith("📍") or p.text.startswith("🗣️") or p.text.startswith("🎙️"):
            print(f"P{i:03d}: {p.text[:90]}")
except Exception as e:
    print("Error reading docx:", e)
