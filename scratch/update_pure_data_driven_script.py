import sys
import os
import docx
from docx.shared import Inches, Pt, RGBColor
import shutil

sys.stdout.reconfigure(encoding='utf-8')

src_template = r"C:\Users\lap4all\Downloads\KICH_BAN_THUYET_TRINH_W38_NAM_TRUNG_BO (3).docx"
if not os.path.exists(src_template):
    print(f"Error: {src_template} not found!")
    sys.exit(1)

doc = docx.Document(src_template)

# Update Title Header paragraphs
for p in doc.paragraphs:
    if "TUẦN W38" in p.text:
        p.text = p.text.replace("TUẦN W38", "TUẦN W39")
    if "14/09/2026 – 20/09/2026" in p.text:
        p.text = p.text.replace("14/09/2026 – 20/09/2026", "21/09/2026 – 27/09/2026")
    if "(W38)" in p.text:
        p.text = p.text.replace("(W38)", "(W39)")

# ==========================================
# 1. TABLE 0: MỤC I - TỔNG QUAN ĐIỀU HÀNH VÙNG W39
# ==========================================
cell0 = doc.tables[0].rows[0].cells[0]
cell0.paragraphs[0].text = "🗣️ LỜI MỞ ĐẦU & TỔNG QUAN ĐIỀU HÀNH VÙNG TUẦN W39:"

sec1_lines = [
    "📍 1. BẢNG CHỈ SỐ TOÀN VÙNG & 5 TỈNH THÀNH (W39 vs W38):",
    "• Khánh Hòa: 92.696 đơn (-1,0% WoW) | %GTC 58,6% (+3,6%p) | %ODR 91,8% | TTS: 18.043 đơn (+449 đơn WoW) ➔ Giữ vững vị trí số 1 sản lượng toàn vùng!",
    "• Lâm Đồng: 88.287 đơn (-7,8% WoW) | Dẫn đầu TTS vùng (19.283 đơn, chiếm 26,5%) | %GTC 47,4% | %ODR 84,7% | Bưu cục cảnh báo tồn đọng cao.",
    "• Bình Thuận: 80.817 đơn (-5,5% WoW) | %GTC 67,5% | %ODR 96,3% (Top 2 vùng) | TTS bùng nổ: 18.146 đơn (+2.273 đơn WoW, +14,3%!).",
    "• Đắk Nông: 33.883 đơn (-2,8% WoW) | %GTC 48,8% (+2,0%p) | %ODR 88,5% | TTS: 9.016 đơn (+702 đơn WoW, +8,4%).",
    "• Ninh Thuận: 33.242 đơn (-2,0% WoW) | Quán quân %GTC toàn vùng (68,0%) | Quán quân %ODR toàn vùng (96,4%) | TTS: 8.293 đơn (+935 đơn WoW, +12,7%).",
    "➔ TOÀN VÙNG: 328.925 đơn Full (-4,3%) | TikTok Shop đạt đỉnh 72.781 đơn (+5,9% WoW, chiếm 22,1% tỷ trọng) | %GTC Full 56,60% (+0,92%p) | %GTC TTS 57,45% (+3,55%p) | %ODR 90,79% | Rớt luân chuyển giảm sâu về 1,52% (-1,81%p WoW) | Truy thu phát sinh 306,9 Tr ₫ trên 4.013 đơn.",
    "",
    "🎙️ LỜI THOẠI THUYẾT TRÌNH MỤC TỔNG QUAN (DÀNH CHO NGƯỜI THUYẾT TRÌNH):",
    "\"Kính thưa Ban Giám Đốc và toàn thể 18 anh chị Quản lý Vận hành (AM),",
    "Nhìn vào bảng chỉ số tổng quan tuần W39, toàn vùng Nam Trung Bộ đạt 328.925 đơn Full hàng, giảm 4,3% so với tuần W38. Tuy nhiên, đi sâu vào cấu trúc số liệu, chúng ta ghi nhận 3 chuyển động vận hành rất đáng chú ý:",
    "• Điểm sáng thứ nhất là dòng hàng TikTok Shop tiếp tục tăng trưởng ngược chiều thị trường, đạt 72.781 đơn (+5,9% WoW) và chính thức lập đỉnh tỷ trọng mới là 22,1% tổng sản lượng toàn vùng. Đi đôi với sản lượng, chất lượng giao hàng TikTok Shop (%GTC) tuần này có bước tiến nhảy vọt +3,55%p, đạt 57,45% – lần đầu tiên vượt qua tỷ lệ GTC của hàng Full (56,60%).",
    "• Điểm sáng thứ hai ở khâu vận chuyển trục: Tỷ lệ rớt luân chuyển KTC giảm ngoạn mục từ 3,32% xuống 1,52% (giảm hơn một nửa, tương ứng -1,81%p WoW). Điều này phản ánh rõ nét việc siết kỷ luật bàn giao ca KTC đã phát huy hiệu quả ngay trong tuần.",
    "• Tuy nhiên, bức tranh vận hành tuần này xuất hiện 2 điểm nghẽn lớn: Thứ nhất, danh sách bưu cục bất ổn tăng từ 13 lên 15 bưu cục, tập trung ở các điểm nghẽn giao hàng vùng núi Lâm Đồng, Đắk Nông; Thứ hai, số liệu phạt truy thu phát sinh tăng vọt lên 306,9 triệu đồng trên 4.013 đơn, trong đó hơn 162 triệu đồng đến từ lỗi backlog giao hàng và tồn đọng luân chuyển trả.",
    "Do đó, trọng tâm điều hành tuần W40 là giải tỏa triệt để các bưu cục nghẽn tải, siết chặt kỷ luật giao ca chiều và xử lý dứt điểm các khoản truy thu phát sinh.\"",
    "",
    "🔍 INSIGHT BẢN CHẤT TỪ SỐ LIỆU ĐIỀU HÀNH:",
    "• Việc tỷ lệ GTC toàn vùng tăng (+0,92%p Full và +3,55%p TTS) giữa lúc sản lượng giảm nhẹ (-4,3%) cho thấy các bưu cục đã tận dụng tốt khoảng giãn áp lực đơn để dọn dẹp hàng phát sinh trong ngày, cộng hưởng với việc luồng hàng luân chuyển trơn tru (rớt LC chỉ còn 1,52%).",
    "• Lâm Đồng và Đắk Nông vẫn là hai vùng trũng về chỉ số chất lượng khi GTC đều nằm dưới mốc 50% (Lâm Đồng 47,4%, Đắk Nông 48,8%), trực tiếp kéo lùi trung bình toàn vùng.",
    "",
    "⚠️ CẢNH BÁO ĐỎ TẬP TRUNG THEO AM:",
    "• 15 bưu cục cảnh báo bất ổn (đặc biệt BC Lâm Viên 2 tụt xuống 19,55% và Đức Trọng 1 kẹt ở 23,78%) có nguy cơ vỡ tải cục bộ kéo dài nếu không được điều tiết hạn mức tiếp nhận.",
    "• 306,9 triệu đồng phát sinh truy thu (nằm đậm nhất tại AM Kim Chi 58,3 Tr, AM Bích Nga 50,7 Tr, AM Văn Trường 50,6 Tr, AM Văn Phước 27,1 Tr) đe dọa trực tiếp đến chỉ số tài chính và trách nhiệm vận hành.",
    "",
    "🎯 GIAO VIỆC & NHIỆM VỤ TÁC CHIẾN ĐÍCH DANH:",
    "• Toàn thể 18 AM thực hiện nghiêm 3 nhiệm vụ tuần W40: Triển khai hạn mức tiếp nhận tải tại bưu cục nghẽn, kiểm soát chặt luân chuyển hàng trả để chặn phát sinh phạt truy thu, và duy trì tỷ lệ rớt luân chuyển KTC dưới 2,0%."
]

for p in cell0.paragraphs[1:]:
    p._p.getparent().remove(p._p)

for line in sec1_lines:
    p = cell0.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.text = line

# ==========================================
# 2. TABLE 1: MỤC II - SẢN LƯỢNG GIAO (3 CHARTS TAB II)
# ==========================================
cell1 = doc.tables[1].rows[0].cells[0]
cell1.paragraphs[0].text = "🗣️ KỊCH BẢN THUYẾT TRÌNH SẢN LƯỢNG GIAO (TAB II — 3 CHARTS):"

sec2_lines = [
    "📍 1. TỔNG HỢP SỐ LIỆU SẢN LƯỢNG 5 TỈNH THÀNH (W39 vs W38):",
    "• Top 1 - Khánh Hòa: 92.696 đơn (giảm -919 đơn, -1,0% WoW) | TTS: 18.043 đơn (+449 đơn WoW) ➔ Giữ vững ngôi đầu toàn vùng.",
    "• Top 2 - Lâm Đồng: 88.287 đơn (giảm sâu -7.440 đơn, -7,8% WoW) | TTS: 19.283 đơn (-304 đơn WoW, đầu tàu TTS vùng).",
    "• Top 3 - Bình Thuận: 80.817 đơn (giảm -4.672 đơn, -5,5% WoW) | TTS: 18.146 đơn (+2.273 đơn WoW, +14,3% ➔ Tăng trưởng TTS mạnh nhất vùng).",
    "• Top 4 - Đắk Nông: 33.883 đơn (giảm -965 đơn, -2,8% WoW) | TTS: 9.016 đơn (+702 đơn WoW, +8,4%).",
    "• Top 5 - Ninh Thuận: 33.242 đơn (giảm -676 đơn, -2,0% WoW) | TTS: 8.293 đơn (+935 đơn WoW, +12,7%).",
    "➔ TOÀN VÙNG: 328.925 đơn Full hàng (-4,3% WoW) | TikTok Shop: 72.781 đơn (+5,9% WoW, lập đỉnh 22,1% tỷ trọng).",
    "",
    "🎙️ LỜI THOẠI THUYẾT TRÌNH SẢN LƯỢNG (DÀNH CHO NGƯỜI THUYẾT TRÌNH):",
    "\"Kính thưa Ban Giám Đốc và các anh chị AM,",
    "Quan sát 3 biểu đồ sản lượng trên màn hình, tổng đơn toàn vùng tuần W39 đạt 328.925 đơn, giảm 4,3% so với tuần trước. Khi đối chiếu tương quan giữa hàng Full và hàng TikTok Shop trên 3 biểu đồ này, số liệu chỉ ra 2 xu hướng và 4 trường hợp phân hóa điển hình giữa các AM:",
    "",
    "Thứ nhất về xu hướng chung: Sự suy giảm sản lượng tuần này thuần túy rơi vào phân khúc hàng truyền thống B2C (-18.727 đơn), trong khi dòng hàng TikTok Shop lại tăng trưởng dương ở 4/5 tỉnh thành (+4.055 đơn, +5,9% WoW), kéo tỷ trọng TikTok Shop trong cơ cấu toàn mạng từ 20,0% lên mức đỉnh lịch sử 22,1%. Bình Thuận, Ninh Thuận và Đắk Nông là những địa bàn có mức tăng trưởng TikTok Shop trên 8% đến 14%.",
    "",
    "Thứ hai, bóc tách theo từng Quản lý Vận hành (AM) trên biểu đồ tăng trưởng WoW:",
    "",
    "🏆 HAI AM DẪN ĐẦU TĂNG TRƯỞNG & GÁNH TẢI (HIGHLIGHT TỐT NHẤT):",
    "• Thứ nhất là AM Thái Thị Thanh Thư (Khánh Hòa): Quán quân tăng trưởng toàn mạng với +2.227 đơn Full (+6,8% WoW, đạt 34.765 đơn). Nhìn trên đồ thị 18 AM, chị Thư là AM DUY NHẤT trong toàn bộ 18 AM đạt tăng trưởng dương về sản lượng Full trong tuần này (17 AM còn lại đều giảm hoặc đi ngang). Nhờ mức tăng bứt phá này, AM Thư đã vượt qua AM Trường và AM Nhựt để chính thức vươn lên vị trí Top 2 sản lượng toàn vùng, giúp Khánh Hòa bảo vệ vững chắc ngôi vị số 1 của tỉnh.",
    "• Thứ hai là AM Nguyễn Duy Long (Bình Thuận): Đầu tàu gánh tải lớn nhất vùng với 42.506 đơn Full (chiếm 12,9% sản lượng toàn mạng) và dẫn đầu phân khúc TikTok Shop với 10.584 đơn TTS (tăng +1.379 đơn WoW, +15,0%). AM Long là trường hợp đầu tiên của toàn vùng vượt mốc 10.000 đơn TikTok Shop chỉ trong 1 tuần, đóng góp chính vào mức tăng 2.273 đơn TTS của tỉnh Bình Thuận.",
    "",
    "🚨 HAI AM SỤT GIẢM NẶNG NỀ CẦN LƯU TÂM (HIGHLIGHT TỆ NHẤT):",
    "• Thứ nhất là AM Nguyễn Thanh Long (Cam Ranh): Cú sụt giảm TikTok Shop sâu nhất toàn mạng (-1.186 đơn TTS, -40,5% WoW, rơi từ 2.930 đơn xuống chỉ còn 1.744 đơn). Trong bối cảnh sản lượng TikTok Shop toàn vùng tăng +5,9% và các AM khác tại Khánh Hòa đều tăng dương TTS (+449 đơn), sự sụt giảm đột biến hơn 40% tại địa bàn Cam Ranh là điểm bất thường lớn nhất trên biểu đồ TTS, cần rà soát lại ngay nguồn phát sinh đơn của địa bàn.",
    "• Thứ hai là AM Lê Văn Trường (Lâm Đồng): Đơn vị có mức sụt giảm Full hàng sâu nhất toàn vùng (-2.132 đơn Full, rơi từ 28.919 đơn về 26.787 đơn). Mức giảm này đóng góp trực tiếp vào mức sụt -7.440 đơn của tỉnh Lâm Đồng. Đáng chú ý, trên các bảng chỉ số vận hành khác, địa bàn của AM Trường cũng đang ghi nhận điểm nghẽn tại bưu cục Lâm Viên 2 (GTC rơi về 19,55%) và phát sinh 50,6 triệu đồng tiền phạt truy thu liên quan đến 722 đơn tồn đọng, cho thấy sự ách tắc giao hàng đang đi liền với sự sụt giảm sản lượng.",
    "",
    "⚠️ ĐIỂM LƯU Ý VỀ TỶ TRỌNG ĐƠN TIKTOK SHOP:",
    "• Trên biểu đồ cơ cấu, hai địa bàn miền núi của Đắk Nông ghi nhận tỷ trọng TikTok Shop ở mức rất cao: AM Trương Quang Linh có tỷ trọng TTS lên tới 34,8% (664 đơn TTS trên 1.906 đơn tổng) và AM Huỳnh Thúc Duân đạt 29,0% (1.449 đơn TTS trên 4.988 đơn tổng). Với tỷ trọng đơn sàn cao trên địa bàn giao nhận xa, việc kiểm soát thời gian xuất tuyến ca sáng có ý nghĩa quyết định đến tỷ lệ giao thành công của 2 AM này.\"",
    "",
    "🔍 INSIGHT BẢN CHẤT RÚT RA TỪ SỐ LIỆU SẢN LƯỢNG:",
    "• Cơ cấu đơn chuyển dịch: Tăng trưởng TikTok Shop (+5,9%) bù đắp một phần cho sự suy giảm của đơn hàng B2C, đưa TTS trở thành động lực giữ nhịp sản lượng cho các tỉnh Bình Thuận, Ninh Thuận và Đắk Nông.",
    "• Tương quan nghịch giữa tồn đọng và sản lượng: Tỉnh Lâm Đồng giảm sâu nhất vùng (-7.440 đơn Full, -7,8%) trùng khớp với việc tỉnh này có số bưu cục cảnh báo cao và tỷ lệ GTC thấp nhất vùng (47,4%).",
    "",
    "🎯 NHIỆM VỤ TÁC CHIẾN GIAO CHO CÁC AM TUẦN W40:",
    "• 1. AM Nguyễn Thanh Long: Rà soát nguyên nhân giảm 1.186 đơn TTS tại Cam Ranh, làm việc với các đầu mối gửi hàng lớn trên địa bàn để khôi phục mức sản lượng trên 2.500 đơn TTS/tuần.",
    "• 2. AM Lê Văn Trường: Tập trung xử lý dứt điểm 722 đơn tồn đọng tại cụm Đà Lạt để thông luồng hàng, đưa sản lượng Full tuần W40 quay trở lại mốc 28.000 đơn.",
    "• 3. 18 AM: Tiếp tục ưu tiên tiến độ xử lý và ca lấy hàng cho luồng đơn TikTok Shop để duy trì đà tăng trưởng trên 22% tỷ trọng toàn mạng."
]

for p in cell1.paragraphs[1:]:
    p._p.getparent().remove(p._p)

for line in sec2_lines:
    p = cell1.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.text = line

# Save to destination master
out_master = 'KICH_BAN_THUYET_TRINH_MOI_NHAT.docx'
doc.save(out_master)
print(f"Generated {out_master} successfully! (Size: {os.path.getsize(out_master)} bytes)")

copies = [
    'KICH_BAN_THUYET_TRINH_W39_INSIGHT_CHUYEN_SAU.docx',
    'KICH_BAN_THUYET_TRINH_W39_INSIGHT_CHUYEN_SAU_CHINH_SUA.docx',
    'KICH_BAN_THUYET_TRINH_W39_NAM_TRUNG_BO.docx'
]
for c in copies:
    shutil.copyfile(out_master, c)
    print(f"Copied to {c}")
