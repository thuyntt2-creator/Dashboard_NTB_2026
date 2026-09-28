import sys
import docx
from docx.shared import Pt, RGBColor
import shutil

sys.stdout.reconfigure(encoding='utf-8')

def update_san_luong_highlight(doc_path):
    doc = docx.Document(doc_path)
    if len(doc.tables) < 2:
        return
    
    cell = doc.tables[1].rows[0].cells[0]
    
    # Clear existing paragraphs
    p0 = cell.paragraphs[0]
    p0.text = ""
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
        
    p0.paragraph_format.space_before = Pt(2)
    p0.paragraph_format.space_after = Pt(4)
    r_t = p0.add_run("🗣️ SẢN LƯỢNG GIAO W39: ĐIỂM SÁNG BỨT PHÁ & BÁO ĐỘNG ĐỎ CÁC AM NGUY HIỂM:")
    r_t.bold = True
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(11)
    r_t.font.color.rgb = RGBColor(194, 65, 12)
    
    lines = [
        ('"Kính thưa Ban Giám Đốc, về sản lượng giao tuần W39 (Full: 328.925 đơn, -4,3% | TTS: 72.781 đơn, +5,9%), em xin không liệt kê dàn trải mà đi thẳng vào 2 AM làm tốt nhất và 2 AM báo động đỏ tệ nhất cần xử lý gấp:', 'intro'),
        
        ('🏆 1. HAI ĐIỂM SÁNG BỨT PHÁ XUẤT SẮC NHẤT MẠNG LƯỚI:', 'hdr_good'),
        ('• 1. AM Thái Thị Thanh Thư (Khánh Hòa) — Quán quân tăng trưởng Full hàng: Là điểm sáng lội ngược dòng duy nhất toàn mạng khi tăng tới +2.227 đơn (+6,8% WoW, đạt 34.765 đơn), chính thức vượt lên vị trí Top 2 sản lượng toàn vùng!', 'bullet_good'),
        ('  ➔ Insight: Tóm trọn sức mua tiêu dùng và thời trang bùng nổ tại nội thị Nha Trang nhờ bố trí ca lấy hàng linh hoạt và chăm sóc rất sát các shop lớn.', 'sub_insight'),
        ('• 2. AM Nguyễn Duy Long (Bình Thuận) — Vô địch gánh tải & Tăng trưởng TikTok Shop kỷ lục: Tiếp tục là đầu tàu số 1 toàn mạng gánh 42.506 đơn Full (13% toàn vùng); riêng TikTok Shop tăng vọt +1.379 đơn (+15,0% WoW, đạt 10.584 đơn) — trở thành AM đầu tiên của NTB vượt mốc 10k đơn TTS/tuần!', 'bullet_good'),
        ('  ➔ Insight: Khai thác cực tốt làn sóng livestream đặc sản nông sản Bình Thuận kết hợp giải tỏa First-mile buổi chiều thần tốc.', 'sub_insight'),
        
        ('🚨 2. HAI ĐIỂM NÓNG BÁO ĐỘNG ĐỎ SỤT GIẢM NGUY HIỂM NHẤT:', 'hdr_bad'),
        ('• 1. AM Nguyễn Thanh Long (Cam Ranh) — Cú sốc rơi tự do TikTok Shop: Đang từ 2.930 đơn TTS rơi thẳng xuống 1.744 đơn, mất tới -1.186 đơn TTS (-40,5% WoW — bốc hơi gần một nửa sản lượng TTS)!', 'bullet_bad'),
        ('  ➔ Insight & Cảnh báo: Nguy cơ cao các shop livestream thời trang lớn tại Cam Ranh đã chuyển sang đơn vị vận chuyển đối thủ hoặc tạm dừng live do đứt gãy hàng. Khối KD phải nhảy vào giữ shop ngay lập tức!', 'sub_warning'),
        ('• 2. AM Lê Văn Trường (Lâm Đồng) — Suy giảm sản lượng Full lớn nhất vùng: Giảm sâu nhất toàn mạng -2.132 đơn Full (từ 28.919 về 26.787 đơn).', 'bullet_bad'),
        ('  ➔ Insight & Cảnh báo: Thời tiết mưa bão vùng cao Đà Lạt làm giảm sức mua tại chỗ, kết hợp tắc nghẽn giao hàng tại điểm nóng Lâm Viên 2 khiến một số shop giảm gửi hàng qua GHN.', 'sub_warning'),
        
        ('🔍 INSIGHT ĐIỀU HÀNH & NGUY CƠ TIỀM ẨN:', 'hdr_in'),
        ('• Rủi ro "sống dựa vào sàn": AM Trương Quang Linh (TTS chiếm tới 34,8% sản lượng) và AM Huỳnh Thúc Duân (TTS chiếm 29,0%) đang phụ thuộc cực nặng vào TikTok Shop. Hai AM này bắt buộc phải kiểm soát giờ xuất tuyến sáng trước 08h15, nếu trễ tuyến là sàn phạt sập ODR ngay!', 'in_item'),
        ('• AM Lê Minh Lợi (Đà Lạt ngoại thành): Tê liệt sức bán khi Full giảm (-639 đơn) và TTS chạm đáy chỉ còn 207 đơn (giảm -75% WoW, chỉ chiếm 8% sản lượng), cần khối KD tái cơ cấu danh mục khách hàng.', 'in_item'),
        
        ('🎯 GIAO VIỆC & MỆNH LỆNH TÁC CHIẾN ĐÍCH DANH (W40):', 'hdr_ac'),
        ('• Khối KD & AM Nguyễn Thanh Long: Trong 24h phải tiếp xúc trực tiếp 3 shop live lớn nhất Cam Ranh để tìm hiểu nguyên nhân mất 1.186 đơn TTS và đưa chính sách giữ chân khẩn cấp.', 'ac_item'),
        ('• AM Lê Văn Trường & Khối Vận Hành: Kích hoạt Cap Volume tại bưu cục Lâm Viên 2 để giải phóng nhanh tồn đọng, lấy lại uy tín dịch vụ với các nhà vườn Đà Lạt."', 'ac_item')
    ]
    
    for text, kind in lines:
        p = cell.add_paragraph()
        p.paragraph_format.line_spacing = 1.25
        
        if kind == 'intro':
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(text)
            r.italic = True
            r.font.name = 'Arial'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(15, 23, 42)
        elif kind == 'hdr_good':
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.bold = True
            r.font.name = 'Arial'
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(5, 150, 105)
        elif kind == 'bullet_good':
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            parts = text.split(':', 1) if ':' in text else [text, '']
            r_pre = p.add_run(parts[0] + (":" if parts[1] else ""))
            r_pre.bold = True
            r_pre.font.name = 'Arial'
            r_pre.font.size = Pt(9.5)
            r_pre.font.color.rgb = RGBColor(15, 23, 42)
            if parts[1]:
                r_post = p.add_run(parts[1])
                r_post.font.name = 'Arial'
                r_post.font.size = Pt(9.5)
                r_post.font.color.rgb = RGBColor(30, 41, 59)
        elif kind == 'sub_insight':
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(text)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(5, 150, 105)
            r.bold = True
        elif kind == 'hdr_bad':
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.bold = True
            r.font.name = 'Arial'
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(220, 38, 38)
        elif kind == 'bullet_bad':
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            parts = text.split(':', 1) if ':' in text else [text, '']
            r_pre = p.add_run(parts[0] + (":" if parts[1] else ""))
            r_pre.bold = True
            r_pre.font.name = 'Arial'
            r_pre.font.size = Pt(9.5)
            r_pre.font.color.rgb = RGBColor(15, 23, 42)
            if parts[1]:
                r_post = p.add_run(parts[1])
                r_post.font.name = 'Arial'
                r_post.font.size = Pt(9.5)
                r_post.font.color.rgb = RGBColor(30, 41, 59)
        elif kind == 'sub_warning':
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(text)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(220, 38, 38)
            r.bold = True
        elif kind == 'hdr_in':
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(text)
            r.bold = True
            r.font.name = 'Arial'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(15, 76, 129)
        elif kind == 'in_item':
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(51, 65, 85)
        elif kind == 'hdr_ac':
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(text)
            r.bold = True
            r.font.name = 'Arial'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(194, 65, 12)
        elif kind == 'ac_item':
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(51, 65, 85)

    doc.save(doc_path)
    print(f"Updated highlight script in {doc_path} successfully!")

targets = [
    'KICH_BAN_THUYET_TRINH_MOI_NHAT.docx',
    'KICH_BAN_THUYET_TRINH_W39_INSIGHT_CHUYEN_SAU.docx',
    'KICH_BAN_THUYET_TRINH_W39_INSIGHT_CHUYEN_SAU_CHINH_SUA.docx',
    'KICH_BAN_THUYET_TRINH_W39_NAM_TRUNG_BO.docx'
]
for t in targets:
    update_san_luong_highlight(t)
