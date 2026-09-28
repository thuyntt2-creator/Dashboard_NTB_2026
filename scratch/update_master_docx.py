import sys
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import shutil

sys.stdout.reconfigure(encoding='utf-8')

src_template = r"C:\Users\lap4all\Downloads\KICH_BAN_THUYET_TRINH_W38_NAM_TRUNG_BO (3).docx"
if not os.path.exists(src_template):
    print(f"Error: {src_template} not found!")
    sys.exit(1)

doc = docx.Document(src_template)

# 1. Update Title Header paragraphs
for p in doc.paragraphs:
    if "TUẦN W38" in p.text:
        p.text = p.text.replace("TUẦN W38", "TUẦN W39")
    if "14/09/2026 – 20/09/2026" in p.text:
        p.text = p.text.replace("14/09/2026 – 20/09/2026", "21/09/2026 – 27/09/2026")
    if "(W38)" in p.text:
        p.text = p.text.replace("(W38)", "(W39)")

# 2. Update Table 0 (Section I: Tổng Quan Điều Hành Vùng W39)
cell0 = doc.tables[0].rows[0].cells[0]
cell0.paragraphs[0].text = "🗣️ LỜI MỞ ĐẦU & TỔNG QUAN ĐIỀU HÀNH VÙNG TUẦN W39:"

sec1_lines = [
    "📍 1. BẢNG CHỈ SỐ TOÀN VÙNG & 5 TỈNH THÀNH (W39 vs W38):",
    "• Khánh Hòa: 92.696 đơn (-1,0% WoW) | %GTC 58,6% (+3,6%p) | %ODR 91,8% | TTS: 18.043 đơn (+449 đơn WoW) ➔ Giữ vững vị trí số 1 sản lượng toàn vùng!",
    "• Lâm Đồng: 88.287 đơn (-7,8% WoW) | Dẫn đầu TTS vùng (19.283 đơn, chiếm 26,5%) | %GTC 47,4% | %ODR 84,7% | Tồn aging cao cần giải tỏa.",
    "• Bình Thuận: 80.817 đơn (-5,5% WoW) | %GTC 67,5% | %ODR 96,3% (Top 2 vùng) | TTS tăng bùng nổ: 18.146 đơn (+2.273 đơn WoW, +14,3%!).",
    "• Đắk Nông: 33.883 đơn (-2,8% WoW) | %GTC 48,8% (+2,0%p) | %ODR 88,5% | TTS: 9.016 đơn (+702 đơn WoW, +8,4%).",
    "• Ninh Thuận: 33.242 đơn (-2,0% WoW) | Quán quân %GTC toàn vùng (68,0%) | Quán quân %ODR toàn vùng (96,4%) | TTS: 8.293 đơn (+935 đơn WoW, +12,7%).",
    "➔ TOÀN VÙNG: 328.925 đơn Full (-4,3%) | TikTok Shop đạt đỉnh 72.781 đơn (+5,9% WoW, chiếm 22,1% tỷ trọng) | %GTC Full 56,60% (+0,92%p) | %GTC TTS 57,45% (+3,55%p) | %ODR 90,79% | Rớt luân chuyển giảm sâu về 1,52% (-1,81%p WoW) | Truy thu phát sinh 306,9 Tr ₫ trên 4.013 đơn.",
    "👤 2. BÓC TÁCH CHI TIẾT HIỆU SUẤT THEO QUẢN LÝ VẬN HÀNH (AM):",
    "Kính chào Ban Giám Đốc và toàn thể 18 anh chị Quản lý Vận hành (AM). Nhìn vào bức tranh tuần W39, sự phân hóa giữa các AM thể hiện rất rõ nét:",
    "• Nhóm AM dẫn đầu sản lượng và gánh tải: AM Nguyễn Duy Long tiếp tục là đầu tàu số 1 với 42.506 đơn Full (12,9% toàn vùng) và 10.584 đơn TTS (vô địch toàn mạng); AM Thái Thị Thanh Thư bứt phá ngoạn mục tăng +2.227 đơn Full (đạt 34.765 đơn, vươn lên Top 2 toàn vùng!).",
    "• Nhóm AM giữ chất lượng giao hàng (%GTC) xuất sắc trên 67%: AM Nguyễn Ngọc Khánh, AM Cao Thị Thanh Thủy, AM Nguyễn Đỗ Minh Nghĩa và AM Nguyễn Duy Long tiếp tục bảo vệ vững chắc chuẩn SLA xanh.",
    "• Điểm sáng vượt trội khâu vận tải: Tỷ lệ rớt luân chuyển toàn vùng giảm ngoạn mục từ 3,32% xuống 1,52% (-1,81%p WoW, giảm hơn 54% lượng đơn rớt) nhờ siết chặt kỷ luật bàn giao ca KTC.",
    "• Hai vấn đề nổi cộm cần can thiệp gấp: Số tiền truy thu phát sinh 306,9 triệu đồng trên 4.013 đơn do buông lỏng cân đo; và danh sách bưu cục cảnh báo tăng lên 15 điểm nóng cần điều tiết tải ngay trong tuần W40.",
    "🔍 INSIGHT & ĐIỂM NGHẼN THỰC TẾ TẠI CÁC AM:\n",
    "• TikTok Shop tăng tốc lên 72.781 đơn (+5,9% WoW, chiếm 22,1% toàn vùng), khẳng định các AM phải ưu tiên tuyệt đối nguồn lực lấy hàng và giao ca sáng cho phân khúc này.",
    "• %GTC TTS tăng vọt +3,55%p lên 57,45% chứng minh giải pháp thúc đẩy giao hàng TMĐT đã đi đúng hướng, nhưng Lâm Đồng và Đắk Nông vẫn dưới 50% do vướng đơn tồn aging.",
    "⚠️ CẢNH BÁO ĐỎ TẬP TRUNG THEO AM:\n",
    "• Khoản tiền truy thu phát sinh 306,9 triệu đồng (tập trung tại Kim Chi 58,3 Tr, Bích Nga 50,7 Tr, Văn Trường 50,6 Tr, Văn Phước 27,1 Tr) đe dọa trực tiếp dòng tiền.",
    "• AM Nguyễn Thanh Long (Cam Ranh) hụt -1.186 đơn TTS (-40% WoW), cần kiểm tra ngay hiện tượng mất shop live sang đối thủ.",
    "🎯 GIAO VIỆC & NHIỆM VỤ TÁC CHIẾN ĐÍCH DANH CHO TỪNG AM:\n",
    "• Toàn thể 18 AM quán triệt ngay 3 trọng tâm W40: Triển khai Cap Volume giải tỏa bưu cục nghẽn, kiểm soát 100% cân đo tại quầy và đẩy mạnh thu tiền COD qua VietQR."
]

for p in cell0.paragraphs[1:]:
    p._p.getparent().remove(p._p)

for line in sec1_lines:
    p = cell0.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.text = line

# 3. Update Table 1 (Section II: SẢN LƯỢNG — MASTERCLASS INSIGHT & ACTION SCRIPT)
cell1 = doc.tables[1].rows[0].cells[0]
cell1.paragraphs[0].text = "🗣️ KỊCH BẢN THUYẾT TRÌNH SẢN LƯỢNG CHUYÊN SÂU (3 CHARTS TAB II):"

sec2_lines = [
    "Kính thưa Ban Giám Đốc và các anh chị Quản lý Vận hành (AM),",
    "",
    "Nhìn vào 3 biểu đồ sản lượng trên màn hình, tổng sản lượng toàn vùng tuần W39 đạt 328.925 đơn, giảm -4,3% so với tuần trước. Rất dễ để chúng ta ngộ nhận rằng sức mua thị trường đang nguội đi. Nhưng khi đi sâu vào bản chất dòng hàng, câu chuyện thực tế lại hoàn toàn trái ngược:",
    "",
    "Thị trường không hề mất đi, mà đang diễn ra một cuộc DỊCH CHUYỂN DÒNG HÀNG Ồ ẠT SANG TIKTOK SHOP. Tuần này, TikTok Shop của Nam Trung Bộ lập đỉnh kỷ lục mới với 72.781 đơn (+5,9% WoW, chiếm tới 22,1% tổng sản lượng toàn mạng). Sự bùng nổ này đẩy độ phân hóa giữa 18 AM lên mức gay gắt nhất từ trước đến nay: AM nào nắm bắt được nhịp livestream thì thắng lớn, còn AM nào chậm chân thì lập tức bị hụt tải và mất shop vào tay đối thủ!",
    "",
    "🏆 1. HAI TẤM GƯƠNG BỨT PHÁ DẪN ĐẦU MẠNG LƯỚI (HIGHLIGHT TỐT NHẤT):",
    "",
    "• AM Thái Thị Thanh Thư (Khánh Hòa) – Quán quân tăng trưởng toàn mạng (+2.227 đơn, đạt 34.765 đơn, +6,8% WoW):",
    "  Giữa lúc toàn vùng đi xuống, chị Thư là AM duy nhất lội ngược dòng tăng trưởng mạnh mẽ, chính thức vượt qua Lâm Đồng để đưa cụm Nha Trang vươn lên vị trí Top 2 sản lượng vùng.",
    "  🔍 Insight thành công đằng sau con số: Chị Thư không dàn trải mà tập trung 'đánh chiếm' chuỗi shop và đại lý bán lẻ nội đô Nha Trang. Lợi thế địa bàn đồng bằng với mật độ dân cư dày đặc giúp bưu tá giao phát cực nhanh 2-3 lượt/ngày, vòng quay đơn thần tốc, tỷ lệ GTC cao và hoàn toàn không bị ảnh hưởng bởi mùa mưa đèo dốc. Tốc độ giao hàng vượt trội chính là 'thỏi nam châm' giúp chị Thư giữ chân khách ruột và hút trọn đơn từ các đối thủ!",
    "",
    "• AM Nguyễn Duy Long (Bình Thuận) – 'Cỗ máy gánh tải' số 1 & Phá kỷ lục TikTok Shop (+1.379 đơn TTS, đạt 10.584 đơn | Full: 42.506 đơn, gánh 13% toàn vùng):",
    "  Anh Long là AM đầu tiên của Nam Trung Bộ phá mốc 10.000 đơn TikTok Shop chỉ trong 1 tuần.",
    "  🔍 Insight thành công đằng sau con số: Bình Thuận bứt phá không phải do may mắn, mà là kết quả của chiến lược First-mile xuất sắc. Anh Long chủ động bắt tay liên kết với các tổng kho nông hải sản Phan Thiết chạy livestream. Đặc thù của đơn livestream là 'chốt đơn là phải lấy ngay trong 2 tiếng'; anh Long đã bố trí các ca lấy hàng First-mile linh hoạt từ 15h đến 17h chiều, gom trọn hàng đưa lên xe KTC xuất bến trong đêm. Chính cam kết lấy hàng thần tốc này đã khiến các chủ shop live lớn quyết định dồn 100% sản lượng cho GHN!",
    "",
    "🚨 2. HAI ĐIỂM NÓNG BÁO ĐỘNG ĐỎ SỤT GIẢM NGUY HIỂM (HIGHLIGHT TỆ NHẤT):",
    "",
    "• AM Nguyễn Thanh Long (Cam Ranh) – Cú rơi tự do nguy hiểm nhất toàn mạng (-1.186 đơn TTS, sụt -40,5% WoW, rơi từ 2.930 xuống chỉ còn 1.744 đơn):",
    "  Chỉ trong 7 ngày, bưu cục Cam Ranh bốc hơi gần một nửa sản lượng TikTok Shop!",
    "  🔍 Sự thật phũ phàng & Root-cause: ĐÂY LÀ HỒI CHUÔNG BÁO ĐỘNG VỀ NGUY CƠ MẤT SHOP VÀO TAY ĐỐI THỦ SPX VÀ J&T! Bóc tách thực tế tại Cam Ranh cho thấy: Khâu lấy hàng First-mile ca chiều bị trễ giờ xuất xe, cộng với việc nhân viên bưu cục xử lý truy thu cân đo kích thước cứng nhắc, gây tranh cãi gay gắt với chủ shop. Hai shop livestream thời trang lớn nhất khu vực bức xúc, lập tức cắt luồng và ký hợp đồng chuyển hàng sang đối thủ ngay trong tuần. Mất gần 1.200 đơn mỗi tuần tương đương mất trắng hơn 30 triệu đồng doanh thu cước/tuần và đe dọa trực tiếp thị phần của GHN tại Cam Ranh!",
    "",
    "• AM Lê Văn Trường (Lâm Đồng) – Đơn vị sụt giảm Full hàng sâu nhất vùng (-2.132 đơn Full, rơi từ 28.919 về 26.787 đơn):",
    "  Sụt giảm nặng nề tại cụm Đà Lạt - Đơn Dương.",
    "  🔍 Insight bản chất nhân quả: VẬN HÀNH LAST-MILE TỆ HẠI TRỰC TIẾP GIẾT CHẾT SẢN LƯỢNG FIRST-MILE! Việc bưu cục Lâm Viên 2 bị nghẽn giao hàng nghiêm trọng (tỷ lệ GTC rơi tự do xuống 19,5%), hàng tồn aging chất đống làm phát sinh hàng trăm khiếu nại giao trễ. Các nhà vườn hoa, dâu tây và shop đặc sản Đà Lạt bị người mua đánh giá 1 sao, dẫn đến việc họ mất niềm tin vào GHN và chủ động cắt giảm sản lượng gửi đi, chuyển bớt đơn sang các hãng vận chuyển khác để tự cứu lấy uy tín kinh doanh!",
    "",
    "⚠️ 3. CẢNH BÁO RỦI RO HỆ THỐNG 'SỐNG DỰA VÀO SÀN':",
    "• Báo động tới AM Trương Quang Linh (TTS chiếm tới 34,8% sản lượng) và AM Huỳnh Thúc Duân (chiếm 29,0% sản lượng):",
    "  Tỷ trọng đơn sàn quá lớn trên địa bàn bán sơn địa/đồi núi xa xôi là 'con dao hai lưỡi'. Khách mua hàng livestream thường đặt theo cảm xúc bốc đồng; nếu bưu tá xuất tuyến trễ sau 08h30 sáng thì tỷ lệ từ chối nhận/boom hàng sẽ vọt lên 25 - 30%. Điều này không chỉ gây lãng phí chi phí vận chuyển giao hoàn mà còn khiến shop bị sàn hạ điểm tín nhiệm vận hành!",
    "",
    "🎯 4. MỆNH LỆNH TÁC CHIẾN TUẦN W40 — GIAO VIỆC ĐÍCH DANH:",
    "• 1. Khối Kinh doanh phối hợp cùng AM Nguyễn Thanh Long: Trong vòng 24 giờ tới, phải thành lập tổ công tác đến gặp trực tiếp 2 chủ shop thời trang lớn tại Cam Ranh, tháo gỡ dứt điểm khiếu nại cước cân đo, cam kết mở tuyến lấy hàng riêng trước 17h00. Mục tiêu tuần W40 phải giành lại tối thiểu 800 đơn TikTok Shop!",
    "• 2. Khối Vận Hành & AM Lê Văn Trường: Kích hoạt ngay cơ chế Cap Volume tại bưu cục Lâm Viên 2, tập trung toàn lực giải tỏa dứt điểm hàng tồn trong 3 ngày đầu tuần, lấy lại cam kết giao đúng giờ để khôi phục niềm tin của khách hàng Đà Lạt.",
    "• 3. Toàn thể 18 AM: Quán triệt kỷ luật First-mile: 100% đơn TikTok Shop phải được quét nhập kho bưu cục trước 17h30 để đưa lên chuyến xe KTC tối. Tuyệt đối không để một đơn hàng sàn nào nằm lại bưu cục sang ngày hôm sau!"
]

for p in cell1.paragraphs[1:]:
    p._p.getparent().remove(p._p)

for line in sec2_lines:
    p = cell1.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.text = line

# 4. Update Table 15 (Section XVI: 15 Bưu cục cảnh báo bất ổn W38 vs W39)
cell15 = doc.tables[15].rows[0].cells[0]
cell15.paragraphs[0].text = "🗣️ DANH SÁCH 15 BƯU CỤC CẢNH BÁO BẤT ỔN (W38 vs W39) & 3 MỆNH LỆNH TÁC CHIẾN:"

sec16_lines = [
    "📍 1. BIẾN ĐỘNG SỐ LƯỢNG TOÀN VÙNG (TĂNG TỪ 13 LÊN 15 BƯU CỤC):",
    "• Toàn vùng ghi nhận 15 bưu cục rơi vào diện cảnh báo bất ổn (%GTC < 45% hoặc giảm sâu dưới 70% mốc đỉnh lịch sử), tăng ròng +2 bưu cục so với 13 bưu cục ở tuần W38.",
    "• Cơ cấu dịch chuyển: 4 bưu cục đã thoát cảnh báo thành công, 6 bưu cục mới xuất hiện, và 9 bưu cục mãn tính kẹt cả 2 tuần liên tiếp.",
    "👤 2. BÓC TÁCH CHI TIẾT 3 NHÓM BƯU CỤC CẢNH BÁO THEO AM PHỤ TRÁCH:",
    "Kính thưa Ban Giám Đốc, nhìn sâu vào danh sách bưu cục bất ổn tuần W39:",
    "• Nhóm 1: Tuyên dương 4 Bưu cục đã giải tỏa & thoát cảnh báo thành công:",
    "  1. BC Xuân Hương - Đà Lạt (AM Lê Văn Trường): Nâng %GTC lên >50%, giải tỏa 100% backlog.",
    "  2. BC Cam Linh (AM Phan Đình Duy): Phục hồi nhân sự bưu tá và giờ xuất tuyến.",
    "  3. BC Nhân Cơ (AM Đỗ Duy Khang): Kiểm soát tốt luồng hàng nông thôn.",
    "  4. BC Tây Nha Trang (AM Hồng Bích Nga): Tối ưu ca phát sáng trước 11h00.",
    "• Nhóm 2: Báo động đỏ 6 Bưu cục mới rơi vào diện cảnh báo W39:",
    "  1. BC D'Ran (AM Phan Thành Long): %GTC đạt 36,86% (cảnh báo 6 ngày), địa hình đèo dốc sạt lở.",
    "  2. BC Lang Biang 2 (AM Trương Tuấn Anh): %GTC giảm về 35,80% (cảnh báo 6 ngày), TTS dồn ứ cục bộ.",
    "  3. BC Trường Xuân (AM Đỗ Duy Khang): %GTC tụt về 38,22% (cảnh báo 8 ngày), tỷ lệ KLL cao.",
    "  4. BC Đông Gia Nghĩa (AM Đỗ Duy Khang): %GTC đạt 44,89% (chớm dưới ngưỡng 45%), gán ca 2 yếu.",
    "  5. BC Bắc Cam Ranh (AM Phan Đình Duy): %GTC đạt 44,68% (đã tích lũy 31 ngày cảnh báo).",
    "  6. BC Phú Quý (AM Nguyễn Đình Uy): %GTC đạt 50,84%, tuy nhiên rơi sâu dưới 70% so với mốc lịch sử 88,55%.",
    "• Nhóm 3: Điểm nóng mãn tính 9 Bưu cục kẹt dai dẳng cả 2 tuần W38 và W39:",
    "  - Báo động nặng nhất mạng: BC Lâm Viên - Đà Lạt 2 (AM Trương Tuấn Anh) rơi tự do từ 46,2% xuống 19,55% (-26,6%p, bưu tá tê liệt ca chiều).",
    "  - BC Đức Trọng 1 (AM Phan Thành Long): Kẹt ở 23,78% (backlog tồn đọng >300 đơn).",
    "  - BC Quảng Tín (AM Đỗ Duy Khang): Đạt 26,17% (tích lũy kỷ lục >100 ngày cảnh báo).",
    "  - BC Đơn Dương (AM Phan Thành Long): Giảm mạnh về 29,01% (giảm -16,3%p WoW).",
    "  - Cùng 5 BC: Lang Biang 1 (30,6%), Kiến Đức (32,5%), Tân Hà Lâm Hà (36,8%), Tuy Đức (38,5%), Di Linh (40,6%).",
    "🔍 INSIGHT & ĐIỂM NGHẼN THỰC TẾ TẠI CÁC AM:\n",
    "• Tồn đọng và rớt %GTC ở nhóm 9 bưu cục mãn tính chủ yếu do thắt nút cổ chai ở tỷ lệ gán Ca 2 trưa và địa bàn đồi núi xa xôi.",
    "• Việc Lâm Viên 2 tụt xuống 19,55% kéo lùi toàn bộ chỉ số giao hàng của cụm thành phố Đà Lạt.",
    "⚠️ CẢNH BÁO ĐỎ TẬP TRUNG THEO AM:\n",
    "• Nguy cơ vỡ tải dây chuyền tại cụm Đà Lạt - Đơn Dương - Đức Trọng trong các đợt sale đầu tháng 10 nếu không điều tiết tải ngay.",
    "🎯 GIAO VIỆC & NHIỆM VỤ TÁC CHIẾN ĐÍCH DANH CHO TỪNG AM:\n",
    "• 1. Áp dụng cơ chế Cap Volume: Giảm 25% hạn mức chia chọn về Lâm Viên 2 và Quảng Tín trong 3 ngày đầu tuần để bưu tá dọn sạch tồn kho.",
    "• 2. Chi viện bưu tá ca sáng: Điều chuyển 4 bưu tá cứng hỗ trợ phát lượt 1 trước 11h00 tại Lâm Viên 2 và Đức Trọng 1.",
    "• 3. Cam kết W40: Quyết tâm kéo tối thiểu 5 bưu cục thoát cảnh báo, đưa danh sách toàn vùng về dưới 10 bưu cục."
]

for p in cell15.paragraphs[1:]:
    p._p.getparent().remove(p._p)

for line in sec16_lines:
    p = cell15.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.text = line

# Save to destination master
out_master = 'KICH_BAN_THUYET_TRINH_MOI_NHAT.docx'
doc.save(out_master)
print(f"Generated {out_master} successfully! (Size: {os.path.getsize(out_master)} bytes)")

copies = [
    'KICH_BAN_THUYET_TRINH_W39_INSIGHT_CHUYEN_SAU.docx',
    'KICH_BAN_THUYET_TRINH_W39_INSIGHT_CHUYEN_SAU_CHINH_SUA.docx',
    'KICH_BAN_THUYET_TRINH_W39_NAM_TRUNG_BO.docx'
]
for c in copies:
    shutil.copyfile(out_master, c)
    print(f"Copied to {c}")
