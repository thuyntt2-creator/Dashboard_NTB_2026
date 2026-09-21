import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_borders_and_shading(cell, border_color="F97316", fill_color="FFF7ED"):
    # Shading
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))
    
    # Borders: thick orange left, none for top/bottom/right
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="none"/>
        <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
        <w:bottom w:val="none"/>
        <w:right w:val="none"/>
    </w:tcBorders>
    '''
    cell._tc.get_or_add_tcPr().append(parse_xml(borders_xml))
    
    # Margins (padding)
    mar_xml = f'''
    <w:tcMar {nsdecls("w")}>
        <w:top w:w="160" w:type="dxa"/>
        <w:bottom w:w="160" w:type="dxa"/>
        <w:left w:w="240" w:type="dxa"/>
        <w:right w:w="240" w:type="dxa"/>
    </w:tcMar>
    '''
    cell._tc.get_or_add_tcPr().append(parse_xml(mar_xml))

def format_run(run, font_name="Times New Roman", font_size_pt=11, bold=False, italic=False, color_rgb=(0x11, 0x18, 0x27)):
    run.font.name = font_name
    run.font.size = Pt(font_size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def add_callout_box(doc, section_title, speech_title, speech_paragraphs, insights=None, warnings=None, actions=None):
    # Section title
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(14)
    p_sec.paragraph_format.space_after = Pt(6)
    p_sec.paragraph_format.keep_with_next = True
    r_sec = p_sec.add_run(section_title)
    format_run(r_sec, font_size_pt=13, bold=True, color_rgb=(0x0F, 0x4C, 0x81)) # Navy
    
    # 1x1 table
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Set full table width
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_borders_and_shading(cell, border_color="F97316", fill_color="FFF7ED")
    
    # First paragraph: Speech title
    p_first = cell.paragraphs[0]
    p_first.paragraph_format.space_before = Pt(2)
    p_first.paragraph_format.space_after = Pt(4)
    p_first.paragraph_format.line_spacing = 1.15
    r_title = p_first.add_run(speech_title)
    format_run(r_title, font_size_pt=11.5, bold=True, color_rgb=(0xEA, 0x58, 0x0C)) # Orange
    
    # Body speech paragraphs
    for sp in speech_paragraphs:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(sp)
        format_run(r, font_size_pt=11, bold=False, color_rgb=(0x1F, 0x29, 0x37))
    
    # Optional Insight
    if insights:
        p_in = cell.add_paragraph()
        p_in.paragraph_format.space_before = Pt(6)
        p_in.paragraph_format.space_after = Pt(3)
        p_in.paragraph_format.line_spacing = 1.15
        r_in_tag = p_in.add_run("🔍 INSIGHT BẢN CHẤT & GỐC RỄ NGUYÊN NHÂN:\n")
        format_run(r_in_tag, font_size_pt=11, bold=True, color_rgb=(0x0F, 0x4C, 0x81))
        for ins in insights:
            p_bullet = cell.add_paragraph()
            p_bullet.paragraph_format.space_before = Pt(1)
            p_bullet.paragraph_format.space_after = Pt(2)
            p_bullet.paragraph_format.line_spacing = 1.15
            p_bullet.paragraph_format.left_indent = Inches(0.2)
            r_b = p_bullet.add_run(f"• {ins}")
            format_run(r_b, font_size_pt=10.5, color_rgb=(0x1F, 0x29, 0x37))

    # Optional Warnings
    if warnings:
        p_w = cell.add_paragraph()
        p_w.paragraph_format.space_before = Pt(6)
        p_w.paragraph_format.space_after = Pt(3)
        p_w.paragraph_format.line_spacing = 1.15
        r_w_tag = p_w.add_run("⚠️ CẢNH BÁO ĐỎ & NGUY CƠ TIỀM ẨN:\n")
        format_run(r_w_tag, font_size_pt=11, bold=True, color_rgb=(0xDC, 0x26, 0x26)) # Red
        for w in warnings:
            p_wbullet = cell.add_paragraph()
            p_wbullet.paragraph_format.space_before = Pt(1)
            p_wbullet.paragraph_format.space_after = Pt(2)
            p_wbullet.paragraph_format.line_spacing = 1.15
            p_wbullet.paragraph_format.left_indent = Inches(0.2)
            r_wb = p_wbullet.add_run(f"• {w}")
            format_run(r_wb, font_size_pt=10.5, color_rgb=(0x99, 0x1B, 0x1B))

    # Optional Actions
    if actions:
        p_act = cell.add_paragraph()
        p_act.paragraph_format.space_before = Pt(6)
        p_act.paragraph_format.space_after = Pt(3)
        p_act.paragraph_format.line_spacing = 1.15
        r_act_tag = p_act.add_run("🎯 QUYẾT SÁCH HÀNH ĐỘNG & MỆNH LỆNH TÁC CHIẾN:\n")
        format_run(r_act_tag, font_size_pt=11, bold=True, color_rgb=(0x15, 0x80, 0x3D)) # Green
        for a in actions:
            p_abullet = cell.add_paragraph()
            p_abullet.paragraph_format.space_before = Pt(1)
            p_abullet.paragraph_format.space_after = Pt(2)
            p_abullet.paragraph_format.line_spacing = 1.15
            p_abullet.paragraph_format.left_indent = Inches(0.2)
            r_ab = p_bullet = p_abullet.add_run(f"• {a}")
            format_run(r_ab, font_size_pt=10.5, color_rgb=(0x14, 0x53, 0x2D))
            
    # Spacer after table
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(4)

print("Helper definitions loaded successfully.")
