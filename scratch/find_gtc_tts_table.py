import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W40_2026.xlsx', data_only=True)
ws = wb['02_GTC tong']

print("=== INSPECT 02_GTC tong FULL ===")
for r in range(1, ws.max_row + 1):
    c1 = ws.cell(r, 1).value
    if c1 and any(kw in str(c1).lower() for kw in ['theo', 'tỉnh', 'am', 'tts', 'tổng quan']):
        row_vals = [str(ws.cell(r, c).value or '')[:20] for c in range(1, 10)]
        print(f"Row {r:2d}: {row_vals}")

# Check columns further right (e.g. col J to Z)
print("\n--- Check right columns in 02_GTC tong ---")
for r in range(1, 15):
    for c in range(10, ws.max_column + 1):
        v = ws.cell(r, c).value
        if v and any(kw in str(v).lower() for kw in ['tts', 'theo am', 'tỉnh']):
            print(f"Cell ({r}, {c}): {v}")
