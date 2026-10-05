import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

file_path = r"C:\Users\lap4all\Downloads\kịch bản.docx"

try:
    doc = docx.Document(file_path)
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    print(f"Total tables: {len(doc.tables)}")
    
    print("\n--- ALL PARAGRAPHS IN kịch bản.docx ---")
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t:
            print(f"P{i:02d}: {t}")
            
    print("\n--- TABLES IN kịch bản.docx ---")
    for ti, tbl in enumerate(doc.tables):
        print(f"\n=== Table {ti+1} ({len(tbl.rows)} rows, {len(tbl.columns)} cols) ===")
        for r_i, r in enumerate(tbl.rows):
            for c_i, c in enumerate(r.cells):
                ct = c.text.strip()
                if ct:
                    print(f"  [R{r_i}C{c_i}]: {ct[:300]}...")
except Exception as e:
    print("Error reading docx:", e)
