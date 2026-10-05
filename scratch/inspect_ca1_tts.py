import openpyxl
import json
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('BaoCao_Tuan_NTB_W40_2026.xlsx', data_only=True)

for sheet_name in ['03_GTC Ca1+Ton', '03b_GTC Ca1 thuan', '04_GTC Ca2']:
    ws = wb[sheet_name]
    print(f"\n==================== SHEET: {sheet_name} ====================")
    for r in range(1, ws.max_row + 1):
        c1 = str(ws.cell(r, 1).value or '')
        if any(kw in c1.lower() for kw in ['tổng quan', 'chỉ tiêu', 'full hàng', 'tts', 'theo am', 'theo tỉnh']):
            vals = [str(ws.cell(r, c).value or '')[:18] for c in range(1, 8)]
            print(f"Row {r:2d}: {vals}")

# Also print AM & Tinh TTS in 03_GTC Ca1+Ton and 03b_GTC Ca1 thuan
for sheet_name in ['03_GTC Ca1+Ton', '03b_GTC Ca1 thuan']:
    ws = wb[sheet_name]
    print(f"\n--- DETAILED TTS IN {sheet_name} ---")
    for r in range(1, ws.max_row + 1):
        c1 = str(ws.cell(r, 1).value or '')
        if 'tts' in c1.lower():
            print(f"Found section header at row {r}: {c1}")
            for r2 in range(r, min(r + 25, ws.max_row + 1)):
                vals = [str(ws.cell(r2, c).value or '')[:18] for c in range(1, 8)]
                if any(vals):
                    print(f"  Row {r2:2d}: {vals}")
            break
