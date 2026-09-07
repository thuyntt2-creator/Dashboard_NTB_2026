import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
INPUT_EXCEL = r"c:\Users\lap4all\Desktop\New folder\downloaded_user_sheet.xlsx"

wb = openpyxl.load_workbook(INPUT_EXCEL, data_only=False)
ws = wb['LTC']
print("Formula in LTC cell A24:", ws['A24'].value)
print("Formula in LTC cell L24:", ws['L24'].value)
