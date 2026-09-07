import openpyxl

file_path = r"c:\Users\lap4all\Desktop\New folder\downloaded_user_sheet.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

output = []
sheets = ['sản lượng', 'gtcnew', 'gtc + tồn', 'ca2', 'LTC', 'ODR']

for sheet_name in sheets:
    ws = wb[sheet_name]
    output.append(f"\n=========================================\nSHEET: {sheet_name}\n=========================================")
    output.append("Col 1 AMs / Provinces:")
    for row in range(1, 40):
        val1 = ws.cell(row=row, column=1).value
        val9 = ws.cell(row=row, column=9).value if ws.max_column >= 9 else None
        if val1 is not None or val9 is not None:
            output.append(f"Row {row:2d}: Col 1: {val1} | Col 9: {val9}")

with open(r"c:\Users\lap4all\Desktop\New folder\am_orders_res.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(output))

print("Results written to am_orders_res.txt")
